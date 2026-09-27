"""Addenda: their files, their revisions, and what they changed (FR-DOC-05)."""

from __future__ import annotations

import io
import uuid
from collections.abc import Callable, Iterator
from datetime import date

import docx
import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from sqlalchemy import Engine, select
from sqlalchemy.orm import Session, sessionmaker

from firebid.api import documents as documents_api
from firebid.api.app import create_app
from firebid.api.deps import get_principal, get_session
from firebid.auth.provisioning import Principal
from firebid.db.models.core import AppUser, Bid, BidMember, Organisation, UserRole
from firebid.db.models.documents import Addendum, Document, SheetRevision
from firebid.domain.actors import Actor
from firebid.domain.state_machines import Role
from firebid.evals import synthetic
from firebid.ingest.scanning import AlwaysCleanScanner
from firebid.services.addenda import affected_items, register_addendum
from firebid.services.classification import classify_in_sandbox
from firebid.services.ingestion import Ingestor
from firebid.services.revisions import current_sheets
from firebid.services.sheets import process_document
from firebid.services.title_blocks import read_title_blocks
from firebid.settings import Settings
from firebid.storage.object_store import MemoryObjectStore

pytestmark = pytest.mark.req("FR-DOC-05")

NUMBER = "FP-L05-201"
SignIn = Callable[[Principal], TestClient]


@pytest.fixture(autouse=True)
def no_tiles(monkeypatch: pytest.MonkeyPatch) -> None:
    from firebid.services import sheets

    monkeypatch.setattr(sheets, "_render_low_levels", lambda *_args, **_kwargs: 0)


@pytest.fixture
def store() -> MemoryObjectStore:
    return MemoryObjectStore()


@pytest.fixture
def person(estimator: Actor) -> Actor:
    return estimator


def addendum(session: Session, bid: Bid, person: Actor, number: str, issued: date) -> Addendum:
    made = register_addendum(
        session, bid.id, number=number, issued_on=issued, summary=None, actor=person
    )
    session.flush()
    return made


def drawing(
    session: Session,
    bid: Bid,
    store: MemoryObjectStore,
    revision: str,
    *,
    within: Addendum | None = None,
) -> SheetRevision:
    plan, _ = synthetic.general_arrangement(sheet_number=NUMBER, revision=revision)
    stored = (
        Ingestor(
            session,
            store,
            AlwaysCleanScanner(),
            bid_id=bid.id,
            tender_package_id=within.tender_package_id if within else None,
        )
        .ingest(f"{NUMBER}.dxf", synthetic.dxf_bytes(plan))
        .stored[0]
    )
    [proposal] = read_title_blocks(
        session, store, stored, process_document(session, store, stored).sheets
    )
    return proposal


class TestAnAddendumsRevisions:
    def test_its_new_revision_is_linked_and_current(
        self, session: Session, bid: Bid, store: MemoryObjectStore, person: Actor
    ) -> None:
        original = drawing(session, bid, store, "R03")
        second = addendum(session, bid, person, "2", date(2026, 7, 1))

        revised = drawing(session, bid, store, "R04", within=second)

        assert revised.addendum_id == second.id
        assert original.addendum_id is None
        assert [row.revision_label for row, _ in current_sheets(session, bid.id)] == ["R04"]

    def test_the_affected_items_query_returns_what_it_changed(
        self, session: Session, bid: Bid, store: MemoryObjectStore, person: Actor
    ) -> None:
        """The acceptance test: an addendum links its revisions, and the query finds them."""
        drawing(session, bid, store, "R03")
        second = addendum(session, bid, person, "2", date(2026, 7, 1))
        drawing(session, bid, store, "R04", within=second)

        items = affected_items(session, second)

        assert [
            (item.kind, item.reference, item.revision, item.replaces, item.state) for item in items
        ] == [("sheet", NUMBER, "R04", "R03", "current")]

    def test_an_addendums_date_orders_what_the_scheme_cannot(
        self, session: Session, bid: Bid, store: MemoryObjectStore, person: Actor
    ) -> None:
        """X1 and Y1 are in no series; the addendum's later date says which came after."""
        drawing(session, bid, store, "X1")
        later = addendum(session, bid, person, "3", date(2026, 8, 1))

        drawing(session, bid, store, "Y1", within=later)

        assert [row.revision_label for row, _ in current_sheets(session, bid.id)] == ["Y1"]

    def test_a_revised_specification_in_an_addendum_is_affected_too(
        self, session: Session, bid: Bid, store: MemoryObjectStore, person: Actor
    ) -> None:
        second = addendum(session, bid, person, "2", date(2026, 7, 1))
        for revision, package in (("1", None), ("2", second.tender_package_id)):
            made = docx.Document()
            made.add_paragraph("PARTICULAR SPECIFICATION FOR FIRE PROTECTION SERVICES")
            made.add_paragraph(f"Revision {revision}")
            for clause in range(1, 7):
                made.add_paragraph(f"1.{clause} Requirement")
            buffer = io.BytesIO()
            made.save(buffer)
            stored = (
                Ingestor(
                    session, store, AlwaysCleanScanner(), bid_id=bid.id, tender_package_id=package
                )
                .ingest(f"Spec Rev {revision}.docx", buffer.getvalue())
                .stored[0]
            )
            classify_in_sandbox(session, stored, buffer.getvalue())

        [item] = affected_items(session, second)

        assert (item.kind, item.revision, item.replaces, item.state) == (
            "document",
            "2",
            "1",
            "current",
        )


class TestThroughTheApi:
    @pytest.fixture
    def app(
        self, engine: Engine, store: MemoryObjectStore, monkeypatch: pytest.MonkeyPatch
    ) -> FastAPI:
        monkeypatch.setattr(documents_api, "get_object_store", lambda: store)
        monkeypatch.setattr(documents_api, "get_scanner", AlwaysCleanScanner)
        factory = sessionmaker(bind=engine, expire_on_commit=False)

        def session_override() -> Iterator[Session]:
            with factory() as session:
                yield session
                session.commit()

        application = create_app(Settings(env="test"), health_checks={})
        application.dependency_overrides[get_session] = session_override
        return application

    @pytest.fixture
    def client(
        self, app: FastAPI, session: Session, organisation: Organisation, bid: Bid
    ) -> TestClient:
        person = AppUser(
            organisation_id=organisation.id,
            external_id=uuid.uuid4().hex,
            username="bella@firebid.test",
            display_name="Bella",
        )
        session.add(person)
        session.flush()
        session.add(UserRole(user_id=person.id, role=str(Role.BID_MANAGER)))
        session.add(BidMember(bid_id=bid.id, user_id=person.id, role=str(Role.BID_MANAGER)))
        session.commit()
        principal = Principal(
            user_id=person.id,
            organisation_id=organisation.id,
            username=person.username,
            display_name=person.display_name,
            roles=frozenset({str(Role.BID_MANAGER)}),
        )
        app.dependency_overrides[get_principal] = lambda: principal
        return TestClient(app)

    def test_an_addendum_is_registered_once(self, client: TestClient, bid: Bid) -> None:
        first = client.post(
            f"/bids/{bid.id}/addenda", json={"number": "2", "issued_on": "2026-07-01"}
        )
        again = client.post(f"/bids/{bid.id}/addenda", json={"number": "2"})

        assert first.status_code == 201, first.text
        assert again.status_code == 409
        assert [row["number"] for row in client.get(f"/bids/{bid.id}/addenda").json()] == ["2"]

    def test_files_uploaded_with_it_go_into_its_package(
        self, client: TestClient, bid: Bid, session: Session
    ) -> None:
        created = client.post(f"/bids/{bid.id}/addenda", json={"number": "2"}).json()
        plan, _ = synthetic.general_arrangement(revision="R05")

        response = client.post(
            f"/bids/{bid.id}/documents",
            files=[("files", (f"{NUMBER}.dxf", synthetic.dxf_bytes(plan), "image/vnd.dxf"))],
            data={"addendum_id": created["id"]},
        )

        assert response.status_code == 201, response.text
        session.expire_all()
        document = session.execute(select(Document)).scalar_one()
        added = session.get(Addendum, uuid.UUID(created["id"]))
        assert added is not None
        assert document.tender_package_id == added.tender_package_id

    def test_another_bids_addendum_is_not_found(self, client: TestClient, bid: Bid) -> None:
        plan, _ = synthetic.general_arrangement()

        response = client.post(
            f"/bids/{bid.id}/documents",
            files=[("files", (f"{NUMBER}.dxf", synthetic.dxf_bytes(plan), "image/vnd.dxf"))],
            data={"addendum_id": str(uuid.uuid4())},
        )

        assert response.status_code == 404

    def test_the_affected_items_are_served(self, client: TestClient, bid: Bid) -> None:
        created = client.post(f"/bids/{bid.id}/addenda", json={"number": "4"}).json()

        response = client.get(f"/bids/{bid.id}/addenda/{created['id']}/affected")

        assert response.status_code == 200
        assert response.json() == []
