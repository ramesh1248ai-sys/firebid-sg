# ruff: noqa: F811  (the fixtures below are imported by name and requested by name)
"""Bills of quantities through the database: built from the verified takeoff, the client's
bill read and mapped to it, reconciled, exported and traced (FR-BOQ-01 to 06, FR-ADM-03).

The takeoff is the synthetic tender `test_qto` measures; the client's bill is
`synthetic_boq`, the same installation as a quantity surveyor billed it.
"""

from __future__ import annotations

import io
import json
import uuid
import zipfile
from datetime import UTC, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any, cast

import pytest
from sqlalchemy import select, text
from sqlalchemy.orm import Session

from firebid.ai_gateway.config import load_config
from firebid.ai_gateway.metering import PostgresMeter
from firebid.ai_gateway.providers.fake import FakeAdapter
from firebid.ai_gateway.router import Router
from firebid.boq import reconcile
from firebid.boq.xlsx_patch import sheet_parts
from firebid.db.models.commercial import BoqLine, BoqLineSource
from firebid.db.models.core import Bid, Organisation
from firebid.db.models.documents import Document
from firebid.db.models.takeoff import QtoItem
from firebid.db.models.workflow import Approval
from firebid.domain.actors import Actor
from firebid.domain.state_machines import Role
from firebid.domain.values import Money
from firebid.evals import synthetic_boq
from firebid.ingest.scanning import AlwaysCleanScanner
from firebid.services import boq, qto
from firebid.services.classification import classify_in_sandbox
from firebid.services.ingestion import Ingestor
from firebid.storage.object_store import MemoryObjectStore
from tests.db.test_qto import tender  # noqa: F401
from tests.db.test_sheet_views import SignIn, app, member, sign_in  # noqa: F401
from tests.db.test_symbol_mapping import no_tiles, store  # noqa: F401

pytestmark = pytest.mark.usefixtures("no_tiles")

CONFIG = """
version: 1
providers:
  primary:
    kind: fake
    approved_data_classes: [internal, confidential]
models:
  first:
    provider: primary
    model_id: mapper-1
    capabilities: [structured_output]
    prices:
      - {effective_from: 2026-01-01, input_per_mtok: 5.0, output_per_mtok: 25.0}
routes:
  boq_mapping:
    requires: [structured_output]
    data_class: confidential
    models: [first]
    reasoning: low
    prompt: boq_mapping
"""


def people(session: Session, organisation: Organisation, bid: Bid) -> tuple[Actor, Actor]:
    estimator = member(session, organisation, bid, "esther", Role.ESTIMATOR).actor()
    senior = member(session, organisation, bid, "sam", Role.SENIOR_ESTIMATOR).actor()
    return estimator, senior


def verify_takeoff(session: Session, bid: Bid) -> None:
    """As the reviewers would: every live item verified (P1-08 tests the doing of it)."""
    for item in qto.live_items(session, bid.id):
        item.state = "verified"
    session.commit()


def client_workbook(session: Session, bid: Bid, store: MemoryObjectStore) -> Document:
    payload = synthetic_boq.client_boq().payload
    document = (
        Ingestor(session, store, AlwaysCleanScanner(), bid_id=bid.id)
        .ingest("Bill of Quantities - Fire Sprinkler.xlsx", payload)
        .stored[0]
    )
    classify_in_sandbox(session, document, payload)
    document.doc_type = "boq"
    session.commit()
    return document


def built(session: Session, bid: Bid, actor: Actor) -> list[BoqLine]:
    made = boq.build_company_boq(session, bid.id, actor)
    session.commit()
    return boq.lines_of(session, made)


def line_for(session: Session, bid: Bid, description: str) -> BoqLine:
    """The company line an item of the takeoff went into."""
    current = boq.current_boq(session, bid.id)
    assert current is not None
    [line] = session.execute(
        select(BoqLine)
        .join(BoqLineSource, BoqLineSource.boq_line_id == BoqLine.id)
        .join(QtoItem, QtoItem.id == BoqLineSource.qto_item_id)
        .where(BoqLine.boq_id == current.id, QtoItem.description == description)
    ).scalars()
    return line


def router(session: Session, tmp_path: Path, reply: dict[str, Any]) -> Router:
    path = tmp_path / "llm.yaml"
    path.write_text(CONFIG, encoding="utf-8")
    return Router(
        config=load_config(path),
        adapters={"primary": FakeAdapter("primary").reply(json.dumps(reply))},
        backoff_base_seconds=0,
        sleep=lambda _s: None,
        meter=PostgresMeter(session),
    )


@pytest.mark.req("FR-BOQ-01")
@pytest.mark.req("FR-ADM-03")
class TestCompanyBoq:
    def test_nothing_is_built_from_unverified_quantities(
        self, session: Session, tender: Bid, organisation: Organisation
    ) -> None:
        estimator, _ = people(session, organisation, tender)

        with pytest.raises(boq.BoqError, match="verify the takeoff first"):
            boq.build_company_boq(session, tender.id, estimator)

    def test_every_line_carries_its_quantity_and_the_items_behind_it(
        self, session: Session, tender: Bid, organisation: Organisation
    ) -> None:
        estimator, _ = people(session, organisation, tender)
        verify_takeoff(session, tender)

        lines = built(session, tender, estimator)

        items = qto.live_items(session, tender.id)
        sources = session.execute(
            select(BoqLineSource.qto_item_id).where(BoqLineSource.bid_id == tender.id)
        ).scalars()
        assert sorted(map(str, sources)) == sorted(str(i.id) for i in items), (
            "every verified item is in exactly one line"
        )
        assert line_for(session, tender, "Sprinkler, pendent").quantity == Decimal(16)
        assert line_for(session, tender, "Pipe, DN50, branch").unit == "m"
        assert all(line.line_key and line.item_no for line in lines)
        current = boq.current_boq(session, tender.id)
        assert current is not None
        assert (current.template_key, current.template_version) == ("company_standard", 1)

    def test_building_again_is_a_new_version_with_the_same_line_keys(
        self, session: Session, tender: Bid, organisation: Organisation
    ) -> None:
        estimator, _ = people(session, organisation, tender)
        verify_takeoff(session, tender)
        first = built(session, tender, estimator)

        second = built(session, tender, estimator)

        assert [line.line_key for line in first] == [line.line_key for line in second]
        versions = session.execute(
            text("SELECT version, is_current FROM boq WHERE bid_id = :b ORDER BY version"),
            {"b": tender.id},
        ).all()
        assert [tuple(v) for v in versions] == [(1, False), (2, True)]

    def test_a_template_edit_is_a_new_version_and_the_old_one_is_kept(
        self, session: Session, tender: Bid, organisation: Organisation
    ) -> None:
        estimator, senior = people(session, organisation, tender)
        verify_takeoff(session, tender)
        seeded = boq.template(session, organisation.id)
        definition = dict(seeded.definition)
        definition["descriptions"] = {
            **cast(dict[str, str], definition["descriptions"]),
            "sprinkler": "Sprinkler head, {orientation}",
        }

        edited = boq.edit_template(
            session,
            organisation.id,
            seeded.key,
            definition=definition,
            actor=senior,
            note="our house wording",
        )
        built(session, tender, estimator)

        history = boq.template_history(session, organisation.id, seeded.key)
        assert [(t.version, t.retired_at is not None) for t in history] == [(1, True), (2, False)]
        assert edited.status == "confirmed"
        current = boq.current_boq(session, tender.id)
        assert current is not None and current.template_version == 2


@pytest.mark.req("FR-BOQ-02")
class TestClientBoq:
    def test_registering_a_client_bill_queues_its_reading(
        self, session: Session, bid: Bid, store: MemoryObjectStore
    ) -> None:
        document = client_workbook(session, bid, store)
        boq.queue_reading(session, document)
        session.commit()

        queued = session.execute(
            text(
                "SELECT count(*) FROM procrastinate_jobs WHERE task_name = 'boq.read' "
                "AND args->>'document_id' = :document"
            ),
            {"document": str(document.id)},
        ).scalar_one()
        assert queued >= 1

    def test_every_line_of_the_bill_is_read_and_the_hidden_sheet_is_not(
        self, session: Session, bid: Bid, store: MemoryObjectStore
    ) -> None:
        document = client_workbook(session, bid, store)

        [sheet] = boq.read_client_boq(session, store, document)
        session.commit()

        fixture = synthetic_boq.client_boq()
        lines = [line for s, line in boq.client_lines(session, bid.id)]
        assert sheet.sheet_name == synthetic_boq.BILL
        assert sheet.header_row == synthetic_boq.HEADER_ROW
        assert {line.item_no for line in lines} == set(fixture.cells)
        by_item = {line.item_no: line for line in lines}
        assert by_item["B3"].quantity == Decimal(70)
        assert by_item["D1"].kind in ("provisional", "lump_sum")
        assert (by_item["A1"].rate_cell, by_item["A1"].amount_cell) == fixture.cells["A1"][1:]

    def test_the_original_workbook_is_never_written(
        self, session: Session, bid: Bid, store: MemoryObjectStore
    ) -> None:
        document = client_workbook(session, bid, store)
        before = store.get(document.storage_key)

        boq.read_client_boq(session, store, document)
        session.commit()

        assert store.get(document.storage_key) == before


@pytest.mark.req("FR-BOQ-02")
@pytest.mark.req("FR-BOQ-03")
class TestMappingAndReconciliation:
    @pytest.fixture
    def mapped(
        self,
        session: Session,
        tender: Bid,
        organisation: Organisation,
        store: MemoryObjectStore,
        tmp_path: Path,
    ) -> Bid:
        estimator, _ = people(session, organisation, tender)
        verify_takeoff(session, tender)
        built(session, tender, estimator)
        boq.read_client_boq(session, store, client_workbook(session, tender, store))
        session.commit()
        return tender

    def test_the_rules_map_the_plain_lines_and_every_proposal_waits_for_a_person(
        self, session: Session, mapped: Bid
    ) -> None:
        boq.propose_mappings(session, mapped.id)
        session.commit()

        rows = boq.mappings(session, mapped.id)
        lines = {line.id: line for _, line in boq.client_lines(session, mapped.id)}
        by_item = {lines[k].item_no: row for k, row in rows.items()}
        assert {row.state for row in rows.values()} == {"proposed"}
        assert (
            by_item["A1"].boq_line_key == line_for(session, mapped, "Sprinkler, pendent").line_key
        )
        assert (
            by_item["B3"].boq_line_key == line_for(session, mapped, "Pipe, DN50, branch").line_key
        )
        assert by_item["A1"].method == "rule"

    def test_the_model_maps_what_the_rules_left_and_a_made_up_line_is_no_answer(
        self, session: Session, mapped: Bid, tmp_path: Path
    ) -> None:
        refs = {line.item_no: str(line.id) for _, line in boq.client_lines(session, mapped.id)}
        boq.propose_mappings(session, mapped.id)
        rows = boq.mappings(session, mapped.id)
        left = [
            item
            for item, ref in refs.items()
            if rows[uuid.UUID(ref)].provenance.get("awaiting_model")
        ]
        # Nobody drew a flow switch: the rules have nothing to map it to. It is still in front
        # of a person ("nothing matched") while the model is asked.
        assert left == ["C3"]
        assert (rows[uuid.UUID(refs["C3"])].boq_line_key, rows[uuid.UUID(refs["C3"])].state) == (
            None,
            "proposed",
        )
        reply = {
            "mappings": [
                {"line": refs["C3"], "maps_to": "no-such-line", "confidence": 0.9, "reason": "x"}
            ]
        }

        boq.propose_mappings(session, mapped.id, router(session, tmp_path, reply))
        session.commit()

        row = boq.mappings(session, mapped.id)[uuid.UUID(refs["C3"])]
        assert (row.method, row.boq_line_key, row.confidence) == ("model", None, 0.0)
        assert not row.provenance.get("awaiting_model")
        assert row.provenance["model"] == "mapper-1"

    def test_a_person_confirms_corrects_or_rejects_and_it_is_audited(
        self, session: Session, mapped: Bid, organisation: Organisation
    ) -> None:
        estimator, _ = people(session, organisation, mapped)
        boq.propose_mappings(session, mapped.id)
        rows = list(boq.mappings(session, mapped.id).values())
        drops = line_for(session, mapped, "Sprinkler drop, DN25 (vertical, not drawn)")

        boq.decide_mapping(session, rows[0], estimator, decision="confirm")
        boq.decide_mapping(
            session, rows[1], estimator, decision="correct", boq_line_key=drops.line_key
        )
        boq.decide_mapping(session, rows[2], estimator, decision="reject", note="not ours")
        session.commit()

        assert [r.state for r in rows[:3]] == ["confirmed", "confirmed", "rejected"]
        assert rows[1].method == "person" and rows[1].boq_line_id == drops.id
        audited = session.execute(
            text(
                "SELECT count(*) FROM audit_event WHERE bid_id = :b "
                "AND action LIKE 'client BOQ mapping:%'"
            ),
            {"b": mapped.id},
        ).scalar_one()
        assert audited == 3

    def test_a_mapping_survives_the_boq_being_built_again(
        self, session: Session, mapped: Bid, organisation: Organisation
    ) -> None:
        estimator, _ = people(session, organisation, mapped)
        boq.propose_mappings(session, mapped.id)
        [row] = [r for r in boq.mappings(session, mapped.id).values() if r.boq_line_key][:1]
        boq.decide_mapping(session, row, estimator, decision="confirm")

        built(session, mapped, estimator)

        assert row.state == "confirmed"
        current = boq.current_boq(session, mapped.id)
        assert current is not None
        target = session.get(BoqLine, row.boq_line_id)
        assert target is not None and target.boq_id == current.id

    def test_reconciliation_flags_variances_over_the_threshold_and_what_one_side_lacks(
        self, session: Session, mapped: Bid
    ) -> None:
        boq.propose_mappings(session, mapped.id)
        session.commit()

        rows = {
            r.client_item or r.line_description: r for r in boq.reconciliation(session, mapped.id)
        }

        branch = rows["B3"]  # 70 m billed, 72 m measured
        assert (branch.variance, branch.variance_percent) == (Decimal(2), Decimal("2.9"))
        assert branch.flagged is False
        assert branch.qto_items and branch.evidence_links
        assert rows["C3"].kind == "client_only" and rows["C3"].flagged
        measured_only = [r for r in rows.values() if r.kind == "measured_only"]
        assert measured_only and all(r.flagged for r in measured_only), "the tees nobody billed"
        assert {r.client_item for r in boq.clarification_candidates(session, mapped.id)} >= {"C3"}

    def test_the_reconciliation_report_is_a_workbook_with_the_flags(
        self, session: Session, mapped: Bid
    ) -> None:
        from openpyxl import load_workbook

        boq.propose_mappings(session, mapped.id)

        book = load_workbook(io.BytesIO(boq.reconciliation_workbook(session, mapped.id)))

        sheet = book["Reconciliation"]
        header = [cell.value for cell in sheet[1]]
        assert header[:5] == [
            "Client ref",
            "Client item",
            "Client description",
            "Client unit",
            "Client qty",
        ]
        clarify = header.index("Clarify")
        assert any(row[clarify] == "yes" for row in sheet.iter_rows(min_row=2, values_only=True))


@pytest.mark.req("FR-BOQ-03")
def test_120_m_billed_against_128_4_m_measured_is_flagged() -> None:
    found = reconcile.variance(Decimal(120), Decimal("128.4"))

    assert (found.difference, found.percent, found.flagged) == (
        Decimal("8.4"),
        Decimal("7.0"),
        True,
    )


@pytest.mark.req("FR-BOQ-04")
class TestExports:
    def test_the_priced_client_workbook_changes_only_the_rate_and_plain_amount_cells(
        self,
        session: Session,
        tender: Bid,
        organisation: Organisation,
        store: MemoryObjectStore,
    ) -> None:
        estimator, _ = people(session, organisation, tender)
        verify_takeoff(session, tender)
        built(session, tender, estimator)
        document = client_workbook(session, tender, store)
        boq.read_client_boq(session, store, document)
        boq.propose_mappings(session, tender.id)
        for row in boq.mappings(session, tender.id).values():
            if row.boq_line_key:
                boq.decide_mapping(session, row, estimator, decision="confirm")
        line_for(session, tender, "Sprinkler, pendent").unit_rate = Money.of("85.50")
        session.commit()

        priced = boq.priced_client_workbook(session, store, tender.id, document.id)

        original = store.get(document.storage_key)
        before, after = zipfile.ZipFile(io.BytesIO(original)), zipfile.ZipFile(io.BytesIO(priced))
        assert before.namelist() == after.namelist()
        changed = {n for n in before.namelist() if before.read(n) != after.read(n)}
        bill = sheet_parts(before)[synthetic_boq.BILL]
        # Only the bill's sheet: the client's workbook already asks Excel to recalculate.
        assert changed == {bill}
        from openpyxl import load_workbook

        sheet = load_workbook(io.BytesIO(priced))[synthetic_boq.BILL]
        _, rate_cell, amount_cell = synthetic_boq.client_boq().cells["A1"]
        assert sheet[rate_cell].value == 85.5
        assert str(sheet[amount_cell].value).startswith("="), "Excel works out the amount"

    def test_the_company_boq_exports_with_its_trace(
        self, session: Session, tender: Bid, organisation: Organisation
    ) -> None:
        from openpyxl import load_workbook

        estimator, _ = people(session, organisation, tender)
        verify_takeoff(session, tender)
        built(session, tender, estimator)

        sheet = load_workbook(io.BytesIO(boq.company_workbook(session, tender.id)))["BOQ"]

        rows = list(sheet.iter_rows(min_row=2, values_only=True))
        assert any(row[1] and "pendent" in str(row[1]).lower() and row[3] == 16 for row in rows)


@pytest.mark.req("FR-BOQ-05")
class TestTrace:
    def test_an_untraced_line_blocks_g1_and_g2_until_it_is_marked(
        self, session: Session, tender: Bid, organisation: Organisation
    ) -> None:
        estimator, senior = people(session, organisation, tender)
        verify_takeoff(session, tender)
        [first, *_] = built(session, tender, estimator)
        session.execute(text("DELETE FROM boq_line_source WHERE boq_line_id = :l"), {"l": first.id})
        session.commit()

        assert [r["id"] for r in qto.g1_blockers(session, tender.id).untraced_lines] == [
            str(first.id)
        ]
        with pytest.raises(boq.BoqError, match="G1 is not approved; 1 BOQ line"):
            boq.approve_g2(session, tender, senior, "senior_estimator")

        boq.mark_line(session, first, "provisional", estimator, "the client's allowance")
        session.commit()

        assert qto.g1_blockers(session, tender.id).untraced_lines == []
        assert boq.g2_blockers(session, tender.id).untraced_lines == []

    def test_g2_is_recorded_once_g1_is_approved_and_every_line_is_traced(
        self, session: Session, tender: Bid, organisation: Organisation
    ) -> None:
        estimator, senior = people(session, organisation, tender)
        verify_takeoff(session, tender)
        # G1 as approved (P1-07 and P1-08 test the approving): this fixture's takeoff still
        # has a duplicate group and unmapped symbols a person would settle first.
        assert senior.id is not None
        session.add(
            Approval(
                bid_id=tender.id,
                gate="G1",
                decision="approved",
                approver_id=senior.id,
                approver_role="senior_estimator",
                decided_at=datetime.now(UTC),
            )
        )
        built(session, tender, estimator)
        boq.add_marked_line(
            session,
            tender.id,
            estimator,
            description="Allow for testing and commissioning",
            unit="sum",
            quantity=Decimal(1),
            marker="provisional",
            note="the client's provisional sum D1",
            amount=Decimal(5000),
        )

        approval = boq.approve_g2(session, tender, senior, "senior_estimator")
        session.commit()

        assert approval.gate == "G2" and approval.snapshot_hash

    def test_a_line_without_a_trace_needs_a_reason(
        self, session: Session, tender: Bid, organisation: Organisation
    ) -> None:
        estimator, _ = people(session, organisation, tender)
        verify_takeoff(session, tender)
        [first, *_] = built(session, tender, estimator)

        with pytest.raises(boq.BoqError, match="say why"):
            boq.mark_line(session, first, "lump_sum", estimator, " ")


@pytest.mark.req("FR-BOQ-06")
def test_conventions_are_set_per_tender_and_worded_for_the_qualifications(
    session: Session, bid: Bid, organisation: Organisation
) -> None:
    estimator = member(session, organisation, bid, "esther", Role.ESTIMATOR).actor()
    defaults = reconcile.defaults()
    choice = {
        "pipe_measurement": next(iter(reconcile.conventions()["pipe_measurement"]["options"]))
    }

    first = boq.set_conventions(session, bid.id, choice, estimator)
    second = boq.set_conventions(session, bid.id, {}, estimator)
    session.commit()

    assert (first.version, second.version) == (1, 2)
    assert set(first.settings) == set(defaults)
    assert boq.qualification_text(session, bid.id)
    with pytest.raises(boq.BoqError):
        boq.set_conventions(session, bid.id, {"pipe_measurement": "by eye"}, estimator)


@pytest.mark.req("FR-BOQ-01")
@pytest.mark.req("FR-BOQ-02")
@pytest.mark.req("FR-BOQ-03")
@pytest.mark.req("FR-BOQ-04")
@pytest.mark.req("FR-BOQ-05")
class TestApi:
    def test_build_map_reconcile_and_export_through_the_api(
        self,
        session: Session,
        tender: Bid,
        organisation: Organisation,
        store: MemoryObjectStore,
        sign_in: SignIn,
        monkeypatch: pytest.MonkeyPatch,
    ) -> None:
        monkeypatch.setattr("firebid.api.boq.get_object_store", lambda: store)
        principal = member(session, organisation, tender, "esther", Role.ESTIMATOR)
        client = sign_in(principal)
        verify_takeoff(session, tender)
        document = client_workbook(session, tender, store)
        boq.read_client_boq(session, store, document)
        session.commit()

        assert client.get(f"/bids/{tender.id}/boq").status_code == 404
        built_boq = client.post(f"/bids/{tender.id}/boq/build", json={})
        assert built_boq.status_code == 201, built_boq.text
        lines = built_boq.json()["lines"]
        assert all(line["traced"] and line["qto_items"] for line in lines)

        [sheet] = client.get(f"/bids/{tender.id}/boq/client").json()
        assert (sheet["status"], sheet["lines"]) == ("read", len(synthetic_boq.client_boq().cells))

        proposed = client.post(f"/bids/{tender.id}/boq/mappings/propose")
        assert proposed.status_code == 200, proposed.text
        pendent = next(m for m in proposed.json() if m["client_item"] == "A1")
        decided = client.post(
            f"/bids/{tender.id}/boq/mappings/{pendent['mapping_id']}",
            json={"decision": "confirm"},
        )
        assert decided.json()["state"] == "confirmed"
        queued = session.execute(
            text("SELECT count(*) FROM procrastinate_jobs WHERE task_name = 'boq.map'")
        ).scalar_one()
        assert queued >= 1, "the flow switch the rules could not map goes to the model"

        rows = client.get(f"/bids/{tender.id}/boq/reconciliation").json()
        assert any(r["client_item"] == "C3" and r["flagged"] for r in rows)
        for path in ("export.xlsx", "reconciliation.xlsx", f"client/{document.id}/priced.xlsx"):
            response = client.get(f"/bids/{tender.id}/boq/{path}")
            assert response.status_code == 200, (path, response.text)
            assert response.content[:2] == b"PK"

        g2 = client.get(f"/bids/{tender.id}/boq/g2").json()
        assert (g2["g1_approved"], g2["boq_built"], g2["untraced_lines"]) == (False, True, [])

    def test_only_a_senior_estimator_changes_a_template_or_approves_g2(
        self, session: Session, bid: Bid, organisation: Organisation, sign_in: SignIn
    ) -> None:
        estimator = sign_in(member(session, organisation, bid, "esther", Role.ESTIMATOR))
        [template] = estimator.get("/boq-templates").json()

        change = estimator.post(
            f"/boq-templates/{template['key']}", json={"definition": template["definition"]}
        )
        approve = estimator.post(f"/bids/{bid.id}/boq/g2/approve", json={})

        assert (change.status_code, approve.status_code) == (403, 403)
        senior = sign_in(member(session, organisation, bid, "sam", Role.SENIOR_ESTIMATOR))
        changed = senior.post(
            f"/boq-templates/{template['key']}",
            json={"definition": template["definition"], "note": "reviewed"},
        )
        assert changed.status_code == 200 and changed.json()["version"] == 2
        refused = senior.post(f"/bids/{bid.id}/boq/g2/approve", json={})
        assert refused.status_code == 409 and "G1 is not approved" in refused.text

    def test_conventions_through_the_api(
        self, session: Session, bid: Bid, organisation: Organisation, sign_in: SignIn
    ) -> None:
        client = sign_in(member(session, organisation, bid, "esther", Role.ESTIMATOR))

        before = client.get(f"/bids/{bid.id}/boq/conventions").json()
        after = client.put(
            f"/bids/{bid.id}/boq/conventions", json={"settings": {"fittings": "deemed_included"}}
        )

        assert before["version"] is None
        assert after.json()["version"] == 1
        assert "deemed included in the rates for pipework" in after.json()["qualification_text"]
        bad = client.put(f"/bids/{bid.id}/boq/conventions", json={"settings": {"x": "y"}})
        assert bad.status_code == 422
