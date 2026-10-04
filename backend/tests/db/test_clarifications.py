# ruff: noqa: F811  (the fixtures below are imported by name and requested by name)
"""Tender clarifications through the database and the API (P2-06).

Two tenders. One has the extended specification and the basement car park plan whose notes
contradict it, so the analysis flags the seeded conflicts. The other has the synthetic
client bill mapped to the measured one, so the reconciliation flags its variances.
"""

from __future__ import annotations

import io
import json
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook
from sqlalchemy import select, text
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from firebid.ai_gateway.config import load_config
from firebid.ai_gateway.metering import PostgresMeter
from firebid.ai_gateway.providers.fake import FakeAdapter
from firebid.ai_gateway.router import Router
from firebid.api import clarifications as clarifications_api
from firebid.auth.provisioning import Principal
from firebid.clarifications.drafting import Candidate
from firebid.db.models.audit import AuditEvent
from firebid.db.models.clarifications import Clarification
from firebid.db.models.core import Bid, Organisation
from firebid.db.models.documents import Document
from firebid.db.models.workflow import HumanTask
from firebid.domain.actors import Actor
from firebid.domain.state_machines import Role
from firebid.evals import synthetic_qto
from firebid.evals.synthetic_network import NETWORK
from firebid.evals.synthetic_spec import CAR_PARK_NOTES
from firebid.ingest.scanning import AlwaysCleanScanner
from firebid.services import boq, qto
from firebid.services import clarifications as service
from firebid.services import spec_analysis as analysis
from firebid.services.detection import detect_bid
from firebid.storage.object_store import MemoryObjectStore
from tests.db.test_boq import built, client_workbook, people, verify_takeoff
from tests.db.test_detection_pipeline import confirm_legend
from tests.db.test_qto import tender  # noqa: F401
from tests.db.test_sheet_views import SignIn, app, member, sign_in  # noqa: F401
from tests.db.test_spec_analysis import SHEET, specification
from tests.db.test_symbol_mapping import from_consultant, no_tiles, read, store  # noqa: F401

pytestmark = pytest.mark.usefixtures("no_tiles")

CUTOFF = datetime(2026, 10, 20, 9, 0, tzinfo=UTC)
MODEL_CONFIG = """
version: 1
providers:
  primary:
    kind: fake
    approved_data_classes: [internal, confidential]
models:
  first:
    provider: primary
    model_id: drafter-1
    capabilities: [structured_output]
    prices:
      - {effective_from: 2026-01-01, input_per_mtok: 5.0, output_per_mtok: 25.0}
routes:
  clarification_draft:
    requires: [structured_output]
    data_class: confidential
    models: [first]
    reasoning: low
    prompt: clarification_draft
"""


@pytest.fixture
def spec_bid(session: Session, bid: Bid, store: MemoryObjectStore) -> Bid:
    """The specification read, the car park plan taken off, and the analysis run."""
    specification(session, bid, store)
    from_consultant(session, bid, NETWORK.name)
    plan, _ = synthetic_qto.car_park_plan(SHEET, CAR_PARK_NOTES)
    read(session, bid, store, SHEET, plan)
    confirm_legend(session, bid, store)
    detect_bid(session, store, bid.id)
    qto.recompute(session, bid.id)
    session.execute(text("DELETE FROM procrastinate_jobs"))
    bid.clarification_cutoff = CUTOFF
    analysis.analyse(session, bid.id)
    session.commit()
    return bid


@pytest.fixture
def bill_bid(
    session: Session, tender: Bid, organisation: Organisation, store: MemoryObjectStore
) -> Bid:
    """The synthetic tender with its client bill read and mapped to the measured bill."""
    estimator, _ = people(session, organisation, tender)
    verify_takeoff(session, tender)
    built(session, tender, estimator)
    boq.read_client_boq(session, store, client_workbook(session, tender, store))
    boq.propose_mappings(session, tender.id)
    tender.clarification_cutoff = CUTOFF
    session.commit()
    return tender


class Team:
    def __init__(self, session: Session, organisation: Organisation, bid: Bid) -> None:
        def one(name: str, role: Role) -> Principal:
            return member(session, organisation, bid, name, role)

        self.estimator_principal = one("edna", Role.ESTIMATOR)
        self.manager_principal = one("bella", Role.BID_MANAGER)
        self.designer_principal = one("dan", Role.DESIGN_MANAGER)
        self.estimator: Actor = self.estimator_principal.actor()
        self.manager: Actor = self.manager_principal.actor()
        self.designer: Actor = self.designer_principal.actor()


def conflict_of(session: Session, bid: Bid, what: str = "pipe_material") -> tuple[str, str]:
    found = next(
        item
        for item in service.candidates(session, bid)
        if item.kind == "spec_issue" and item.topic == what
    )
    return found.kind, found.ref


def issued(session: Session, bid: Bid, team: Team, row: Clarification) -> Clarification:
    service.transition(session, bid, row, "internal_review", team.estimator)
    if row.engineering_content:
        service.design_approval(session, bid, row, team.designer)
    service.transition(session, bid, row, "approved_to_issue", team.manager)
    service.transition(session, bid, row, "issued", team.manager)
    session.commit()
    return row


@pytest.mark.req("FR-RFI-02")
class TestDrafting:
    def test_a_seeded_spec_conflict_drafts_with_every_field_and_its_evidence(
        self, session: Session, spec_bid: Bid, organisation: Organisation
    ) -> None:
        team = Team(session, organisation, spec_bid)

        row = service.draft(session, spec_bid, [conflict_of(session, spec_bid)], team.estimator)
        session.commit()

        assert (service.label(row), row.kind, row.state) == (
            "TC-001",
            "tender_clarification",
            "draft",
        )
        assert row.subject.startswith("Pipe material:")
        assert row.project == "Example Commercial Tower"
        assert [(s["sheet_number"], s["revision"]) for s in row.sheets] == [(SHEET, "R01")]
        assert "Specification clause 2.1.3 states:" in row.problem
        assert f"Drawing {SHEET} revision R01 notes:" in row.problem
        assert "Please confirm which governs." in row.problem
        assert [(e["kind"], e["label"]) for e in row.evidence] == [
            ("clause", "Specification clause 2.1.3"),
            ("sheet", f"Drawing {SHEET}"),
        ]
        assert all(e["quote"] and e["link"] and e["revision"] for e in row.evidence)
        assert [o["text"] for o in row.options] == [
            "Recommendation: the specification governs (galvanised steel), and the drawing is "
            "to be revised",
            "Recommendation: the drawing governs (black steel), and the specification is to be "
            "amended",
        ]
        assert all(o["recommendation"] for o in row.options)
        # Potential impact is the estimator's to say; the reviewer follows from the content.
        assert (row.cost_impact, row.programme_impact) == ("", "")
        assert (row.required_reviewer, row.engineering_content) == ("design_manager", True)
        assert row.engineering_reason == "it concerns pipe material"
        # Once drafted, the issue is not offered again.
        assert ("spec_issue", conflict_of.__name__) != ("", "")
        assert all(c.topic != "pipe_material" for c in service.candidates(session, spec_bid))

    def test_a_bill_variance_drafts_with_its_evidence(
        self, session: Session, bill_bid: Bid, organisation: Organisation
    ) -> None:
        team = Team(session, organisation, bill_bid)
        variances = [c for c in service.candidates(session, bill_bid) if c.kind == "boq_variance"]
        flow = next(c for c in variances if "Flow switch" in c.subject)

        row = service.draft(session, bill_bid, [(flow.kind, flow.ref)], team.estimator)
        session.commit()

        assert row.subject.startswith("Bill item not found on the drawings: Flow switch, 150 mm")
        assert "The tender drawings show nothing it measures." in row.problem
        [evidence] = row.evidence
        assert evidence["kind"] == "client_boq_line"
        assert evidence["label"].startswith("Client bill ") and "C3" in evidence["quote"]
        assert (row.required_reviewer, row.engineering_content) == ("bid_manager", False)
        assert row.project and row.due_at is not None
        measured = next(
            c for c in variances if c.subject.startswith("Measured but not in the bill")
        )
        assert {e.kind for e in measured.evidence} == {"qto_item"}
        assert all(e.link for e in measured.evidence)

    def test_a_draft_with_no_evidence_is_refused_and_nothing_is_saved(
        self,
        session: Session,
        bill_bid: Bid,
        organisation: Organisation,
        monkeypatch: pytest.MonkeyPatch,
    ) -> None:
        team = Team(session, organisation, bill_bid)
        bare = Candidate(
            kind="boq_variance",
            ref="line:Z9",
            subject="Something nobody can point to",
            problem="Please confirm.",
            evidence=(),
        )
        monkeypatch.setattr(service, "candidates", lambda _session, _bid: [bare])

        with pytest.raises(service.ClarificationError, match="no evidence reference"):
            service.draft(session, bill_bid, [("boq_variance", "line:Z9")], team.estimator)

        assert service.register(session, bill_bid.id) == []

    def test_the_database_refuses_a_clarification_without_evidence(
        self, session: Session, spec_bid: Bid, organisation: Organisation
    ) -> None:
        team = Team(session, organisation, spec_bid)
        row = service.draft(session, spec_bid, [conflict_of(session, spec_bid)], team.estimator)
        session.commit()

        with pytest.raises(service.ClarificationError, match="at least one evidence reference"):
            service.edit(session, spec_bid, row, {"evidence": []}, team.estimator)
        with pytest.raises(IntegrityError, match="has_evidence"):
            session.execute(
                text("UPDATE clarification SET evidence = '[]'::jsonb WHERE id = :c"), {"c": row.id}
            )
        session.rollback()

    def test_the_model_rewords_a_draft_only_when_it_cites_the_draft_s_evidence(
        self, session: Session, spec_bid: Bid, organisation: Organisation, tmp_path: Path
    ) -> None:
        team = Team(session, organisation, spec_bid)
        row = service.draft(session, spec_bid, [conflict_of(session, spec_bid)], team.estimator)
        session.commit()
        before = (row.subject, row.problem)
        uncited = {
            "subject": "Pipe material in the car park",
            "problem": "Kindly advise.",
            "options": [],
            "evidence": [],
            "confidence": 0.9,
        }

        with pytest.raises(service.ClarificationError, match="nothing was saved"):
            service.draft_with_model(
                session, spec_bid, row, router(session, tmp_path, uncited), team.estimator
            )
        session.rollback()
        session.refresh(row)
        assert (row.subject, row.problem) == before

        with pytest.raises(service.ClarificationError, match="evidence the draft does not have"):
            service.draft_with_model(
                session,
                spec_bid,
                row,
                router(session, tmp_path, {**uncited, "evidence": [7], "problem": "Advise (2)."}),
                team.estimator,
            )
        session.rollback()
        session.refresh(row)

        cited = {
            **uncited,
            "problem": "Clause 2.1.3 requires galvanised pipework in the basement car park; "
            f"{SHEET} R01 notes black steel. Please confirm which governs.",
            "options": ["the specification governs"],
            "evidence": [0, 1],
        }
        row.subject = "Pipe material (basement)"  # another request: not the same model call
        service.draft_with_model(
            session, spec_bid, row, router(session, tmp_path, cited), team.estimator
        )
        session.commit()

        assert row.subject == "Pipe material in the car park"
        assert [o["text"] for o in row.options] == ["Recommendation: the specification governs"]
        assert row.drafting["method"] == "model" and row.drafting["model"] == "drafter-1"
        assert row.drafting["evidence_relied_on"] == [0, 1]
        assert len(row.evidence) == 2 and row.engineering_content is True


def router(session: Session, tmp_path: Path, reply: dict[str, Any]) -> Router:
    path = tmp_path / "llm.yaml"
    path.write_text(MODEL_CONFIG, encoding="utf-8")
    return Router(
        config=load_config(path),
        adapters={"primary": FakeAdapter("primary").reply(json.dumps(reply))},
        backoff_base_seconds=0,
        sleep=lambda _s: None,
        meter=PostgresMeter(session),
    )


@pytest.mark.req("FR-RFI-06")
class TestApproval:
    def test_an_engineering_clarification_needs_the_design_manager_before_issue(
        self, session: Session, spec_bid: Bid, organisation: Organisation
    ) -> None:
        team = Team(session, organisation, spec_bid)
        row = service.draft(session, spec_bid, [conflict_of(session, spec_bid)], team.estimator)
        service.transition(session, spec_bid, row, "internal_review", team.estimator)
        session.commit()

        with pytest.raises(
            service.ClarificationError, match="Design Manager's approval is missing"
        ):
            service.transition(session, spec_bid, row, "approved_to_issue", team.manager)
        with pytest.raises(service.ClarificationError, match="approved by the Design Manager"):
            service.design_approval(session, spec_bid, row, team.manager)
        with pytest.raises(
            service.ClarificationError, match="needs one of these roles: bid_manager"
        ):
            service.transition(session, spec_bid, row, "approved_to_issue", team.designer)

        service.design_approval(
            session, spec_bid, row, team.designer, qp_input_needed=True, note="QP to confirm"
        )
        service.transition(session, spec_bid, row, "approved_to_issue", team.manager)
        session.commit()

        assert row.state == "approved_to_issue"
        assert (row.design_approved_by, row.qp_input_needed) == (team.designer.label, True)
        assert row.approved_by == team.manager.label
        actions = [
            event.action
            for event in session.execute(
                select(AuditEvent).where(AuditEvent.entity_id == str(row.id))
            ).scalars()
        ]
        assert sorted(actions) == sorted(
            [
                "clarification: drafted",
                "clarification: send for internal review",
                "clarification: Design Manager approval",
                "clarification: approve to issue",
            ]
        )

    def test_a_plain_clarification_needs_only_the_bid_manager(
        self, session: Session, bill_bid: Bid, organisation: Organisation
    ) -> None:
        team = Team(session, organisation, bill_bid)
        flow = next(c for c in service.candidates(session, bill_bid) if "Flow switch" in c.subject)
        row = service.draft(session, bill_bid, [(flow.kind, flow.ref)], team.estimator)
        service.transition(session, bill_bid, row, "internal_review", team.estimator)

        with pytest.raises(
            service.ClarificationError, match="needs one of these roles: bid_manager"
        ):
            service.transition(session, bill_bid, row, "approved_to_issue", team.estimator)
        service.transition(session, bill_bid, row, "approved_to_issue", team.manager)

        assert row.state == "approved_to_issue" and row.design_approved_by is None

    def test_engineering_words_added_to_a_plain_draft_bring_the_design_manager_in(
        self, session: Session, bill_bid: Bid, organisation: Organisation
    ) -> None:
        team = Team(session, organisation, bill_bid)
        flow = next(c for c in service.candidates(session, bill_bid) if "Flow switch" in c.subject)
        row = service.draft(session, bill_bid, [(flow.kind, flow.ref)], team.estimator)

        service.edit(
            session,
            bill_bid,
            row,
            {"options": ["omit the flow switch, as the hydraulic design does not need it"]},
            team.estimator,
        )

        assert (row.engineering_content, row.required_reviewer) == (True, "design_manager")
        assert row.engineering_reason == "it mentions hydraulic"
        assert row.options[0]["text"].startswith("Recommendation: omit the flow switch")


@pytest.mark.req("FR-RFI-04")
class TestRegister:
    def test_due_dates_follow_the_cut_off(
        self, session: Session, spec_bid: Bid, organisation: Organisation
    ) -> None:
        team = Team(session, organisation, spec_bid)
        row = service.draft(session, spec_bid, [conflict_of(session, spec_bid)], team.estimator)

        # Two days before the cut-off of 20 October, to leave time to send it.
        assert row.due_at == datetime(2026, 10, 18, 9, 0, tzinfo=UTC)
        assert service.overdue(row, datetime(2026, 10, 18, 9, 1, tzinfo=UTC))
        assert not service.overdue(row, datetime(2026, 10, 17, tzinfo=UTC))

        spec_bid.clarification_cutoff = datetime(2026, 11, 3, 9, 0, tzinfo=UTC)
        later = service.draft(
            session, spec_bid, [conflict_of(session, spec_bid, "joining_method")], team.estimator
        )
        assert later.due_at == datetime(2026, 11, 1, 9, 0, tzinfo=UTC)
        assert service.label(later) == "TC-002"

    def test_a_response_links_back_and_raises_an_impact_task(
        self,
        session: Session,
        bill_bid: Bid,
        organisation: Organisation,
        store: MemoryObjectStore,
    ) -> None:
        team = Team(session, organisation, bill_bid)
        flow = next(c for c in service.candidates(session, bill_bid) if "Flow switch" in c.subject)
        row = service.draft(session, bill_bid, [(flow.kind, flow.ref)], team.estimator)

        with pytest.raises(service.ClarificationError, match="against an issued clarification"):
            service.record_response(session, bill_bid, row, team.estimator, summary="Too soon")
        issued(session, bill_bid, team, row)
        service.record_response(
            session,
            bill_bid,
            row,
            team.estimator,
            summary="The flow switch is at the L05 zone valve; see the revised schematic.",
            payload=b"%PDF-1.4\n% the consultant's reply\n",
            filename="TC-001 response.pdf",
            store=store,
            scanner=AlwaysCleanScanner(),
        )
        session.commit()

        assert row.state == "responded" and row.responded_at is not None
        document = session.get(Document, row.response_document_id)
        assert document is not None and document.filename == "TC-001 response.pdf"
        assert document.doc_type == "clarification_response"
        task = session.get(HumanTask, row.impact_task_id)
        assert task is not None
        assert (task.kind, task.state) == ("clarification.impact", "open")
        assert task.payload == {"clarification_id": str(row.id), "number": "TC-001"}
        assert task.title.startswith("Assess the impact of the response to TC-001")

        # It is closed only by saying what it changed; takeoff and pricing run again first.
        with pytest.raises(service.ClarificationError, match="those set this state"):
            service.transition(session, bill_bid, row, "closed_no_change", team.estimator)
        service.assess_impact(
            session,
            bill_bid,
            row,
            team.estimator,
            outcome="no_change",
            note="Already measured on FP-L05-201; the bill item maps to it.",
            rerun=True,
        )
        session.commit()

        assert row.state == "closed_no_change" and row.impact_by == team.estimator.label
        assert row.impact_detail["rerun"] is True
        assert row.impact_detail["takeoff"]["created"] == 0
        assert row.impact_detail["takeoff"]["unchanged"] > 0
        session.refresh(task)
        assert task.state == "done" and task.completed_by_id == team.estimator.id


@pytest.mark.req("FR-RFI-07")
class TestGroupingAndExport:
    def test_related_issues_are_proposed_together_and_merged_on_confirmation(
        self, session: Session, spec_bid: Bid, organisation: Organisation
    ) -> None:
        team = Team(session, organisation, spec_bid)

        groups = service.proposed_groups(session, spec_bid)
        related = next(group for group in groups if len(group.candidates) > 1)

        topics = {c.topic for c in related.candidates}
        assert {"pipe_material", "joining_method"} <= topics
        assert related.key == f"sprinkler|{SHEET}"
        # Proposing changes nothing: the group is one clarification only once confirmed.
        assert service.register(session, spec_bid.id) == []

        row = service.draft(
            session, spec_bid, [(c.kind, c.ref) for c in related.candidates], team.estimator
        )
        session.commit()

        assert row.subject == (
            f"{len(related.candidates)} queries on the sprinkler installation on {SHEET}"
        )
        assert len(row.problem.splitlines()) == len(related.candidates)
        assert len(service.sources(session, row)) == len(related.candidates)
        left = {(c.kind, c.ref) for c in service.candidates(session, spec_bid)}
        assert left.isdisjoint({(c.kind, c.ref) for c in related.candidates})
        with pytest.raises(service.ClarificationError, match="already in a clarification"):
            service.draft(
                session,
                spec_bid,
                [(related.candidates[0].kind, related.candidates[0].ref)],
                team.estimator,
            )

    def test_the_export_matches_the_chosen_template_and_is_only_a_download(
        self,
        session: Session,
        spec_bid: Bid,
        organisation: Organisation,
        sign_in: SignIn,
    ) -> None:
        team = Team(session, organisation, spec_bid)
        row = service.draft(session, spec_bid, [conflict_of(session, spec_bid)], team.estimator)
        session.commit()
        client = sign_in(team.manager_principal)

        response = client.get(
            f"/bids/{spec_bid.id}/clarifications/export",
            params={"template": "client_query_log", "format": "xlsx"},
        )
        unknown = client.get(
            f"/bids/{spec_bid.id}/clarifications/export", params={"template": "nobody"}
        )

        assert response.status_code == 200, response.text
        assert response.headers["content-disposition"].startswith("attachment;")
        sheet = load_workbook(io.BytesIO(response.content)).active
        assert sheet is not None
        rows = [list(r) for r in sheet.iter_rows(values_only=True)]
        assert rows[3] == [
            "Query No",
            "Drawing / Spec Ref",
            "Location",
            "Tenderer's Query",
            "Tenderer's Proposal",
            "Consultant's Response",
        ]
        assert rows[4][0] == "TC-001" and rows[4][1] == f"{SHEET} rev R01"
        assert rows[4][3] == row.problem
        assert str(rows[4][4]).startswith("Recommendation: the specification governs")
        assert unknown.status_code == 422
        # The download is the person's, and is recorded as theirs. Nothing else left.
        event = session.execute(
            select(AuditEvent).where(AuditEvent.action == "clarifications: downloaded")
        ).scalar_one()
        assert event.actor_label == team.manager.label
        assert event.after == {
            "template": "client_query_log",
            "format": "xlsx",
            "clarifications": ["TC-001"],
        }
        queued = session.execute(text("SELECT count(*) FROM procrastinate_jobs")).scalar_one()
        assert queued == 0, "no job is queued that could send it"
        # Issuing is a person saying they sent it: the platform sends nothing.
        issued(session, spec_bid, team, row)
        assert row.issued_by == team.manager.label
        assert session.execute(text("SELECT count(*) FROM procrastinate_jobs")).scalar_one() == 0


@pytest.mark.req("FR-RFI-05")
class TestQualifications:
    def test_unresolved_clarifications_become_proposed_qualifications_at_submission(
        self, session: Session, spec_bid: Bid, organisation: Organisation
    ) -> None:
        team = Team(session, organisation, spec_bid)
        sent = service.draft(session, spec_bid, [conflict_of(session, spec_bid)], team.estimator)
        unsent = service.draft(
            session, spec_bid, [conflict_of(session, spec_bid, "joining_method")], team.estimator
        )
        closed = service.draft(
            session,
            spec_bid,
            [next((c.kind, c.ref) for c in service.candidates(session, spec_bid))],
            team.estimator,
        )
        issued(session, spec_bid, team, sent)
        issued(session, spec_bid, team, closed)
        service.record_response(session, spec_bid, closed, team.estimator, summary="Confirmed.")
        service.assess_impact(
            session, spec_bid, closed, team.estimator, outcome="no_change", note="As priced."
        )
        session.commit()

        with pytest.raises(
            service.ClarificationError, match="needs one of these roles: bid_manager"
        ):
            service.prepare_submission(session, spec_bid, team.estimator)
        session.rollback()
        made = service.prepare_submission(session, spec_bid, team.manager)
        session.commit()

        by_clarification = {q.clarification_id: q for q in made}
        assert set(by_clarification) == {sent.id, unsent.id}
        assert all(q.state == "proposed" for q in made)
        asked = by_clarification[sent.id]
        assert asked.kind == "qualification"
        assert "tender clarification TC-001 was raised and is not resolved" in asked.text
        assert "Our offer is on the basis that the specification governs" in asked.text
        never = by_clarification[unsent.id]
        assert never.kind == "assumption" and "we have assumed that" in never.text
        assert (sent.state, unsent.state) == ("converted_to_qualification",) * 2
        assert closed.state == "closed_no_change"

        service.decide_qualification(
            session, spec_bid, asked, team.manager, decision="accepted", note="as drafted"
        )
        assert (asked.state, asked.decided_by) == ("accepted", team.manager.label)
        # Preparing again proposes nothing new.
        assert service.prepare_submission(session, spec_bid, team.manager) == []


@pytest.mark.req("FR-RFI-01")
class TestTenderClarificationsOnly:
    def test_a_construction_rfi_is_reserved_and_never_raised(
        self, session: Session, spec_bid: Bid, organisation: Organisation, sign_in: SignIn
    ) -> None:
        team = Team(session, organisation, spec_bid)
        kind, ref = conflict_of(session, spec_bid)
        client = sign_in(team.estimator_principal)

        with pytest.raises(service.ClarificationError, match="construction RFIs are out of scope"):
            service.draft(session, spec_bid, [(kind, ref)], team.estimator, kind="construction_rfi")
        offered = client.get(f"/bids/{spec_bid.id}/clarifications/candidates").json()
        refused = client.post(
            f"/bids/{spec_bid.id}/clarifications",
            json={"candidates": [{"kind": kind, "ref": ref}], "kind": "construction_rfi"},
        )
        drafted = client.post(
            f"/bids/{spec_bid.id}/clarifications",
            json={"candidates": [{"kind": kind, "ref": ref}]},
        )
        listed = client.get(f"/bids/{spec_bid.id}/clarifications").json()

        assert offered["kinds"] == ["tender_clarification"]
        assert refused.status_code == 422 and "out of scope" in refused.json()["detail"]
        assert drafted.status_code == 201, drafted.text
        assert drafted.json()["kind"] == "tender_clarification"
        assert [c["number"] for c in listed["clarifications"]] == ["TC-001"]
        assert listed["due_at"] == "2026-10-18T09:00:00Z"
        assert {t["key"] for t in listed["templates"]} == {"company_default", "client_query_log"}
        # The reserved kind exists in the data model, and nothing has used it.
        kinds = session.execute(text("SELECT DISTINCT kind FROM clarification")).scalars().all()
        assert kinds == ["tender_clarification"]


class TestApi:
    @pytest.mark.req("FR-RFI-04")
    @pytest.mark.req("FR-RFI-06")
    def test_a_clarification_goes_from_draft_to_closed_through_the_api(
        self,
        session: Session,
        spec_bid: Bid,
        organisation: Organisation,
        sign_in: SignIn,
        monkeypatch: pytest.MonkeyPatch,
    ) -> None:
        memory = MemoryObjectStore()
        monkeypatch.setattr(clarifications_api, "get_object_store", lambda: memory)
        monkeypatch.setattr(clarifications_api, "get_scanner", AlwaysCleanScanner)
        team = Team(session, organisation, spec_bid)
        kind, ref = conflict_of(session, spec_bid)
        base = f"/bids/{spec_bid.id}/clarifications"

        estimator = sign_in(team.estimator_principal)
        row = estimator.post(base, json={"candidates": [{"kind": kind, "ref": ref}]}).json()
        one = f"{base}/{row['id']}"
        edited = estimator.patch(one, json={"cost_impact": "Galvanised costs more per metre."})
        estimator.post(f"{one}/transition", json={"target": "internal_review"})
        too_soon = estimator.post(f"{one}/transition", json={"target": "approved_to_issue"})

        manager = sign_in(team.manager_principal)
        blocked = manager.post(f"{one}/transition", json={"target": "approved_to_issue"})
        not_theirs = manager.post(f"{one}/design-approval", json={})

        designer = sign_in(team.designer_principal)
        approved = designer.post(f"{one}/design-approval", json={"qp_input_needed": True})

        manager = sign_in(team.manager_principal)
        to_issue = manager.post(f"{one}/transition", json={"target": "approved_to_issue"})
        sent = manager.post(f"{one}/transition", json={"target": "issued"})

        estimator = sign_in(team.estimator_principal)
        responded = estimator.post(
            f"{one}/response",
            data={"summary": "The specification governs."},
            files={"file": ("reply.pdf", b"%PDF-1.4\n% reply\n", "application/pdf")},
        )
        closed = estimator.post(
            f"{one}/impact",
            json={"outcome": "incorporated", "note": "Rates changed to galvanised."},
        )

        assert edited.json()["cost_impact"] == "Galvanised costs more per metre."
        assert too_soon.status_code == 409 and "bid_manager" in too_soon.json()["detail"]
        assert blocked.status_code == 409
        assert "Design Manager's approval is missing" in blocked.json()["detail"]
        assert not_theirs.status_code == 403
        assert approved.status_code == 200 and approved.json()["qp_input_needed"] is True
        assert to_issue.json()["state_label"] == "Approved to issue"
        assert sent.json()["state"] == "issued" and sent.json()["issued_by"] == "Bella"
        assert responded.status_code == 200, responded.text
        assert responded.json()["response_document_id"] and responded.json()["impact_task_id"]
        assert closed.status_code == 200, closed.text
        assert closed.json()["state_label"] == "Closed: incorporated"
