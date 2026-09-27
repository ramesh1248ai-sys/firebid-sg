"""The registers through the API: reading them, exporting them, and deciding on them (FR-DOC-03)."""

from __future__ import annotations

import io
import uuid
from collections.abc import Callable, Iterator
from datetime import datetime

import docx
import openpyxl
import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from sqlalchemy import Engine, select
from sqlalchemy.orm import Session, sessionmaker

from firebid.api.app import create_app
from firebid.api.deps import get_principal, get_session
from firebid.auth.provisioning import Principal
from firebid.db.models.audit import AuditEvent
from firebid.db.models.core import AppUser, Bid, BidMember, Organisation, UserRole
from firebid.db.models.documents import SheetRevision, TitleBlockLayout
from firebid.domain.state_machines import Role
from firebid.evals import synthetic
from firebid.ingest.scanning import AlwaysCleanScanner
from firebid.services.classification import classify_in_sandbox
from firebid.services.ingestion import Ingestor
from firebid.services.registers import DOCUMENT_COLUMNS, DRAWING_COLUMNS
from firebid.services.sheets import process_document
from firebid.services.title_blocks import read_title_blocks
from firebid.settings import Settings
from firebid.storage.object_store import MemoryObjectStore

pytestmark = pytest.mark.req("FR-DOC-03")

SignIn = Callable[[Principal], TestClient]
NUMBER = "FP-L05-201"


@pytest.fixture(autouse=True)
def no_tiles(monkeypatch: pytest.MonkeyPatch) -> None:
    from firebid.services import sheets

    monkeypatch.setattr(sheets, "_render_low_levels", lambda *_args, **_kwargs: 0)


@pytest.fixture
def store() -> MemoryObjectStore:
    return MemoryObjectStore()


@pytest.fixture
def app(engine: Engine) -> FastAPI:
    factory = sessionmaker(bind=engine, expire_on_commit=False)

    def session_override() -> Iterator[Session]:
        with factory() as session:
            yield session
            session.commit()

    application = create_app(Settings(env="test"), health_checks={})
    application.dependency_overrides[get_session] = session_override
    return application


@pytest.fixture
def sign_in(app: FastAPI) -> SignIn:
    def _sign_in(principal: Principal) -> TestClient:
        app.dependency_overrides[get_principal] = lambda: principal
        return TestClient(app)

    return _sign_in


def member(
    session: Session, organisation: Organisation, bid: Bid, name: str, role: Role
) -> Principal:
    person = AppUser(
        organisation_id=organisation.id,
        external_id=uuid.uuid4().hex,
        username=f"{name}@firebid.test",
        display_name=name.title(),
    )
    session.add(person)
    session.flush()
    session.add(UserRole(user_id=person.id, role=str(role)))
    session.add(BidMember(bid_id=bid.id, user_id=person.id, role=str(role)))
    session.commit()
    return Principal(
        user_id=person.id,
        organisation_id=organisation.id,
        username=person.username,
        display_name=person.display_name,
        roles=frozenset({str(role)}),
    )


@pytest.fixture
def estimator(session: Session, organisation: Organisation, bid: Bid) -> Principal:
    return member(session, organisation, bid, "ethan", Role.ESTIMATOR)


@pytest.fixture
def design_manager(session: Session, organisation: Organisation, bid: Bid) -> Principal:
    return member(session, organisation, bid, "dora", Role.DESIGN_MANAGER)


def drawing(
    session: Session,
    bid: Bid,
    store: MemoryObjectStore,
    revision: str,
    *,
    number: str = NUMBER,
    filename: str | None = None,
    document: object | None = None,
) -> SheetRevision:
    if document is None:
        document, _ = synthetic.general_arrangement(sheet_number=number, revision=revision)
    payload = synthetic.dxf_bytes(document)  # type: ignore[arg-type]
    stored = (
        Ingestor(session, store, AlwaysCleanScanner(), bid_id=bid.id)
        .ingest(filename or f"{number}.dxf", payload)
        .stored[0]
    )
    [proposal] = read_title_blocks(
        session, store, stored, process_document(session, store, stored).sheets
    )
    session.commit()
    return proposal


def specification(session: Session, bid: Bid, store: MemoryObjectStore, revision: str) -> None:
    made = docx.Document()
    made.add_paragraph("PARTICULAR SPECIFICATION FOR FIRE PROTECTION SERVICES")
    made.add_paragraph(f"Revision {revision}")
    for clause in range(1, 7):
        made.add_paragraph(f"1.{clause} Requirement {clause}")
    buffer = io.BytesIO()
    made.save(buffer)
    payload = buffer.getvalue()
    stored = (
        Ingestor(session, store, AlwaysCleanScanner(), bid_id=bid.id)
        .ingest(f"Particular Specification Rev {revision}.docx", payload)
        .stored[0]
    )
    classify_in_sandbox(session, stored, payload)
    session.commit()


def unreadable(session: Session, bid: Bid, store: MemoryObjectStore) -> SheetRevision:
    plan, space = synthetic._new_drawing()
    space.add_text("FIRE SPRINKLER LAYOUT", height=250).set_placement((1_000, 1_000))
    return drawing(session, bid, store, "", filename="untitled.dxf", document=plan)


class TestTheDrawingRegister:
    def test_every_revision_is_listed_with_the_current_first(
        self,
        session: Session,
        bid: Bid,
        store: MemoryObjectStore,
        sign_in: SignIn,
        estimator: Principal,
    ) -> None:
        drawing(session, bid, store, "R03")
        drawing(session, bid, store, "R04")

        rows = sign_in(estimator).get(f"/bids/{bid.id}/registers/drawings").json()

        assert [(row["revision"], row["state"]) for row in rows] == [
            ("R04", "current"),
            ("R03", "superseded"),
        ]
        assert rows[0]["title"] == "FIRE SPRINKLER LAYOUT"
        assert rows[0]["level"] == "L05"

    def test_it_filters_by_state_level_and_discipline(
        self,
        session: Session,
        bid: Bid,
        store: MemoryObjectStore,
        sign_in: SignIn,
        estimator: Principal,
    ) -> None:
        drawing(session, bid, store, "R03")
        drawing(session, bid, store, "R04")
        drawing(session, bid, store, "R01", number="FP-L06-201")
        client = sign_in(estimator)

        def numbers(**query: str) -> list[tuple[str, str]]:
            rows = client.get(f"/bids/{bid.id}/registers/drawings", params=query).json()
            return [(row["sheet_number"], row["revision"]) for row in rows]

        assert numbers(state="current") == [(NUMBER, "R04"), ("FP-L06-201", "R01")]
        assert numbers(level="L06") == [("FP-L06-201", "R01")]
        assert len(numbers(discipline="fire protection")) == 3
        assert numbers(discipline="mechanical") == []


class TestExport:
    """The acceptance test: the register export opens in Excel with the right columns."""

    def test_the_drawing_register_opens_as_a_workbook_with_its_columns(
        self,
        session: Session,
        bid: Bid,
        store: MemoryObjectStore,
        sign_in: SignIn,
        estimator: Principal,
    ) -> None:
        drawing(session, bid, store, "R03")
        drawing(session, bid, store, "R04")

        response = sign_in(estimator).get(f"/bids/{bid.id}/registers/drawings.xlsx")

        assert response.status_code == 200
        assert response.headers["content-type"].startswith(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        assert "Drawing register" in response.headers["content-disposition"]
        book = openpyxl.load_workbook(io.BytesIO(response.content))
        sheet = book["Drawing register"]
        header = [cell.value for cell in sheet[1]]
        assert header == [heading for heading, _, _ in DRAWING_COLUMNS]
        first = {heading: cell.value for heading, cell in zip(header, sheet[2], strict=True)}
        assert first["Drawing number"] == NUMBER
        assert first["Revision"] == "R04"
        assert first["Status"] == "current"
        assert isinstance(first["Revision date"], datetime), "a date, so Excel sorts it"
        assert first["Manual takeoff recommended"] == "No"
        assert sheet.freeze_panes == "A2"
        assert sheet.max_row == 3
        assert book["About"]["B2"].value == bid.human_id

    def test_the_specification_register_exports_too(
        self,
        session: Session,
        bid: Bid,
        store: MemoryObjectStore,
        sign_in: SignIn,
        estimator: Principal,
    ) -> None:
        specification(session, bid, store, "2")

        response = sign_in(estimator).get(f"/bids/{bid.id}/registers/documents.xlsx")

        book = openpyxl.load_workbook(io.BytesIO(response.content))
        sheet = book["Specification register"]
        assert [cell.value for cell in sheet[1]] == [heading for heading, _, _ in DOCUMENT_COLUMNS]
        assert sheet["C2"].value == "specification"

    def test_an_empty_register_still_has_its_header(
        self, bid: Bid, sign_in: SignIn, estimator: Principal
    ) -> None:
        response = sign_in(estimator).get(f"/bids/{bid.id}/registers/drawings.xlsx")

        sheet = openpyxl.load_workbook(io.BytesIO(response.content))["Drawing register"]
        assert sheet["A1"].value == "Drawing number"
        assert sheet.max_row == 1


class TestTheSpecificationRegister:
    def test_one_current_revision_per_document(
        self,
        session: Session,
        bid: Bid,
        store: MemoryObjectStore,
        sign_in: SignIn,
        estimator: Principal,
    ) -> None:
        specification(session, bid, store, "1")
        specification(session, bid, store, "2")

        rows = sign_in(estimator).get(f"/bids/{bid.id}/registers/documents").json()

        assert [(row["revision"], row["state"]) for row in rows] == [
            ("2", "current"),
            ("1", "superseded"),
        ]
        assert rows[0]["doc_type"] == "specification"
        assert rows[0]["doc_key"] == "PARTICULAR SPECIFICATION FOR FIRE PROTECTION SERVICES"


class TestDecisions:
    @pytest.mark.req("FR-DOC-02")
    def test_a_confirmed_reading_teaches_the_consultants_layout(
        self,
        session: Session,
        bid: Bid,
        store: MemoryObjectStore,
        sign_in: SignIn,
        estimator: Principal,
        monkeypatch: pytest.MonkeyPatch,
    ) -> None:
        """Confirm one sheet, and the next from that consultant needs nobody."""
        from firebid.drawings import title_block

        # Make the reader unsure of this consultant's first sheet, as with a new layout.
        with monkeypatch.context() as unsure:
            unsure.setattr(title_block, "LABELLED", 0.6)
            first = drawing(session, bid, store, "R04")
        assert first.state == "received"

        response = sign_in(estimator).post(
            f"/bids/{bid.id}/sheet-revisions/{first.id}/confirm",
            json={"sheet_number": NUMBER, "revision": "R04", "consultant": "Synthetic Consultants"},
        )

        assert response.status_code == 200, response.text
        assert response.json()["state"] == "current"
        assert response.json()["read_by"] == "person"
        assert session.execute(select(TitleBlockLayout)).scalar_one().consultant == (
            "Synthetic Consultants"
        )
        second = drawing(session, bid, store, "R01", number="FP-L06-201")
        assert second.extraction_method == "layout"
        assert second.state == "current"

    @pytest.mark.req("FR-DOC-04")
    def test_a_person_resolves_a_conflict(
        self,
        session: Session,
        bid: Bid,
        store: MemoryObjectStore,
        sign_in: SignIn,
        estimator: Principal,
    ) -> None:
        disputed = drawing(session, bid, store, "R04", filename=f"{NUMBER}-R03.dxf")

        response = sign_in(estimator).post(
            f"/bids/{bid.id}/sheet-revisions/{disputed.id}/resolve",
            json={"outcome": "current", "reason": "misnamed file", "revision": "R04"},
        )

        assert response.status_code == 200, response.text
        assert response.json()["state"] == "current"
        assert response.json()["conflict_reason"] is None

    def test_a_role_without_the_right_cannot_decide(
        self,
        session: Session,
        bid: Bid,
        store: MemoryObjectStore,
        sign_in: SignIn,
        design_manager: Principal,
    ) -> None:
        disputed = drawing(session, bid, store, "R04", filename=f"{NUMBER}-R03.dxf")

        response = sign_in(design_manager).post(
            f"/bids/{bid.id}/sheet-revisions/{disputed.id}/resolve",
            json={"outcome": "current", "reason": "not mine to decide"},
        )

        assert response.status_code == 403

    def test_a_page_that_is_not_a_drawing_can_be_withdrawn(
        self,
        session: Session,
        bid: Bid,
        store: MemoryObjectStore,
        sign_in: SignIn,
        estimator: Principal,
    ) -> None:
        cover = unreadable(session, bid, store)

        response = sign_in(estimator).post(
            f"/bids/{bid.id}/sheet-revisions/{cover.id}/withdraw",
            json={"reason": "cover sheet, not a drawing"},
        )

        assert response.status_code == 200, response.text
        assert response.json()["state"] == "withdrawn"


class TestRegisterConfirmed:
    def test_confirmation_is_refused_while_something_is_undecided(
        self,
        session: Session,
        bid: Bid,
        store: MemoryObjectStore,
        sign_in: SignIn,
        estimator: Principal,
    ) -> None:
        drawing(session, bid, store, "R04", filename=f"{NUMBER}-R03.dxf")
        client = sign_in(estimator)

        status_now = client.get(f"/bids/{bid.id}/registers/status").json()
        response = client.post(f"/bids/{bid.id}/registers/confirm", json={})

        assert status_now["ready"] is False
        assert status_now["conflicts"] == 1
        assert response.status_code == 409
        assert "1 in Conflict" in response.json()["detail"]

    def test_the_estimator_confirms_and_the_bid_moves_to_takeoff(
        self,
        session: Session,
        bid: Bid,
        store: MemoryObjectStore,
        sign_in: SignIn,
        estimator: Principal,
    ) -> None:
        bid.stage = "S1"
        session.commit()
        drawing(session, bid, store, "R03")
        drawing(session, bid, store, "R04")
        client = sign_in(estimator)

        response = client.post(
            f"/bids/{bid.id}/registers/confirm", json={"comment": "Set complete per transmittal"}
        )

        assert response.status_code == 200, response.text
        confirmation = response.json()
        assert confirmation["drawings"] == 1, "one Current revision of one drawing"
        assert confirmation["confirmed_role"] == "estimator"
        assert len(confirmation["snapshot_hash"]) == 64
        session.expire_all()
        assert session.get(Bid, bid.id).stage == "S2"  # type: ignore[union-attr]
        assert client.get(f"/bids/{bid.id}/registers/status").json()["confirmation"] is not None
        audited = session.execute(
            select(AuditEvent).where(AuditEvent.action == "register: confirmed")
        ).scalar_one()
        assert audited.actor_id == estimator.user_id

    def test_only_an_estimator_confirms_the_register(
        self, bid: Bid, sign_in: SignIn, design_manager: Principal
    ) -> None:
        response = sign_in(design_manager).post(f"/bids/{bid.id}/registers/confirm", json={})

        assert response.status_code == 403
