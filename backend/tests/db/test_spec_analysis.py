# ruff: noqa: F811  (the fixtures below are imported by name and requested by name)
"""Full specification analysis through the database and the API (FR-SPEC-02, 03, 04).

The extended synthetic specification is read as an uploaded specification is, and the
basement car park plan, whose notes contradict it, is taken off. The analysis then finds
every obligation, every seeded issue with both citations, and the scope matrix, which a
person edits, confirms and exports.
"""

from __future__ import annotations

import io
import json
from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook
from sqlalchemy import select, text
from sqlalchemy.orm import Session

from firebid.ai_gateway.config import load_config
from firebid.ai_gateway.metering import PostgresMeter
from firebid.ai_gateway.providers.fake import FakeAdapter
from firebid.ai_gateway.router import Router
from firebid.auth.provisioning import Principal
from firebid.db.models.audit import AuditEvent
from firebid.db.models.core import Bid, Organisation
from firebid.db.models.documents import DocumentRevision
from firebid.db.models.specs import SpecClause
from firebid.domain.state_machines import Role
from firebid.evals import synthetic_qto
from firebid.evals.synthetic_network import NETWORK
from firebid.evals.synthetic_spec import (
    CAR_PARK_NOTES,
    EXPECTED_INTERFACES,
    EXPECTED_OBLIGATIONS,
    SEEDED_ISSUES,
    extended,
    specification_docx,
)
from firebid.ingest.scanning import AlwaysCleanScanner
from firebid.services import qto, specs
from firebid.services import spec_analysis as analysis
from firebid.services.classification import classify_in_sandbox
from firebid.services.detection import detect_bid
from firebid.services.ingestion import Ingestor
from firebid.specs import obligations
from firebid.storage.object_store import MemoryObjectStore
from tests.db.test_detection_pipeline import confirm_legend
from tests.db.test_sheet_views import SignIn, app, member, sign_in  # noqa: F401
from tests.db.test_symbol_mapping import from_consultant, no_tiles, read, store  # noqa: F401

SHEET = "FP-B1-201"
MODEL_CONFIG = """
version: 1
providers:
  primary:
    kind: fake
    approved_data_classes: [internal, confidential]
models:
  first:
    provider: primary
    model_id: reader-1
    capabilities: [structured_output]
    prices:
      - {effective_from: 2026-01-01, input_per_mtok: 5.0, output_per_mtok: 25.0}
routes:
  spec_obligation_extract:
    requires: [structured_output]
    data_class: confidential
    models: [first]
    reasoning: low
    prompt: spec_obligation_extract
"""


def specification(session: Session, bid: Bid, store: MemoryObjectStore) -> DocumentRevision:
    payload = specification_docx(clauses=extended())
    document = (
        Ingestor(session, store, AlwaysCleanScanner(), bid_id=bid.id)
        .ingest("Particular Specification Fire Protection Rev B.docx", payload)
        .stored[0]
    )
    classify_in_sandbox(session, document, payload)
    revision = specs.read_specification(session, store, document)
    assert revision is not None
    revision.state = "current"
    session.commit()
    return revision


@pytest.fixture
def reviewer(session: Session, organisation: Organisation, bid: Bid) -> Principal:
    return member(session, organisation, bid, "esther", Role.ESTIMATOR)


@pytest.fixture
def tender(session: Session, bid: Bid, store: MemoryObjectStore) -> Bid:
    """The specification read, and the car park plan taken off."""
    specification(session, bid, store)
    from_consultant(session, bid, NETWORK.name)
    plan, _ = synthetic_qto.car_park_plan(SHEET, CAR_PARK_NOTES)
    read(session, bid, store, SHEET, plan)
    confirm_legend(session, bid, store)
    detect_bid(session, store, bid.id)
    qto.recompute(session, bid.id)
    session.execute(text("DELETE FROM procrastinate_jobs"))
    session.commit()
    return bid


@pytest.mark.req("FR-SPEC-02")
class TestObligations:
    def test_reading_a_specification_stores_every_obligation_with_its_citation(
        self, session: Session, bid: Bid, store: MemoryObjectStore
    ) -> None:
        revision = specification(session, bid, store)

        rows = analysis.list_obligations(session, bid.id)

        assert sorted((r.category, r.clause_number, r.quantities) for r in rows) == sorted(
            EXPECTED_OBLIGATIONS, key=lambda e: (e[0], e[1])
        )
        assert {r.category for r in rows} == set(obligations.CATEGORY_KEYS)
        clauses = {
            c.id: c
            for c in session.execute(
                select(SpecClause).where(SpecClause.document_revision_id == revision.id)
            ).scalars()
        }
        for row in rows:
            assert (row.state, row.method, row.citation_ok) == ("proposed", "rule", True)
            assert row.clause_id is not None and row.quote in clauses[row.clause_id].text
            assert row.document_revision_id == revision.id

    def test_reading_again_changes_nothing(
        self, session: Session, bid: Bid, store: MemoryObjectStore
    ) -> None:
        revision = specification(session, bid, store)
        before = [row.id for row in analysis.list_obligations(session, bid.id)]

        assert analysis.read_obligations(session, revision) == []
        analysis.analyse(session, bid.id)

        assert [row.id for row in analysis.list_obligations(session, bid.id)] == before

    def test_a_person_confirms_or_rejects_and_it_is_audited(
        self,
        session: Session,
        bid: Bid,
        store: MemoryObjectStore,
        reviewer: Principal,
        sign_in: SignIn,
    ) -> None:
        specification(session, bid, store)
        client = sign_in(reviewer)
        listed = client.get(f"/bids/{bid.id}/spec/obligations").json()
        testing = next(o for o in listed if o["category"] == "testing")

        confirmed = client.post(
            f"/bids/{bid.id}/spec/obligations/{testing['id']}/decide",
            json={"decision": "confirm"},
        )

        assert confirmed.status_code == 200, confirmed.text
        assert (confirmed.json()["state"], confirmed.json()["decided_by"]) == ("verified", "Esther")
        assert testing["quantities"] == {
            "pressure": 14,
            "pressure_unit": "bar",
            "duration_hours": 2,
        }
        event = session.execute(
            select(AuditEvent).where(AuditEvent.action == "specification obligation: confirm")
        ).scalar_one()
        assert event.after["clause"] == "6.1"  # type: ignore[index]

    def test_a_model_answer_is_checked_against_the_clause_it_cites(
        self, session: Session, bid: Bid, store: MemoryObjectStore, tmp_path: Path
    ) -> None:
        revision = specification(session, bid, store)
        clause = session.execute(
            select(SpecClause).where(
                SpecClause.document_revision_id == revision.id, SpecClause.number == "7.6"
            )
        ).scalar_one()
        clause.text += " The header shall be witnessed by the Engineer."
        session.flush()
        reply = {
            "obligations": [
                {
                    "category": "testing",
                    "summary": "The flow test at the header is witnessed.",
                    "clause": "7.6",
                    "quote": "The header shall be witnessed by the Engineer.",
                    "quantities": {},
                    "confidence": 0.9,
                },
                {
                    "category": "spares",
                    "summary": "Ten spare gauges.",
                    "clause": "7.6",
                    "quote": "ten spare gauges shall be supplied",
                    "quantities": {"count": 10},
                    "confidence": 0.9,
                },
            ]
        }

        stored = analysis.read_obligations_with_model(
            session, revision.id, router(session, tmp_path, reply)
        )

        assert [(r.category, r.method, r.citation_ok) for r in stored] == [
            ("testing", "model", True),
            ("spares", "model", False),
        ]
        assert stored[0].confidence == 0.9 and stored[1].confidence == 0.2
        assert stored[1].citation_reason == "clause 7.6 does not contain the words quoted"
        assert stored[0].provenance["model"] == "reader-1"


def router(session: Session, tmp_path: Path, reply: dict[str, Any]) -> Router:
    path = tmp_path / "llm.yaml"
    path.write_text(MODEL_CONFIG, encoding="utf-8")
    return Router(
        config=load_config(path),
        adapters={"primary": FakeAdapter("primary").reply(json.dumps(reply))},
        backoff_base_seconds=0,
        sleep=lambda _s: None,
        meter=PostgresMeter(session),
    )


@pytest.mark.req("FR-SPEC-03")
class TestIssues:
    def test_every_seeded_issue_is_flagged_with_both_citations(
        self, session: Session, tender: Bid
    ) -> None:
        revision = analysis.current_specifications(session, tender.id)[0]

        outcome = analysis.analyse(session, tender.id)
        issues = analysis.list_issues(session, tender.id)

        assert sorted(
            (i.rule, i.spec_ref.get("clause"), i.detail["what"]) for i in issues
        ) == sorted(SEEDED_ISSUES, key=lambda s: (s[0], s[1] or "", s[2]))
        assert outcome.issues_new == len(SEEDED_ISSUES) == outcome.issues_open
        for issue in issues:
            # The specification side: document, revision, and the clause where one speaks.
            assert issue.spec_ref["document_revision_id"] == str(revision.id)
            assert issue.spec_ref["revision"] == revision.revision_label
            # The drawing side: the sheet and its revision, or the sheets looked through.
            cited: list[Any] = (
                [issue.drawing_ref]
                if issue.drawing_ref.get("sheet_number")
                else list(issue.drawing_ref["sheets"])  # type: ignore[call-overload]
            )
            assert cited
            assert all(c["sheet_number"] == SHEET and c["revision"] == "R01" for c in cited)

    def test_the_p1_06_contradiction_cites_the_clause_and_the_note(
        self, session: Session, tender: Bid
    ) -> None:
        analysis.analyse(session, tender.id)

        [conflict] = [
            i
            for i in analysis.list_issues(session, tender.id)
            if i.rule == "conflict:pipe_material"
        ]

        assert (conflict.category, conflict.severity, conflict.state) == (
            "conflict",
            "high",
            "open",
        )
        assert conflict.spec_ref["clause"] == "2.1.3"
        assert "hot-dip galvanised" in str(conflict.spec_ref["quote"])
        assert conflict.drawing_ref["note"] == CAR_PARK_NOTES[0]
        assert conflict.drawing_ref["sheet_id"] and conflict.drawing_ref["x"] is not None

    def test_open_issues_are_the_clarification_candidates(
        self, session: Session, tender: Bid, reviewer: Principal, sign_in: SignIn
    ) -> None:
        client = sign_in(reviewer)
        ran = client.post(f"/bids/{tender.id}/spec/analysis/run")
        assert ran.status_code == 200 and ran.json()["issues_open"] == len(SEEDED_ISSUES)
        listed = client.get(f"/bids/{tender.id}/spec/issues").json()
        ambiguous = next(i for i in listed if i["detail"]["what"] == "or equal")

        refused = client.post(
            f"/bids/{tender.id}/spec/issues/{ambiguous['id']}/decide", json={"decision": "dismiss"}
        )
        dismissed = client.post(
            f"/bids/{tender.id}/spec/issues/{ambiguous['id']}/decide",
            json={"decision": "dismiss", "note": "the schedule names the make"},
        )
        candidates = client.get(f"/bids/{tender.id}/spec/clarification-candidates").json()

        assert refused.status_code == 422, "a dismissal says why"
        assert dismissed.json()["decided_by"] == "Esther"
        assert len(candidates) == len(SEEDED_ISSUES) - 1
        assert ambiguous["id"] not in {c["id"] for c in candidates}
        assert [c["severity"] for c in candidates][:3] == ["high", "high", "high"]

    def test_running_again_keeps_a_dismissal_and_resolves_what_is_no_longer_found(
        self, session: Session, tender: Bid, reviewer: Principal
    ) -> None:
        analysis.analyse(session, tender.id)
        issues = {
            i.rule + ":" + str(i.detail["what"]): i
            for i in analysis.list_issues(session, tender.id)
        }
        analysis.decide_issue(
            session, issues["ambiguous_clause:or equal"], "dismiss", reviewer.actor(), "scheduled"
        )
        # A person rejects the clause the car park conflict rests on.
        galvanised = next(
            a for a in specs.current_attributes(session, tender.id) if a.clause_number == "2.1.3"
        )
        specs.decide(session, galvanised.lineage_id, reviewer.actor(), verdict="reject")

        again = analysis.analyse(session, tender.id)

        after = {
            i.rule + ":" + str(i.detail["what"]): i.state
            for i in analysis.list_issues(session, tender.id)
        }
        assert after["ambiguous_clause:or equal"] == "dismissed"
        assert after["conflict:pipe_material:pipe_material"] == "resolved"
        assert (again.issues_new, again.issues_resolved) == (0, 1)
        assert again.issues_open == len(SEEDED_ISSUES) - 2


@pytest.mark.req("FR-SPEC-04")
class TestScopeMatrix:
    def test_a_row_per_obligation_and_interface_with_status_and_clause(
        self, session: Session, tender: Bid
    ) -> None:
        analysis.analyse(session, tender.id)

        rows = analysis.matrix(session, tender.id)

        assert sorted({r.system for r in rows}) == ["hose_reel", "hydrant", "sprinkler"]
        sprinkler = [r for r in rows if r.system == "sprinkler"]
        interfaces = [
            (r.key, r.status, r.clause_number) for r in sprinkler if r.kind == "interface"
        ]
        assert interfaces == list(EXPECTED_INTERFACES)
        assert {r.key for r in sprinkler if r.kind == "obligation"} == set(
            obligations.CATEGORY_KEYS
        )
        for row in rows:
            assert row.status in ("included", "excluded", "by_others", "unclear")
            assert row.document_revision_id is not None, "every row links to its specification"
            if row.clause_number is not None:
                assert row.clause_id is not None and row.quote
            else:
                assert row.reason == "the specification does not mention it"

    def test_the_estimator_edits_confirms_and_exports(
        self, session: Session, tender: Bid, reviewer: Principal, sign_in: SignIn
    ) -> None:
        client = sign_in(reviewer)
        client.post(f"/bids/{tender.id}/spec/analysis/run")
        rows = client.get(f"/bids/{tender.id}/spec/scope-matrix").json()
        access = next(
            r for r in rows if r["system"] == "sprinkler" and r["key"] == "ceiling_access"
        )

        edited = client.post(
            f"/bids/{tender.id}/spec/scope-matrix/rows/{access['id']}",
            json={"status": "by_others", "note": "architect's package, per the tender briefing"},
        )
        confirmed = client.post(f"/bids/{tender.id}/spec/scope-matrix/confirm")
        export = client.get(f"/bids/{tender.id}/spec/scope-matrix/export.xlsx")

        assert (edited.json()["status"], edited.json()["source"]) == ("by_others", "person")
        assert edited.json()["proposed_status"] == "unclear"
        assert {r["confirmed_by"] for r in confirmed.json()} == {"Esther"}
        assert export.status_code == 200
        assert export.headers["content-type"].startswith(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        sheet = load_workbook(io.BytesIO(export.content))["Scope matrix"]
        table = list(sheet.iter_rows(values_only=True))
        assert table[0][:5] == ("System", "Kind", "Item", "Status", "Clause")
        assert len(table) == 1 + len(rows)
        by_item = {(r[0], r[2]): r for r in table[1:]}
        power = by_item[("sprinkler", "Power supply to fire pumps and control panels")]
        assert (power[3], power[4], power[6]) == ("By others", "7.1", "B")
        assert "electrical contractor" in str(power[7])
        panels = by_item[("sprinkler", "Ceiling access panels")]
        assert (panels[3], panels[9], panels[11]) == ("By others", "Esther", "Esther")

    def test_running_again_keeps_a_person_s_status(
        self, session: Session, tender: Bid, reviewer: Principal
    ) -> None:
        analysis.analyse(session, tender.id)
        row = next(
            r
            for r in analysis.matrix(session, tender.id)
            if r.system == "hydrant" and r.key == "excavation"
        )
        analysis.edit_row(session, row, "excluded", reviewer.actor(), "by the civil contractor")

        analysis.analyse(session, tender.id)

        assert (row.status, row.source, row.proposed_status) == ("excluded", "person", "unclear")
        assert len(analysis.matrix(session, tender.id)) == 3 * (13 + 8) + 1

    def test_confirming_needs_a_matrix(
        self, session: Session, bid: Bid, reviewer: Principal
    ) -> None:
        with pytest.raises(analysis.AnalysisError, match="run the analysis first"):
            analysis.confirm_matrix(session, bid.id, reviewer.actor())
