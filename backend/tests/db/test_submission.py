# ruff: noqa: F811  (the fixtures below are imported by name and requested by name)
"""The review pack, gates G2 to G4, the frozen submission, outcomes and library governance
through the database and the API (P2-08).

The risk fixture's bid (a specification analysed, a basement plan taken off and billed, a
productivity library) is priced from the synthetic rate list and taken through its gates.
"""

from __future__ import annotations

import ast
import io
import json
from datetime import UTC, date, datetime
from decimal import Decimal
from pathlib import Path

import pdfplumber
import pytest
from openpyxl import load_workbook
from sqlalchemy import select, text
from sqlalchemy.exc import DBAPIError, IntegrityError
from sqlalchemy.orm import Session

from firebid.api import submission as submission_api
from firebid.db.models.audit import AuditEvent
from firebid.db.models.commercial import Rate
from firebid.db.models.core import Bid, Organisation
from firebid.db.models.workflow import Approval
from firebid.domain.state_machines import BidState
from firebid.evals import synthetic_rates
from firebid.services import boq, costing, labour, pricing, review_pack
from firebid.services import library_governance as governance
from firebid.services import risk as risk_service
from firebid.services import submission as service
from firebid.storage.object_store import MemoryObjectStore
from tests.db.test_risk import Team, risk_bid  # noqa: F401
from tests.db.test_sheet_views import SignIn, app, sign_in  # noqa: F401
from tests.db.test_symbol_mapping import no_tiles, store  # noqa: F401

pytestmark = pytest.mark.usefixtures("no_tiles")

TODAY = date(2026, 10, 4)
SECTIONS = [
    "estimate_by_component",
    "estimate_by_system",
    "cost_drivers",
    "risk_allowances",
    "variances",
    "open_items",
    "unpriced_lines",
    "g1_coverage",
]


@pytest.fixture
def estimated(
    session: Session,
    risk_bid: tuple[Bid, Team],
    organisation: Organisation,
) -> tuple[Bid, Team]:
    """The bid priced from the rate library, with labour, a margin, risks found and G1 on
    record: an estimate to review."""
    bid, team = risk_bid
    report = pricing.import_rates(
        session, organisation.id, synthetic_rates.rate_list(), team.senior, "rates.xlsx"
    )
    assert report.imported, report.problems
    pricing.price_boq(session, bid, team.estimator)
    costing.enter(
        session,
        bid,
        team.senior,
        component="margin",
        basis="percentage",
        percent=Decimal(8),
        base="cost",
    )
    assert team.senior.id is not None
    session.add(
        Approval(
            bid_id=bid.id,
            gate="G1",
            decision="approved",
            approver_id=team.senior.id,
            approver_role="senior_estimator",
            decided_at=datetime.now(UTC),
        )
    )
    risk_service.build_checklist(session, bid, team.estimator)
    risk_service.find(session, bid, team.estimator)
    session.commit()
    return bid, team


def ready_for_g3(session: Session, bid: Bid, team: Team) -> None:
    """Everything G3 waits for: G2, the checklist resolved, every risk treated, and the bid
    under review."""
    boq.approve_g2(session, bid, team.senior, "senior_estimator")
    for check in risk_service.checks(session, bid.id):
        if check.status == "open":
            risk_service.resolve_check(
                session, bid, check, team.estimator, status="by_others", note="main contract"
            )
    for item in risk_service.risks(session, bid.id):
        risk_service.treat(session, bid, item, team.manager, treatment=item.proposed_treatment)
    bid.state = str(BidState.UNDER_REVIEW)
    session.commit()


def ready_for_g4(session: Session, bid: Bid, team: Team) -> None:
    ready_for_g3(session, bid, team)
    service.approve_g3(session, bid, team.director, "approved as reviewed")
    risk_service.propose_qualifications(session, bid, team.manager)
    for entry in risk_service.qualifications(session, bid.id):
        entry.state, entry.decided_by, entry.decided_by_id = "accepted", "Bella", team.manager.id
    session.commit()


@pytest.mark.req("FR-PKG-01")
class TestReviewPack:
    def test_the_pack_has_every_section_and_its_figures_are_the_estimate_s(
        self, session: Session, estimated: tuple[Bid, Team]
    ) -> None:
        bid, _ = estimated

        pack = review_pack.build(session, bid, TODAY)

        assert [section.key for section in pack.sections] == SECTIONS
        assert all(section.link.startswith(f"/bids/{bid.id}/") for section in pack.sections)
        built = costing.build_up(session, bid, TODAY)
        current = boq.current_boq(session, bid.id)
        assert current is not None
        totals = pricing.boq_totals(session, current)
        estimate = labour.estimate(session, bid, TODAY)
        assert pack.figures["total_excluding_gst"] == str(built.total.amount)
        assert pack.figures["gst"] == str(built.gst.gst.amount)
        assert pack.figures["total_including_gst"] == str(built.gst.inclusive.amount)
        assert pack.figures["direct"] == str(built.direct.amount)
        assert pack.figures["priced_bill"] == str(totals.grand.amount)
        assert pack.figures["labour_cost"] == str(estimate.cost.amount)
        assert Decimal(pack.figures["total_excluding_gst"]) > Decimal(pack.figures["direct"]) > 0

        # By component: one row a component, adding up to the total; the totals rows after.
        components = pack.section("estimate_by_component")
        amounts = {row[0]: row[3] for row in components.rows}
        assert len(components.rows) == len(built.lines) + 4
        assert amounts["Total (excluding GST)"] == f"{built.total.amount:,.2f}"
        assert amounts["Labour"] == f"{estimate.cost.amount:,.2f}"
        # By system: the bill and the labour of each section.
        [system] = pack.section("estimate_by_system").rows
        assert system[1] == f"{totals.grand.amount:,.2f}"
        assert system[3] == f"{estimate.cost.amount:,.2f}"
        # Drivers are the largest lines first, and the margin is stated.
        drivers = [Decimal(row[4].replace(",", "")) for row in pack.section("cost_drivers").rows]
        assert drivers and drivers == sorted(drivers, reverse=True)
        margin = next(line for line in built.lines if line.component == "margin")
        assert margin.amount is not None
        assert pack.section("cost_drivers").note.startswith(
            f"Margin: SGD {margin.amount.amount:,.2f} ("
        )
        # Risks with their treatments; what is unpriced; issues open; coverage.
        risks = pack.section("risk_allowances")
        assert {row[2] for row in risks.rows} == {"none"} and "have no treatment" in risks.note
        assert pack.figures["unpriced_lines"] == str(len(pack.section("unpriced_lines").rows))
        assert int(pack.figures["unpriced_lines"]) == totals.unpriced > 0
        assert int(pack.figures["open_issues"]) > 0
        assert pack.section("open_items").rows[0][0] == "specification issue"
        assert pack.section("g1_coverage").rows[0][0] == "Items verified"

    def test_the_pack_is_a_workbook_and_a_pdf_with_every_section(
        self, session: Session, estimated: tuple[Bid, Team]
    ) -> None:
        bid, _ = estimated
        pack = review_pack.build(session, bid, TODAY)

        book = load_workbook(io.BytesIO(review_pack.as_workbook(pack)))
        pdf = review_pack.as_pdf(pack)

        assert book.sheetnames == ["Summary", *(section.title[:31] for section in pack.sections)]
        summary = {row[0]: row[1] for row in book["Summary"].iter_rows(values_only=True) if row[0]}
        assert summary["Total excluding GST (SGD)"] == pack.figures["total_excluding_gst"]
        assert summary["Risk allowances and treatments"] == f"/bids/{bid.id}/risk"
        component = book["Estimate by component"]
        assert component["A2"].value == f"In the platform: /bids/{bid.id}/boq"
        assert [cell.value for cell in component[4]][:4] == [
            "Component",
            "Basis",
            "Source",
            "Amount (SGD)",
        ]
        assert pdf.startswith(b"%PDF")
        with pdfplumber.open(io.BytesIO(pdf)) as document:
            pages = [page.extract_text() or "" for page in document.pages]
        assert len(pages) >= len(pack.sections) + 1
        for section in pack.sections:
            assert any(section.title in page and section.link in page for page in pages)
        assert pack.figures["total_excluding_gst"] in pages[0]


@pytest.mark.req("FR-PKG-02")
class TestGates:
    def test_g3_is_blocked_with_reasons_until_it_is_ready_and_records_what_it_approved(
        self, session: Session, estimated: tuple[Bid, Team]
    ) -> None:
        bid, team = estimated

        waiting = {gate.gate: gate for gate in service.gate_status(session, bid)}
        assert waiting["G1"].approved and not waiting["G2"].approved
        assert waiting["G3"].blockers[0] == "G2 is not approved"
        assert "risk(s) with no treatment" in waiting["G3"].blockers[1]
        assert waiting["G4"].blockers == ["G3 is not approved"]
        with pytest.raises(service.GateError, match="G3 is blocked: G2 is not approved; "):
            service.approve_g3(session, bid, team.director)

        ready_for_g3(session, bid, team)
        with pytest.raises(service.GateError, match="G3 is approved by the commercial director"):
            service.approve_g3(session, bid, team.senior)
        with pytest.raises(service.GateError, match="G3 is approved by the commercial director"):
            service.approve_g3(session, bid, team.manager)
        expected = service.estimate_hash(session, bid)
        approval = service.approve_g3(session, bid, team.director, "margin agreed at 8%")
        session.commit()

        assert bid.state == "approved_for_submission"
        assert (approval.gate, approval.approver_id, approval.approver_role) == (
            "G3",
            team.director.id,
            "commercial_director",
        )
        assert approval.snapshot_hash == expected and len(expected) == 64
        assert approval.comment == "margin agreed at 8%" and approval.decided_at is not None
        now = {gate.gate: gate for gate in service.gate_status(session, bid)}
        assert now["G3"].approved and now["G3"].snapshot_hash == expected
        assert now["G2"].approved and now["G2"].snapshot_hash

    def test_each_gate_is_approvable_only_by_its_role_through_the_api(
        self,
        session: Session,
        estimated: tuple[Bid, Team],
        sign_in: SignIn,
        monkeypatch: pytest.MonkeyPatch,
    ) -> None:
        bid, team = estimated
        memory = MemoryObjectStore()
        monkeypatch.setattr(submission_api, "get_snapshot_store", lambda: memory)
        ready_for_g4(session, bid, team)
        base = f"/bids/{bid.id}"

        answers = {}
        for name, principal in (
            ("estimator", team.estimator_principal),
            ("senior", team.senior_principal),
            ("manager", team.manager_principal),
        ):
            client = sign_in(principal)
            answers[name] = (
                client.post(f"{base}/gates/g3/approve", json={}).status_code,
                client.post(f"{base}/gates/g4/approve", json={}).status_code,
            )
        estimator = sign_in(team.estimator_principal)
        early = estimator.get(f"{base}/submission/files/company-boq.xlsx")
        director = sign_in(team.director_principal)
        approved = director.post(f"{base}/gates/g4/approve", json={"comment": "submit"})
        again = director.post(f"{base}/gates/g4/approve", json={})
        file = director.get(f"{base}/submission/files/company-boq.xlsx")

        assert answers == {"estimator": (403, 403), "senior": (403, 403), "manager": (403, 403)}
        assert early.status_code == 409 and "once G4 is approved" in early.json()["detail"]
        assert approved.status_code == 201, approved.text
        gates = {gate["gate"]: gate for gate in approved.json()["gates"]}
        assert approved.json()["state"] == "submitted" and gates["G4"]["approved"]
        assert len(gates["G4"]["snapshot_hash"]) == 64
        assert again.status_code == 409
        assert file.status_code == 200 and file.content.startswith(b"PK")
        assert file.headers["content-disposition"].startswith("attachment;")
        downloads = (
            session.execute(
                select(AuditEvent).where(AuditEvent.action == "submission: file downloaded")
            )
            .scalars()
            .all()
        )
        assert [event.actor_label for event in downloads] == [team.director.label]

    def test_g4_waits_for_unresolved_clarifications_and_undecided_qualifications(
        self,
        session: Session,
        estimated: tuple[Bid, Team],
        store: MemoryObjectStore,
    ) -> None:
        bid, team = estimated
        ready_for_g3(session, bid, team)
        service.approve_g3(session, bid, team.director)
        risk_service.propose_qualifications(session, bid, team.manager)
        session.commit()

        with pytest.raises(service.GateError, match=r"qualification\(s\) neither accepted nor"):
            service.approve_g4(session, bid, team.director, store)

        assert service.snapshot_of(session, bid.id) is None and bid.state != "submitted"

    def test_nothing_can_send_a_bid_anywhere(
        self, session: Session, estimated: tuple[Bid, Team], store: MemoryObjectStore
    ) -> None:
        """After G4 the files are downloads for a person on the bid. No module of the gates,
        the pack or the snapshot opens a connection, sends mail or queues a job."""
        bid, team = estimated
        ready_for_g4(session, bid, team)
        service.approve_g4(session, bid, team.director, store)
        session.commit()

        assert session.execute(text("SELECT count(*) FROM procrastinate_jobs")).scalar_one() == 0
        root = Path(service.__file__).resolve().parents[1]
        forbidden = {
            "smtplib",
            "email",
            "socket",
            "http",
            "urllib",
            "httpx",
            "requests",
            "aiohttp",
            "ftplib",
            "firebid.jobs",
        }
        for file in (
            root / "services" / "submission.py",
            root / "services" / "review_pack.py",
            root / "api" / "submission.py",
        ):
            for node in ast.walk(ast.parse(file.read_text(encoding="utf-8"))):
                names = (
                    [alias.name for alias in node.names]
                    if isinstance(node, ast.Import)
                    else [node.module or ""]
                    if isinstance(node, ast.ImportFrom)
                    else []
                )
                for name in names:
                    assert not any(
                        name == bad or name.startswith(bad + ".") for bad in forbidden
                    ), f"{file.name} imports {name}"


@pytest.mark.req("FR-PKG-03")
class TestFreeze:
    def test_g4_freezes_the_submission_and_the_snapshot_verifies(
        self,
        session: Session,
        estimated: tuple[Bid, Team],
        store: MemoryObjectStore,
    ) -> None:
        bid, team = estimated
        ready_for_g4(session, bid, team)

        approval, snapshot = service.approve_g4(session, bid, team.director, store, "submit")
        session.commit()

        assert bid.state == "submitted"
        assert approval.snapshot_hash == snapshot.manifest_sha256
        manifest = json.loads(store.get(snapshot.manifest_key))
        assert manifest["bid"]["human_id"] == bid.human_id
        assert manifest["frozen_by"] == team.director.label
        # Every kind of record the submission rests on, each with its content hash.
        for kind in (
            "qto_item",
            "evidence",
            "boq",
            "boq_line",
            "rate",
            "clarification",
            "qualification",
            "risk",
            "scope_check",
            "approval",
        ):
            assert kind in manifest["records"], kind
        for kind in (
            "qto_item",
            "evidence",
            "boq_line",
            "rate",
            "qualification",
            "risk",
            "approval",
        ):
            rows = manifest["records"][kind]
            assert rows and all(len(row["sha256"]) == 64 and row["id"] for row in rows), kind
        assert {row["id"] for row in manifest["records"]["approval"]} >= {
            str(a.id)
            for a in session.execute(
                select(Approval).where(
                    Approval.bid_id == bid.id, Approval.gate.in_(["G1", "G2", "G3"])
                )
            ).scalars()
        }
        assert snapshot.record_counts["boq_line"] == len(manifest["records"]["boq_line"])
        assert [f["name"] for f in snapshot.files] == [
            "company-boq.xlsx",
            "review-pack.xlsx",
            "review-pack.pdf",
            "qualifications.xlsx",
            "clarifications.xlsx",
        ]
        assert all(store.exists(f["key"]) for f in snapshot.files)

        found = service.verify_snapshot(session, store, snapshot, bid)
        assert found.ok and found.problems == []
        assert found.files_checked == 5 and found.records_listed > 50
        assert found.changed_since == []

    def test_a_tampered_object_fails_verification(
        self,
        session: Session,
        estimated: tuple[Bid, Team],
        store: MemoryObjectStore,
    ) -> None:
        bid, team = estimated
        ready_for_g4(session, bid, team)
        _, snapshot = service.approve_g4(session, bid, team.director, store)
        session.commit()
        boq_file = next(f for f in snapshot.files if f["name"] == "company-boq.xlsx")
        original = store.objects[boq_file["key"]]

        store.objects[boq_file["key"]] = original + b"tampered"
        file_changed = service.verify_snapshot(session, store, snapshot)
        store.objects[boq_file["key"]] = original
        del store.objects[next(f["key"] for f in snapshot.files if f["name"] == "review-pack.pdf")]
        file_missing = service.verify_snapshot(session, store, snapshot)
        manifest = json.loads(store.objects[snapshot.manifest_key])
        manifest["records"]["boq_line"][0]["sha256"] = "0" * 64
        store.objects[snapshot.manifest_key] = json.dumps(manifest).encode()
        manifest_changed = service.verify_snapshot(session, store, snapshot)

        assert not file_changed.ok
        assert file_changed.problems == [
            "company-boq.xlsx is not the file that was frozen: its hash differs"
        ]
        assert file_missing.problems == ["review-pack.pdf is missing from the snapshot store"]
        assert manifest_changed.problems == [
            "the manifest is not the one that was frozen: its hash differs"
        ]

    def test_writes_to_the_snapshot_are_refused(
        self,
        session: Session,
        estimated: tuple[Bid, Team],
        store: MemoryObjectStore,
    ) -> None:
        bid, team = estimated
        ready_for_g4(session, bid, team)
        _, snapshot = service.approve_g4(session, bid, team.director, store)
        session.commit()

        for statement in (
            "UPDATE submission_snapshot SET manifest_sha256 = repeat('0', 64) WHERE id = :s",
            "DELETE FROM submission_snapshot WHERE id = :s",
        ):
            with pytest.raises(DBAPIError, match="append-only"):
                session.execute(text(statement), {"s": snapshot.id})
            session.rollback()
        # The store writes each object once: it cannot be put again.
        with pytest.raises(Exception, match=snapshot.manifest_key):
            store.put_once(snapshot.manifest_key, b"{}", content_type="application/json")
        with pytest.raises(service.GateError, match="already frozen"):
            service.freeze(session, bid, team.director, store)
        # A record changed after the freeze shows up against the snapshot; the snapshot stands.
        entry = risk_service.qualifications(session, bid.id)[0]
        entry.text = entry.text + " (changed afterwards)"
        session.commit()
        found = service.verify_snapshot(session, store, snapshot, bid)
        assert found.ok and found.changed_since == [f"qualification {entry.id}"]


@pytest.mark.req("FR-LRN-02")
class TestOutcome:
    def test_an_outcome_moves_the_lifecycle_and_is_reported(
        self,
        session: Session,
        estimated: tuple[Bid, Team],
        store: MemoryObjectStore,
        organisation: Organisation,
        sign_in: SignIn,
    ) -> None:
        bid, team = estimated
        ready_for_g4(session, bid, team)
        service.approve_g4(session, bid, team.director, store)
        session.commit()
        before = service.outcome_report(session, organisation.id)

        with pytest.raises(service.GateError, match="needs one of these roles"):
            service.record_outcome(session, bid, team.estimator, outcome="awarded")
        session.rollback()
        with pytest.raises(service.GateError, match="say why the tender was lost"):
            service.record_outcome(session, bid, team.manager, outcome="lost")
        row = service.record_outcome(
            session,
            bid,
            team.manager,
            outcome="awarded",
            awarded_price=Decimal("412500"),
            reasons="Lowest compliant offer; programme accepted.",
            competitor_feedback="Two others bid; the next was 4% higher.",
        )
        session.commit()

        assert bid.state == "awarded"
        assert row.awarded_price is not None and row.awarded_price.amount == Decimal("412500.00")
        assert (before.submitted, before.awaiting, before.win_rate_percent) == (1, 1, None)
        with pytest.raises(service.GateError, match="recorded as awarded: that stands"):
            service.record_outcome(session, bid, team.manager, outcome="lost", reasons="no")
        report = sign_in(team.manager_principal).get("/outcomes").json()
        assert (report["awarded"], report["lost"], report["awaiting"]) == (1, 0, 0)
        assert report["win_rate_percent"] == 100.0
        assert Decimal(report["awarded_value"]) == Decimal("412500.00")
        [listed] = report["rows"]
        assert (listed["human_id"], listed["outcome"], listed["state"]) == (
            bid.human_id,
            "awarded",
            "awarded",
        )
        event = session.execute(
            select(AuditEvent).where(AuditEvent.action == "outcome: recorded")
        ).scalar_one()
        assert event.after is not None and event.after["awarded_price"] == "412500.00"

    def test_a_lost_tender_records_why(
        self,
        session: Session,
        estimated: tuple[Bid, Team],
        store: MemoryObjectStore,
        sign_in: SignIn,
    ) -> None:
        bid, team = estimated
        ready_for_g4(session, bid, team)
        service.approve_g4(session, bid, team.director, store)
        session.commit()
        client = sign_in(team.manager_principal)

        recorded = client.post(
            f"/bids/{bid.id}/outcome",
            json={"outcome": "lost", "reasons": "Price", "competitor_feedback": "8% under us"},
        )
        shown = client.get(f"/bids/{bid.id}/outcome").json()

        assert recorded.status_code == 200, recorded.text
        assert (shown["outcome"], shown["state"], shown["reasons"]) == ("lost", "lost", "Price")
        assert shown["competitor_feedback"] == "8% under us" and shown["recorded_by"] == "Bella"


RATE = {
    "item_key": "check_valve|150||||",
    "unit": "no",
    "unit_rate": "640.00",
    "source_type": "purchase_order",
    "source_reference": "PO-26-0455",
    "description": "Check valve DN150, flanged",
    "effective_from": "2026-10-01",
    "valid_until": None,
}


@pytest.mark.req("FR-LRN-03")
class TestLibraryGovernance:
    def test_a_proposed_rate_changes_nothing_until_an_estimator_approves_it(
        self, session: Session, estimated: tuple[Bid, Team], organisation: Organisation
    ) -> None:
        _, team = estimated
        before = {rate.id for rate in pricing.current_rates(session, organisation.id)}

        proposal = governance.propose(
            session,
            organisation.id,
            team.estimator,
            library="rate",
            payload=RATE,
            source="estimator",
            reason="the PO price from the Tampines job",
        )
        session.commit()

        assert proposal.state == "proposed" and proposal.applied_entry_id is None
        assert {rate.id for rate in pricing.current_rates(session, organisation.id)} == before
        with pytest.raises(governance.ProposalError, match="approved by the senior estimator"):
            governance.decide(session, proposal, team.estimator, approve=True)
        with pytest.raises(governance.ProposalError, match="approved by the senior estimator"):
            governance.decide(session, proposal, team.director, approve=True)

        governance.decide(session, proposal, team.senior, approve=True, note="checked the PO")
        again = governance.propose(
            session,
            organisation.id,
            team.estimator,
            library="rate",
            payload={**RATE, "unit_rate": "655.00", "effective_from": "2026-10-03"},
            source="quotation",
            source_ref="Q-2026-2001",
            reason="the supplier's new price",
        )
        governance.decide(session, again, team.senior, approve=True)
        session.commit()

        versions = list(
            session.execute(
                select(Rate)
                .where(
                    Rate.organisation_id == organisation.id, Rate.source_reference == "PO-26-0455"
                )
                .order_by(Rate.version)
            ).scalars()
        )
        assert [(r.version, r.unit_rate.amount, r.retired_at is None) for r in versions] == [
            (1, Decimal("640.00"), False),
            (2, Decimal("655.00"), True),
        ]
        assert versions[1].supersedes_id == versions[0].id
        assert (proposal.state, proposal.decided_by) == ("approved", team.senior.label)
        assert proposal.applied_entry_id == versions[0].id
        actions = [
            event.action
            for event in session.execute(
                select(AuditEvent).where(AuditEvent.entity_id == str(proposal.id))
            ).scalars()
        ]
        assert sorted(actions) == [
            "rate library: change proposed",
            "rate library: proposed change approved",
        ]
        with pytest.raises(governance.ProposalError, match="already approved"):
            governance.decide(session, proposal, team.senior, approve=False, note="no")

    def test_a_rejected_proposal_leaves_the_library_as_it_was(
        self, session: Session, estimated: tuple[Bid, Team], organisation: Organisation
    ) -> None:
        _, team = estimated
        before = [
            (e.id, e.hours_per_unit) for e in labour.current_entries(session, organisation.id)
        ]
        proposal = governance.propose(
            session,
            organisation.id,
            team.estimator,
            library="productivity",
            payload={
                "item_type": "pipe",
                "dn": "50",
                "unit": "m",
                "hours_per_unit": "0.26",
                "trade": "pipefitter",
                "description": "Steel pipe DN50",
                "source_type": "historical_project",
                "source_reference": "Punggol Mall (2026)",
            },
            source="outcome",
            reason="actual hours on the last job",
        )

        with pytest.raises(governance.ProposalError, match="say why the change is not made"):
            governance.decide(session, proposal, team.senior, approve=False)
        governance.decide(
            session, proposal, team.senior, approve=False, note="one job is not a trend"
        )
        session.commit()

        after = [(e.id, e.hours_per_unit) for e in labour.current_entries(session, organisation.id)]
        assert after == before and proposal.state == "rejected"
        with pytest.raises(IntegrityError, match="applied_approved"):
            session.execute(
                text(
                    "UPDATE library_proposal SET applied_entry_id = gen_random_uuid() WHERE id = :p"
                ),
                {"p": proposal.id},
            )
        session.rollback()

    def test_an_approved_productivity_change_is_a_new_version_through_the_api(
        self,
        session: Session,
        estimated: tuple[Bid, Team],
        organisation: Organisation,
        sign_in: SignIn,
    ) -> None:
        _, team = estimated
        payload = {
            "item_type": "pipe",
            "dn": "50",
            "unit": "m",
            "hours_per_unit": "0.26",
            "trade": "pipefitter",
            "description": "Steel pipe DN50",
            "source_type": "historical_project",
            "source_reference": "Punggol Mall (2026)",
        }
        estimator = sign_in(team.estimator_principal)
        proposed = estimator.post(
            "/library-proposals",
            json={
                "library": "productivity",
                "payload": payload,
                "source": "outcome",
                "reason": "actual hours on the last job",
            },
        )
        unsourced = estimator.post(
            "/library-proposals",
            json={
                "library": "productivity",
                "payload": {**payload, "source_reference": " "},
                "source": "estimator",
                "reason": "a hunch",
            },
        )
        refused = estimator.post(
            f"/library-proposals/{proposed.json()['id']}/decide", json={"approve": True}
        )
        senior = sign_in(team.senior_principal)
        approved = senior.post(
            f"/library-proposals/{proposed.json()['id']}/decide", json={"approve": True}
        )
        queue = senior.get("/library-proposals", params={"state": "proposed"}).json()

        assert proposed.status_code == 201, proposed.text
        assert unsourced.status_code == 422 and "says which project" in unsourced.json()["detail"]
        assert refused.status_code == 403
        assert approved.status_code == 200 and approved.json()["state"] == "approved"
        assert queue == []
        current = next(
            e
            for e in labour.current_entries(session, organisation.id)
            if (e.item_type, e.dn) == ("pipe", "50")
        )
        assert (current.version, current.hours_per_unit) == (2, Decimal("0.2600"))
        assert str(current.id) == approved.json()["applied_entry_id"]
        assert len(labour.history(session, current)) == 2
