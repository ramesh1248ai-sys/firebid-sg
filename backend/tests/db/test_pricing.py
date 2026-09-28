# ruff: noqa: F811  (the fixtures below are imported by name and requested by name)
"""Pricing through the database: the rate library, matching, provenance, validity, totals
and the priced exports (FR-CST-01).

The synthetic tender's takeoff is verified and built into a BOQ; the synthetic rate list is
imported. It prices most lines exactly, leaves a tee to a partial match and a check valve and
another tee unpriced, and has one expired rate and one ending before the tender validity.
"""

from __future__ import annotations

import io
import json
import zipfile
from datetime import UTC, date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any

import pytest
from sqlalchemy import text
from sqlalchemy.exc import DBAPIError, IntegrityError
from sqlalchemy.orm import Session

from firebid.ai_gateway.config import load_config
from firebid.ai_gateway.metering import PostgresMeter
from firebid.ai_gateway.providers.fake import FakeAdapter
from firebid.ai_gateway.router import Router
from firebid.boq.xlsx_patch import sheet_parts
from firebid.db.models.commercial import BoqLine, Rate
from firebid.db.models.core import Bid, Organisation
from firebid.domain.actors import Actor
from firebid.domain.state_machines import Role
from firebid.domain.values import Money
from firebid.evals import synthetic_boq, synthetic_rates
from firebid.services import boq, pricing, review
from firebid.storage.object_store import MemoryObjectStore
from tests.db.test_boq import built, client_workbook, line_for, people, verify_takeoff
from tests.db.test_qto import tender  # noqa: F401
from tests.db.test_sheet_views import SignIn, app, member, sign_in  # noqa: F401
from tests.db.test_symbol_mapping import no_tiles, store  # noqa: F401

pytestmark = [pytest.mark.usefixtures("no_tiles"), pytest.mark.req("FR-CST-01")]

TODAY = date(2026, 9, 28)
CONFIG = """
version: 1
providers:
  primary:
    kind: fake
    approved_data_classes: [internal, confidential]
models:
  first:
    provider: primary
    model_id: matcher-1
    capabilities: [structured_output]
    prices:
      - {effective_from: 2026-01-01, input_per_mtok: 5.0, output_per_mtok: 25.0}
routes:
  rate_match:
    requires: [structured_output]
    data_class: confidential
    models: [first]
    reasoning: low
    prompt: rate_match
"""

TEE = "Tee, DN100xDN50 (rule-derived: not drawn)"
OTHER_TEE = "Tee, DN150xDN50 (rule-derived: not drawn)"


def router(session: Session, tmp_path: Path, reply: dict[str, Any]) -> Router:
    path = tmp_path / "llm.yaml"
    path.write_text(CONFIG, encoding="utf-8")
    return Router(
        config=load_config(path),
        adapters={"primary": FakeAdapter("primary").reply(json.dumps(reply))},
        backoff_base_seconds=0,
        sleep=lambda _s: None,
        meter=PostgresMeter(session),
    )


@pytest.fixture
def priced(session: Session, tender: Bid, organisation: Organisation) -> tuple[Bid, Actor, Actor]:
    """Verified takeoff, a built BOQ, the synthetic rate list, a tender valid for 90 days
    from 15 October 2026 (to 13 January 2027), and the rules run."""
    estimator, senior = people(session, organisation, tender)
    tender.submission_deadline = datetime(2026, 10, 15, 12, 0, tzinfo=UTC)
    tender.tender_validity_days = 90
    verify_takeoff(session, tender)
    built(session, tender, estimator)
    report = pricing.import_rates(
        session, organisation.id, synthetic_rates.rate_list(), senior, "rates.xlsx"
    )
    assert report.imported, report.problems
    pricing.price_boq(session, tender, estimator)
    session.commit()
    return tender, estimator, senior


def prices(session: Session, bid: Bid) -> dict[str, pricing.LinePrice]:
    current = boq.current_boq(session, bid.id)
    assert current is not None
    found = pricing.line_prices(session, bid, current, TODAY)
    lines = {line.id: line for line in boq.lines_of(session, current)}
    return {lines[line_id].description: price for line_id, price in found.items()}


def described(session: Session, bid: Bid, item: str) -> BoqLine:
    return line_for(session, bid, item)


class TestProvenance:
    def test_the_domain_refuses_a_price_without_a_rate_entry(
        self, session: Session, priced: tuple[Bid, Actor, Actor]
    ) -> None:
        bid, estimator, _ = priced
        line = described(session, bid, "Check valve, DN150")

        with pytest.raises(pricing.PricingError, match="needs a rate library entry"):
            pricing.price_line(session, line, None, method="person", actor=estimator)

    def test_the_database_refuses_a_price_without_a_rate_entry(
        self, session: Session, priced: tuple[Bid, Actor, Actor]
    ) -> None:
        bid, _, _ = priced
        line = described(session, bid, "Check valve, DN150")

        with pytest.raises(IntegrityError, match="priced_from_rate"):
            session.execute(
                text("UPDATE boq_line SET unit_rate = 120.00, amount = 120.00 WHERE id = :l"),
                {"l": line.id},
            )
        session.rollback()

    def test_the_database_refuses_a_rate_that_is_not_its_entry_s(
        self, session: Session, priced: tuple[Bid, Actor, Actor]
    ) -> None:
        bid, _, _ = priced
        line = described(session, bid, "Sprinkler, pendent")
        assert line.rate_id is not None

        with pytest.raises(DBAPIError, match="priced at its rate entry"):
            session.execute(
                text("UPDATE boq_line SET unit_rate = 1.00 WHERE id = :l"), {"l": line.id}
            )
        session.rollback()
        with pytest.raises(DBAPIError, match="quantity times its rate"):
            session.execute(text("UPDATE boq_line SET amount = 1.00 WHERE id = :l"), {"l": line.id})
        session.rollback()

    def test_an_allowance_is_the_estimator_s_own_and_labelled_so(
        self, session: Session, priced: tuple[Bid, Actor, Actor]
    ) -> None:
        bid, estimator, _ = priced
        allowance = boq.add_marked_line(
            session,
            bid.id,
            estimator,
            description="Allow for testing and commissioning",
            unit="sum",
            quantity=Decimal(1),
            marker="provisional",
            note="client's provisional sum",
            amount=Decimal(5000),
        )
        session.commit()

        assert prices(session, bid)[allowance.description].status == "allowance"


class TestMatching:
    def test_exact_keys_price_lines_with_the_entry_s_rate(
        self, session: Session, priced: tuple[Bid, Actor, Actor]
    ) -> None:
        bid, _, _ = priced

        pendent = described(session, bid, "Sprinkler, pendent")
        branch = described(session, bid, "Pipe, DN50, branch")

        assert (pendent.unit_rate, pendent.amount, pendent.price_method) == (
            Money.of("38.50"),
            Money.of("616.00"),  # 16 heads
            "rule",
        )
        assert (branch.unit_rate, branch.amount) == (Money.of("21.35"), Money.of("1537.20"))
        rate = session.get(Rate, branch.rate_id)
        assert rate is not None and rate.source_reference == "CS-2026"

    def test_a_partial_match_is_proposed_and_waits_for_an_estimator(
        self, session: Session, priced: tuple[Bid, Actor, Actor], tmp_path: Path
    ) -> None:
        bid, estimator, _ = priced
        tee = described(session, bid, TEE)
        [candidate] = [
            r
            for r in pricing.current_rates(session, bid.organisation_id)
            if r.item_key.startswith("fitting_tee|100x50")
        ]
        reply = {
            "choices": [
                {
                    "line": str(tee.id),
                    "entry": str(candidate.id),
                    "confidence": 0.8,
                    "reason": "the same tee; the entry names a brand the line leaves open",
                }
            ]
        }

        pricing.price_boq(session, bid, estimator, router(session, tmp_path, reply))
        session.commit()

        assert (tee.proposed_rate_id, tee.rate_id, tee.unit_rate, tee.amount) == (
            candidate.id,
            None,
            None,
            None,
        )
        assert prices(session, bid)[tee.description].status == "proposed"

        pricing.confirm_proposal(session, bid, tee, estimator)
        session.commit()

        assert (tee.unit_rate, tee.amount, tee.price_method) == (
            Money.of("41.80"),
            Money.of("125.40"),  # 3 tees
            "model",
        )
        assert tee.price_provenance["model"] == "matcher-1"

    def test_a_line_with_no_entry_shows_unpriced(
        self, session: Session, priced: tuple[Bid, Actor, Actor]
    ) -> None:
        bid, _, _ = priced

        found = prices(session, bid)

        assert found[described(session, bid, "Check valve, DN150").description].status == "unpriced"
        assert found[described(session, bid, OTHER_TEE).description].status == "unpriced"

    def test_a_person_s_choice_survives_rules_and_a_rebuild(
        self, session: Session, priced: tuple[Bid, Actor, Actor]
    ) -> None:
        bid, estimator, _ = priced
        tee = described(session, bid, TEE)
        [candidate] = [
            r
            for r in pricing.current_rates(session, bid.organisation_id)
            if r.item_key.startswith("fitting_tee|100x50")
        ]
        pricing.choose_rate(session, bid, tee, candidate.id, estimator, "checked the catalogue")
        session.commit()

        built(session, bid, estimator)

        again = described(session, bid, TEE)
        assert again.id != tee.id, "a new BOQ version"
        assert (again.rate_id, again.price_method) == (candidate.id, "person")


class TestValidity:
    def test_a_rate_ending_before_the_tender_validity_warns(
        self, session: Session, priced: tuple[Bid, Actor, Actor]
    ) -> None:
        bid, _, _ = priced

        main = prices(session, bid)[described(session, bid, "Pipe, DN150, main").description]

        assert [w.code for w in main.warnings] == ["ends_before_tender_validity"]

    def test_an_expired_rate_warns(
        self, session: Session, priced: tuple[Bid, Actor, Actor]
    ) -> None:
        bid, _, _ = priced

        main = prices(session, bid)[described(session, bid, "Pipe, DN100, main").description]

        assert "expired" in [w.code for w in main.warnings]


class TestLibrary:
    def test_a_changed_rate_is_a_new_version_and_the_old_one_is_kept(
        self, session: Session, priced: tuple[Bid, Actor, Actor], organisation: Organisation
    ) -> None:
        bid, _, senior = priced
        changed = tuple(
            row
            if row.dn != "50"
            else synthetic_rates.RateRow(
                row.type,
                row.dn,
                row.description,
                row.unit,
                "22.10",
                row.source_type,
                row.source_reference,
                date(2026, 9, 1),
                row.valid_until,
            )
            for row in synthetic_rates.RATES
        )

        report = pricing.import_rates(
            session, organisation.id, synthetic_rates.rate_list(changed), senior, "rates-v2.xlsx"
        )
        session.commit()

        assert (report.created, report.superseded, report.unchanged) == (0, 1, 9)
        branch = described(session, bid, "Pipe, DN50, branch")
        old = session.get(Rate, branch.rate_id)
        assert old is not None and old.retired_at is not None, "the line keeps its source"
        versions = pricing.history(session, old)
        assert [(v.version, str(v.unit_rate.amount)) for v in versions] == [
            (1, "21.35"),
            (2, "22.10"),
        ]

    def test_a_list_with_any_error_imports_nothing(
        self, session: Session, organisation: Organisation
    ) -> None:
        senior = member_actor(session, organisation)
        bad = synthetic_rates.rate_list(
            extra=[
                (
                    "pipe",
                    "65",
                    "",
                    "",
                    "",
                    "",
                    "Pipe DN65",
                    "m",
                    -3,
                    "PO",
                    "PO-1",
                    date(2026, 1, 1),
                    None,
                )
            ]
        )

        report = pricing.import_rates(session, organisation.id, bad, senior, "bad.xlsx")

        assert report.imported is False and report.problems
        assert pricing.current_rates(session, organisation.id) == []


def member_actor(session: Session, organisation: Organisation) -> Actor:
    return Actor(label="Sam Senior", roles=frozenset({str(Role.SENIOR_ESTIMATOR)}), id=None)


class TestTotals:
    def test_the_boq_adds_up_line_by_line(
        self, session: Session, priced: tuple[Bid, Actor, Actor]
    ) -> None:
        bid, _, _ = priced
        current = boq.current_boq(session, bid.id)
        assert current is not None

        found = pricing.boq_totals(session, current)

        by_hand = sum(
            (
                line.amount.amount
                for line in boq.lines_of(session, current)
                if line.amount and line.rate_id
            ),
            Decimal("0.00"),
        )
        assert found.priced.amount == by_hand
        assert found.grand == found.priced + found.allowances
        assert found.unpriced == 3, "the check valve and two tees"


class TestExports:
    def test_the_priced_client_workbook_round_trips_with_rates_filled(
        self,
        session: Session,
        priced: tuple[Bid, Actor, Actor],
        store: MemoryObjectStore,
    ) -> None:
        bid, estimator, _ = priced
        document = client_workbook(session, bid, store)
        boq.read_client_boq(session, store, document)
        boq.propose_mappings(session, bid.id)
        for row in boq.mappings(session, bid.id).values():
            if row.boq_line_key:
                boq.decide_mapping(session, row, estimator, decision="confirm")
        session.commit()

        priced_book = boq.priced_client_workbook(session, store, bid.id, document.id)

        original = store.get(document.storage_key)
        before = zipfile.ZipFile(io.BytesIO(original))
        after = zipfile.ZipFile(io.BytesIO(priced_book))
        assert before.namelist() == after.namelist()
        changed = {n for n in before.namelist() if before.read(n) != after.read(n)}
        assert changed == {sheet_parts(before)[synthetic_boq.BILL]}
        from openpyxl import load_workbook

        sheet = load_workbook(io.BytesIO(priced_book))[synthetic_boq.BILL]
        cells = synthetic_boq.client_boq().cells
        assert sheet[cells["A1"][1]].value == 38.5  # pendent heads
        assert sheet[cells["B3"][1]].value == 21.35  # DN50 branch
        assert sheet[cells["C2"][1]].value is None, "the check valve is unpriced"

    def test_our_boq_export_carries_rates_sources_and_unpriced(
        self, session: Session, priced: tuple[Bid, Actor, Actor]
    ) -> None:
        from openpyxl import load_workbook

        bid, _, _ = priced

        sheet = load_workbook(io.BytesIO(boq.company_workbook(session, bid.id)))["BOQ"]

        rows = [list(row) for row in sheet.iter_rows(values_only=True)]
        header = rows[0]
        rate, source = header.index("Rate"), header.index("Source")
        pendent = next(r for r in rows if r[1] and str(r[1]).startswith("Pendent"))
        valve = next(r for r in rows if r[1] and "check valve" in str(r[1]).lower())
        assert pendent[rate] == 38.5 and "CS-2026" in str(pendent[source])
        assert valve[rate] is None and valve[source] == "unpriced"


class TestQueue:
    def test_a_priced_item_is_weighed_by_its_rate(
        self, session: Session, priced: tuple[Bid, Actor, Actor]
    ) -> None:
        bid, _, _ = priced

        rows = {row.item.description: row for row in review.queue(session, bid.id)}
        found = review.coverage(session, bid.id)

        assert rows["Sprinkler, pendent"].impact == pytest.approx(16 * 38.50)
        assert rows["Check valve, DN150"].impact == pytest.approx(1 * 25.0 * 40.0)
        assert found["value_from_rates"] > 0 and "from rates" in found["value_basis"]


class TestApi:
    def test_only_a_senior_estimator_imports_rates_and_a_bad_list_is_reported(
        self, session: Session, bid: Bid, organisation: Organisation, sign_in: SignIn
    ) -> None:
        files = {"file": ("rates.xlsx", synthetic_rates.rate_list(), "application/octet-stream")}
        # One person signed in at a time: sign-in replaces the app's principal.
        estimator = sign_in(member(session, organisation, bid, "esther", Role.ESTIMATOR))
        refused = estimator.post("/rates/import", files=files)
        senior = sign_in(member(session, organisation, bid, "sam", Role.SENIOR_ESTIMATOR))
        imported = senior.post("/rates/import", files=files)
        bad = senior.post(
            "/rates/import",
            files={
                "file": (
                    "bad.xlsx",
                    synthetic_rates.rate_list(
                        extra=[
                            (
                                "pipe",
                                "65",
                                "",
                                "",
                                "",
                                "",
                                "Pipe",
                                "m",
                                -3,
                                "PO",
                                "PO-1",
                                date(2026, 1, 1),
                                None,
                            )
                        ]
                    ),
                    "application/octet-stream",
                )
            },
        )

        assert refused.status_code == 403
        assert imported.status_code == 200, imported.text
        assert imported.json()["created"] == len(synthetic_rates.RATES)
        assert bad.json()["imported"] is False
        assert {p["column"] for p in bad.json()["problems"]} == {"rate"}
        listed = senior.get("/rates").json()
        history = senior.get(f"/rates/{listed[0]['id']}/history").json()
        assert len(listed) == len(synthetic_rates.RATES) and len(history) == 1

    def test_an_estimator_prices_the_boq_and_sees_what_is_unpriced(
        self,
        session: Session,
        priced: tuple[Bid, Actor, Actor],
        organisation: Organisation,
        sign_in: SignIn,
    ) -> None:
        bid, _, _ = priced
        client = sign_in(member(session, organisation, bid, "esther", Role.ESTIMATOR))

        run = client.post(f"/bids/{bid.id}/pricing/run")

        assert run.status_code == 200, run.text
        body = run.json()
        by_line = {line["description"]: line for line in body["lines"]}
        assert by_line["Pendent sprinkler head"]["status"] == "priced" or any(
            line["status"] == "priced" and line["unit_rate"] == "38.50" for line in body["lines"]
        )
        statuses = {line["status"] for line in body["lines"]}
        assert {"priced", "unpriced"} <= statuses
        assert body["tender_validity_end"] == "2027-01-13"
        assert body["totals"]["gst_included"] is False
        assert Decimal(body["totals"]["grand"]) == Decimal(body["totals"]["priced"]) + Decimal(
            body["totals"]["allowances"]
        )
        tee = next(line for line in body["lines"] if line["awaiting_model"])
        options = client.get(f"/bids/{bid.id}/pricing/lines/{tee['line_id']}/candidates").json()
        assert [o["key_parts"]["brand"] for o in options] == ["victaulic"]

        chosen = client.post(
            f"/bids/{bid.id}/pricing/lines/{tee['line_id']}/rate",
            json={"rate_id": options[0]["id"], "note": "checked the catalogue"},
        )

        line = next(x for x in chosen.json()["lines"] if x["line_id"] == tee["line_id"])
        assert (line["status"], line["method"], line["unit_rate"]) == ("priced", "person", "41.80")
        queued = session.execute(
            text("SELECT count(*) FROM procrastinate_jobs WHERE task_name = 'pricing.match'")
        ).scalar_one()
        assert queued >= 1
