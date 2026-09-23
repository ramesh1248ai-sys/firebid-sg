"""The estimator template and its importer (decision D3, FR-LRN-01).

The importer's error messages are the feature. An estimator who fills this in after hours and
gets a traceback will not fill in the second one, so most of these tests are about what a bad
workbook says rather than what a good one does.
"""

from __future__ import annotations

from datetime import date
from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook

from firebid.evals.cli import main
from firebid.evals.importer import ImportFailed, import_workbook
from firebid.evals.schema import InputClass, ObjectType, RevisionStatus
from firebid.evals.template import (
    BOQ,
    COUNTS,
    DUPLICATES,
    INSTRUCTIONS,
    LISTS,
    PIPES,
    SHEETS,
    TENDER,
    build_workbook,
    write_template,
)

pytestmark = pytest.mark.req("FR-LRN-01")

FIRST = 3  # the first data row, below the header and the hint


def filled(tmp_path: Path, **changes: Any) -> Path:
    """A workbook filled the way an estimator would fill it, with optional damage."""
    book = build_workbook()

    tender = book[TENDER]
    tender.cell(row=FIRST, column=1, value=changes.get("tender_id", "MC-2024-014"))
    tender.cell(row=FIRST, column=2, value="Consultant Engineers LLP")
    tender.cell(row=FIRST, column=3, value=changes.get("received_on", "2024-06-03"))
    tender.cell(row=FIRST, column=4, value=changes.get("input_class", "vector_pdf"))
    tender.cell(row=FIRST, column=5, value="Two enlarged plans repeat the level 5 GA.")

    sheets = book[SHEETS]
    rows = changes.get(
        "sheets",
        [
            ("FP-L05-201", "R04", "current", "vector_pdf", "no"),
            ("FP-L05-201", "R03", "superseded", "vector_pdf", "no"),
            ("FP-L05-202", "R01", "current", "vector_pdf", "no"),
            ("FP-SCH-001", "R01", "current", "vector_pdf", "yes"),
        ],
    )
    for offset, values in enumerate(rows):
        for column, value in enumerate(values, start=1):
            sheets.cell(row=FIRST + offset, column=column, value=value)

    counts = book[COUNTS]
    for offset, values in enumerate(
        changes.get(
            "counts",
            [
                ("FP-L05-201", "sprinkler_pendent", 124),
                ("FP-L05-201", "hose_reel", 4),
                ("FP-L05-202", "sprinkler_pendent", 38),
                ("FP-SCH-001", "landing_valve", 0),
            ],
        )
    ):
        for column, value in enumerate(values, start=1):
            counts.cell(row=FIRST + offset, column=column, value=value)

    pipes = book[PIPES]
    for offset, values in enumerate(
        changes.get(
            "pipes",
            [("FP-L05-201", 150, 84_200), ("FP-L05-201", 100, 39_500), ("FP-L05-202", 50, 12_000)],
        )
    ):
        for column, value in enumerate(values, start=1):
            pipes.cell(row=FIRST + offset, column=column, value=value)

    duplicates = book[DUPLICATES]
    for offset, values in enumerate(changes.get("duplicates", [("FP-L05-202", "FP-L05-201")])):
        for column, value in enumerate(values, start=1):
            duplicates.cell(row=FIRST + offset, column=column, value=value)

    boq = book[BOQ]
    for offset, values in enumerate(
        changes.get(
            "boq",
            [
                ("2.4.1", "Pendent sprinkler heads", "nr", 162, "sprinkler_pendent"),
                ("2.4.2", "Hose reel assemblies", "nr", 4, "hose_reel"),
                ("1.1.1", "Preliminaries", "item", 1, None),
            ],
        )
    ):
        for column, value in enumerate(values, start=1):
            boq.cell(row=FIRST + offset, column=column, value=value)

    path = tmp_path / "filled.xlsx"
    book.save(path)
    return path


class TestTheTemplateItself:
    def test_it_has_every_tab_an_estimator_needs(self, tmp_path: Path) -> None:
        path = write_template(tmp_path / "blank.xlsx")
        names = load_workbook(path).sheetnames
        assert {TENDER, SHEETS, COUNTS, PIPES, DUPLICATES, BOQ, INSTRUCTIONS} <= set(names)

    def test_the_instructions_come_first(self, tmp_path: Path) -> None:
        """Whoever opens it should land on what 'verified' means, not on an empty grid."""
        path = write_template(tmp_path / "blank.xlsx")
        assert load_workbook(path).sheetnames[0] == INSTRUCTIONS

    def test_the_machinery_tab_is_hidden(self, tmp_path: Path) -> None:
        path = write_template(tmp_path / "blank.xlsx")
        assert load_workbook(path)[LISTS].sheet_state == "hidden"

    def test_the_dropdowns_refuse_anything_off_the_list(self, tmp_path: Path) -> None:
        """A warning that can be clicked past is how 'DN65 ' gets into a dataset."""
        path = write_template(tmp_path / "blank.xlsx")
        book = load_workbook(path)
        validations = book[COUNTS].data_validations.dataValidation
        assert validations, "the object type column has no dropdown"
        assert all(validation.showErrorMessage for validation in validations)

    def test_a_blank_template_lists_every_object_type_the_platform_counts(
        self, tmp_path: Path
    ) -> None:
        """If the vocabulary drifts, an estimator cannot record what the platform detects."""
        path = write_template(tmp_path / "blank.xlsx")
        lists = load_workbook(path)[LISTS]
        offered = {row[0] for row in lists.iter_rows(min_col=1, max_col=1, values_only=True)}
        assert {str(value) for value in ObjectType} <= offered


class TestAGoodWorkbook:
    def test_it_imports(self, tmp_path: Path) -> None:
        truth = import_workbook(filled(tmp_path))
        assert truth.tender_id == "MC-2024-014"
        assert truth.consultant == "Consultant Engineers LLP"
        assert truth.received_on == date(2024, 6, 3)
        assert truth.input_class is InputClass.VECTOR_PDF

    def test_superseded_revisions_are_kept(self, tmp_path: Path) -> None:
        """How the platform handles them is itself being measured, so they cannot be dropped."""
        truth = import_workbook(filled(tmp_path))
        statuses = {(s.sheet_number, s.revision): s.status for s in truth.sheets}
        assert statuses[("FP-L05-201", "R03")] is RevisionStatus.SUPERSEDED
        assert statuses[("FP-L05-201", "R04")] is RevisionStatus.CURRENT
        assert len(truth.current_sheets()) == 3

    def test_counts_land_on_the_current_revision(self, tmp_path: Path) -> None:
        truth = import_workbook(filled(tmp_path))
        sheet = next(s for s in truth.current_sheets() if s.sheet_number == "FP-L05-201")
        assert sheet.count_of(ObjectType.SPRINKLER_PENDENT) == 124
        assert sheet.count_of(ObjectType.HOSE_REEL) == 4

    def test_a_zero_count_is_kept_as_a_real_answer(self, tmp_path: Path) -> None:
        """Zero means 'I looked and there were none', which is not the same as silence."""
        truth = import_workbook(filled(tmp_path))
        schedule = next(s for s in truth.sheets if s.sheet_number == "FP-SCH-001")
        assert schedule.count_of(ObjectType.LANDING_VALVE) == 0
        assert any(entry.object_type is ObjectType.LANDING_VALVE for entry in schedule.counts)

    def test_pipe_lengths_stay_whole_millimetres(self, tmp_path: Path) -> None:
        truth = import_workbook(filled(tmp_path))
        sheet = next(s for s in truth.current_sheets() if s.sheet_number == "FP-L05-201")
        assert sheet.length_at(150) == 84_200
        assert all(isinstance(entry.length_mm, int) for entry in sheet.pipe_lengths)

    def test_duplicates_are_recorded(self, tmp_path: Path) -> None:
        truth = import_workbook(filled(tmp_path))
        sheet = next(s for s in truth.current_sheets() if s.sheet_number == "FP-L05-202")
        assert sheet.duplicates_of == ("FP-L05-201",)

    def test_a_boq_line_that_maps_to_nothing_is_allowed(self, tmp_path: Path) -> None:
        """Preliminaries and PC sums map to no object; forcing a choice would be a lie."""
        truth = import_workbook(filled(tmp_path))
        prelims = next(line for line in truth.boq_lines if line.line_reference == "1.1.1")
        assert prelims.maps_to is None

    def test_a_spreadsheet_integer_stored_as_a_float_is_accepted(self, tmp_path: Path) -> None:
        """Excel stores 124 as 124.0; that is not a fractional answer."""
        path = filled(tmp_path, counts=[("FP-L05-201", "sprinkler_pendent", 124.0)])
        truth = import_workbook(path)
        sheet = next(s for s in truth.current_sheets() if s.sheet_number == "FP-L05-201")
        assert sheet.count_of(ObjectType.SPRINKLER_PENDENT) == 124


class TestABadWorkbookExplainsItself:
    def _problems(self, path: Path) -> tuple[list[str], str]:
        with pytest.raises(ImportFailed) as caught:
            import_workbook(path)
        return [str(problem) for problem in caught.value.problems], caught.value.report()

    def test_a_count_for_a_sheet_that_was_never_listed(self, tmp_path: Path) -> None:
        path = filled(tmp_path, counts=[("FP-L09-999", "sprinkler_pendent", 10)])
        problems, report = self._problems(path)
        assert any("FP-L09-999" in problem and SHEETS in problem for problem in problems)
        assert "row 3" in report

    def test_a_misspelled_object_type_lists_what_is_allowed(self, tmp_path: Path) -> None:
        path = filled(tmp_path, counts=[("FP-L05-201", "sprinkler pendant", 10)])
        problems, _ = self._problems(path)
        assert any("sprinkler_pendent" in problem for problem in problems), (
            "the message should say what is allowed, not only that the value is wrong"
        )

    def test_a_count_that_is_not_a_number(self, tmp_path: Path) -> None:
        path = filled(tmp_path, counts=[("FP-L05-201", "sprinkler_pendent", "about 120")])
        problems, _ = self._problems(path)
        assert any("not a number" in problem for problem in problems)

    def test_a_fractional_count(self, tmp_path: Path) -> None:
        path = filled(tmp_path, counts=[("FP-L05-201", "sprinkler_pendent", 12.5)])
        problems, _ = self._problems(path)
        assert any("whole number" in problem for problem in problems)

    def test_two_current_revisions_of_one_sheet(self, tmp_path: Path) -> None:
        """The platform has to work this out later; the yardstick cannot be ambiguous."""
        path = filled(
            tmp_path,
            sheets=[
                ("FP-L05-201", "R04", "current", "vector_pdf", "no"),
                ("FP-L05-201", "R03", "current", "vector_pdf", "no"),
            ],
            counts=[("FP-L05-201", "sprinkler_pendent", 10)],
            pipes=[],
            duplicates=[],
        )
        problems, _ = self._problems(path)
        assert any("current revision" in problem for problem in problems)

    def test_lengths_on_a_not_to_scale_sheet(self, tmp_path: Path) -> None:
        path = filled(tmp_path, pipes=[("FP-SCH-001", 150, 1000)])
        problems, _ = self._problems(path)
        assert any("not-to-scale" in problem for problem in problems)

    def test_a_sheet_that_duplicates_itself(self, tmp_path: Path) -> None:
        path = filled(tmp_path, duplicates=[("FP-L05-201", "FP-L05-201")])
        problems, _ = self._problems(path)
        assert any("cannot duplicate itself" in problem for problem in problems)

    def test_the_same_object_counted_twice_on_one_sheet(self, tmp_path: Path) -> None:
        path = filled(
            tmp_path,
            counts=[
                ("FP-L05-201", "sprinkler_pendent", 124),
                ("FP-L05-201", "sprinkler_pendent", 38),
            ],
        )
        problems, _ = self._problems(path)
        assert any("already counted on row" in problem for problem in problems)

    def test_an_empty_workbook_says_what_is_missing(self, tmp_path: Path) -> None:
        path = tmp_path / "blank.xlsx"
        write_template(path)
        problems, _ = self._problems(path)
        assert any("no tender row" in problem for problem in problems)
        assert any("no sheets listed" in problem for problem in problems)

    def test_the_wrong_file_entirely(self, tmp_path: Path) -> None:
        path = tmp_path / "notes.txt"
        path.write_text("these are my notes", encoding="utf-8")
        problems, _ = self._problems(path)
        assert any("cannot be opened" in problem for problem in problems)

    def test_every_problem_is_reported_at_once(self, tmp_path: Path) -> None:
        """Fixing one error, re-running, and finding the next is how goodwill runs out."""
        path = filled(
            tmp_path,
            counts=[
                ("FP-L09-999", "sprinkler_pendent", 10),
                ("FP-L05-201", "not_a_type", 5),
                ("FP-L05-201", "hose_reel", "lots"),
            ],
        )
        problems, report = self._problems(path)
        assert len(problems) >= 3
        assert report.count("row ") >= 3

    def test_the_report_is_grouped_by_tab_and_ordered_by_row(self, tmp_path: Path) -> None:
        path = filled(
            tmp_path,
            counts=[
                ("FP-L05-201", "bad_type_two", 5),
                ("FP-L05-201", "bad_type_one", 5),
            ],
        )
        _, report = self._problems(path)
        assert COUNTS in report
        assert report.index("row 3") < report.index("row 4")


class TestTheCommandLine:
    def test_template_writes_a_workbook(self, tmp_path: Path, capsys: Any) -> None:
        path = tmp_path / "out" / "golden_takeoff.xlsx"
        assert main(["template", "--out", str(path)]) == 0
        assert path.exists()
        assert "wrote" in capsys.readouterr().out

    def test_import_of_a_good_workbook_succeeds_and_can_write_json(
        self, tmp_path: Path, capsys: Any
    ) -> None:
        source = filled(tmp_path)
        destination = tmp_path / "truth.json"
        assert main(["import", str(source), "--out", str(destination)]) == 0
        assert destination.exists()
        assert "MC-2024-014" in capsys.readouterr().out

    def test_import_of_a_bad_workbook_fails_and_writes_nothing(
        self, tmp_path: Path, capsys: Any
    ) -> None:
        """A partial import is worse than none: it looks like it worked."""
        source = filled(tmp_path, counts=[("FP-L09-999", "sprinkler_pendent", 10)])
        destination = tmp_path / "truth.json"
        assert main(["import", str(source), "--out", str(destination)]) == 1
        assert not destination.exists()
        assert "problem" in capsys.readouterr().err
