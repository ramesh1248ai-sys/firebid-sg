"""Tender clarifications without a database: the draft and its evidence, engineering
content, grouping, the lifecycle's guards, and the exported register."""

from __future__ import annotations

import ast
import io
from pathlib import Path

import pytest
from docx import Document
from openpyxl import load_workbook
from pydantic import ValidationError

from firebid.clarifications import drafting, export
from firebid.clarifications.drafting import Candidate, EvidenceRef, SheetRef
from firebid.domain.state_machines import (
    CLARIFICATION,
    ClarificationState,
    Role,
    TransitionError,
    plan_transition,
)

CLAUSE = EvidenceRef(
    kind="clause",
    label="Specification clause 2.1.3",
    quote="All sprinkler pipework within the basement car park shall be hot-dip galvanised.",
    revision="B",
)
SHEET = EvidenceRef(
    kind="sheet",
    label="Drawing FP-B1-201",
    quote="ALL SPRINKLER PIPEWORK TO BE BLACK STEEL",
    revision="R01",
)


def conflict(**values: object) -> Candidate:
    base: dict[str, object] = {
        "kind": "spec_issue",
        "ref": "issue-1",
        "subject": "Pipe material: the specification says galvanised steel, sheet FP-B1-201 "
        "says black steel",
        "problem": "The two do not agree on the pipe material. Please confirm which governs.",
        "evidence": (CLAUSE, SHEET),
        "system": "sprinkler",
        "level_grid": "B1",
        "sheets": (SheetRef(sheet_number="FP-B1-201", revision="R01"),),
        "options": ("the specification governs (galvanised steel)",),
        "topic": "pipe_material",
        "group_key": "sprinkler|FP-B1-201",
    }
    return Candidate(**{**base, **values})  # type: ignore[arg-type]


def variance(**values: object) -> Candidate:
    base: dict[str, object] = {
        "kind": "boq_variance",
        "ref": "client:Bill 1!12",
        "subject": "Quantity of 50 mm diameter pipe",
        "problem": "The bill gives 70 m; the drawings measure 72 m. Please confirm the quantity.",
        "evidence": (
            EvidenceRef(kind="client_boq_line", label="Client bill Bill 1!12", quote="B1: 70 m"),
            EvidenceRef(kind="qto_item", label="Takeoff QTO-000007"),
        ),
        "options": ("the quantity measured from the drawings (72 m) is priced",),
        "cost_impact": "2 m at the bill's rate",
        "topic": "quantity",
    }
    return Candidate(**{**base, **values})  # type: ignore[arg-type]


@pytest.mark.req("FR-RFI-02")
class TestDraft:
    def test_a_spec_conflict_drafts_with_every_field_and_its_evidence(self) -> None:
        found = drafting.compose([conflict()], "Marina Bay Commercial Tower")

        assert found.subject.startswith("Pipe material")
        assert found.project == "Marina Bay Commercial Tower"
        assert found.level_grid == "B1"
        assert [(s.sheet_number, s.revision) for s in found.sheets] == [("FP-B1-201", "R01")]
        assert "Please confirm which governs" in found.problem
        assert [e.label for e in found.evidence] == [
            "Specification clause 2.1.3",
            "Drawing FP-B1-201",
        ]
        assert found.required_reviewer == "design_manager"

    def test_a_bill_variance_drafts_with_its_cost_impact_for_the_bid_manager(self) -> None:
        found = drafting.compose([variance()], "Marina Bay Commercial Tower")

        assert [e.kind for e in found.evidence] == ["client_boq_line", "qto_item"]
        assert found.cost_impact == "2 m at the bill's rate"
        assert (found.required_reviewer, found.engineering_content) == ("bid_manager", False)

    def test_a_draft_with_no_evidence_is_refused(self) -> None:
        with pytest.raises(ValidationError, match="evidence"):
            drafting.compose([variance(evidence=())], "Tower")
        with pytest.raises(ValueError, match="at least one flagged issue"):
            drafting.compose([], "Tower")


@pytest.mark.req("FR-RFI-06")
class TestOptionsAndEngineering:
    def test_every_option_is_worded_as_a_recommendation(self) -> None:
        found = drafting.compose([conflict()], "Tower")

        assert [option.text for option in found.options] == [
            "Recommendation: the specification governs (galvanised steel)"
        ]
        assert all(option.recommendation for option in found.options)
        with pytest.raises(ValidationError, match="recommendation only"):
            drafting.Option(text="Use black steel", recommendation=False)

    @pytest.mark.parametrize(
        ("candidate", "words", "flagged", "why"),
        [
            (conflict(), "", True, "it concerns pipe material"),
            (variance(), "", False, ""),
            (variance(), "The pump duty should be confirmed", True, "it mentions pump duty"),
            (variance(), "confirm the hydraulically remote area", False, ""),
            (variance(), "to SS CP 52", True, "it mentions SS CP 52"),
        ],
    )
    def test_engineering_content_is_told_by_subject_or_by_wording(
        self, candidate: Candidate, words: str, flagged: bool, why: str
    ) -> None:
        assert drafting.engineering([candidate], words) == (flagged, why)

    def test_an_engineering_clarification_cannot_be_approved_without_the_design_manager(
        self,
    ) -> None:
        def approve(context: dict[str, object], roles: set[str]) -> None:
            plan_transition(
                CLARIFICATION,
                source=ClarificationState.INTERNAL_REVIEW,
                target=ClarificationState.APPROVED_TO_ISSUE,
                actor_roles=roles,
                context=context,
            )

        bid_manager = {str(Role.BID_MANAGER)}
        engineering = {"has_evidence": True, "engineering_content": True}
        with pytest.raises(TransitionError, match="the Design Manager's approval is missing"):
            approve({**engineering, "design_manager_approved": False}, bid_manager)
        approve({**engineering, "design_manager_approved": True}, bid_manager)
        # A plain one needs only the Bid Manager.
        approve({"has_evidence": True, "engineering_content": False}, bid_manager)
        # Nobody else approves, and the Design Manager's approval does not stand in for it.
        for other in (Role.ESTIMATOR, Role.SENIOR_ESTIMATOR, Role.DESIGN_MANAGER):
            with pytest.raises(TransitionError, match="needs one of these roles: bid_manager"):
                approve({**engineering, "design_manager_approved": True}, {str(other)})
        with pytest.raises(TransitionError, match="no evidence reference"):
            approve({"has_evidence": False}, bid_manager)


@pytest.mark.req("FR-RFI-04")
class TestLifecycle:
    def test_the_lifecycle_is_the_model_of_the_requirements(self) -> None:
        def after(state: ClarificationState) -> set[str]:
            return {str(target) for target in CLARIFICATION.allowed_targets(state)}

        assert after(ClarificationState.DRAFT) == {"internal_review", "converted_to_qualification"}
        assert after(ClarificationState.INTERNAL_REVIEW) == {
            "draft",
            "approved_to_issue",
            "converted_to_qualification",
        }
        assert after(ClarificationState.APPROVED_TO_ISSUE) == {
            "internal_review",
            "issued",
            "converted_to_qualification",
        }
        assert after(ClarificationState.ISSUED) == {"responded", "converted_to_qualification"}
        assert after(ClarificationState.RESPONDED) == {
            "closed_incorporated",
            "closed_no_change",
            "converted_to_qualification",
        }
        for closed in CLARIFICATION.terminal:
            assert after(closed) == set()  # type: ignore[arg-type]

    def test_a_response_is_closed_only_once_its_impact_is_assessed(self) -> None:
        def close(assessed: bool) -> None:
            plan_transition(
                CLARIFICATION,
                source=ClarificationState.RESPONDED,
                target=ClarificationState.CLOSED_INCORPORATED,
                actor_roles={str(Role.ESTIMATOR)},
                context={"impact_assessed": assessed},
            )

        with pytest.raises(TransitionError, match="impact is not assessed"):
            close(False)
        close(True)


@pytest.mark.req("FR-RFI-07")
class TestGroupingAndExport:
    def test_related_issues_are_proposed_as_one_clarification(self) -> None:
        joining = conflict(ref="issue-2", subject="Joining method", topic="joining_method")
        elsewhere = conflict(ref="issue-3", group_key="sprinkler|FP-L05-201")

        groups = drafting.propose_groups([conflict(), variance(), joining, elsewhere])

        assert [(g.key, [c.ref for c in g.candidates]) for g in groups] == [
            ("sprinkler|FP-B1-201", ["issue-1", "issue-2"]),
            ("boq_variance:client:Bill 1!12", ["client:Bill 1!12"]),
            ("sprinkler|FP-L05-201", ["issue-3"]),
        ]
        assert groups[0].reason == "2 issues about sprinkler, FP-B1-201"
        assert groups[1].reason == "on its own"

    def test_a_confirmed_group_is_one_draft_with_every_issue_and_all_the_evidence(self) -> None:
        joining = conflict(
            ref="issue-2",
            subject="Joining method",
            problem="The two do not agree on the joining method.",
            evidence=(SHEET, EvidenceRef(kind="clause", label="Specification clause 2.1.2")),
            options=("the specification governs (grooved)",),
        )

        found = drafting.compose([conflict(), joining], "Tower")

        assert found.subject == "2 queries on the sprinkler installation on FP-B1-201"
        assert found.problem.splitlines() == [
            "1. The two do not agree on the pipe material. Please confirm which governs.",
            "2. The two do not agree on the joining method.",
        ]
        assert [e.label for e in found.evidence] == [
            "Specification clause 2.1.3",
            "Drawing FP-B1-201",
            "Specification clause 2.1.2",
        ]
        assert len(found.options) == 2 and len(found.sheets) == 1

    ROWS = (
        {
            "number": "TC-001",
            "subject": "Pipe material",
            "level_grid": "B1",
            "sheets": "FP-B1-201 rev R01",
            "problem": "Please confirm which governs.",
            "evidence": "Specification clause 2.1.3 rev B",
            "options": "Recommendation: the specification governs",
            "status": "Approved to issue",
            "due": "2026-10-18",
        },
    )

    @pytest.mark.parametrize("key", ["company_default", "client_query_log"])
    def test_the_workbook_has_the_chosen_template_s_headings_and_fields(self, key: str) -> None:
        chosen = export.template(key)

        sheet = load_workbook(
            io.BytesIO(export.build(key, "xlsx", "Tower", list(self.ROWS)))
        ).active
        assert sheet is not None
        rows = [[cell or "" for cell in row] for row in sheet.iter_rows(values_only=True)]

        assert rows[0][0] == chosen.title and rows[1][0] == "Tower"
        assert rows[3] == [heading for heading, _ in chosen.columns]
        assert rows[4] == [self.ROWS[0].get(name, "") for _, name in chosen.columns]

    def test_the_client_s_layout_differs_from_the_company_s(self) -> None:
        client = export.template("client_query_log")

        assert [heading for heading, _ in client.columns] == [
            "Query No",
            "Drawing / Spec Ref",
            "Location",
            "Tenderer's Query",
            "Tenderer's Proposal",
            "Consultant's Response",
        ]
        assert client.columns != export.template("company_default").columns

    def test_the_document_has_the_same_table(self) -> None:
        file = Document(
            io.BytesIO(export.build("client_query_log", "docx", "Tower", list(self.ROWS)))
        )

        [table] = file.tables
        assert [cell.text for cell in table.rows[0].cells][:2] == ["Query No", "Drawing / Spec Ref"]
        assert table.rows[1].cells[0].text == "TC-001"

    def test_an_unknown_template_or_format_is_refused(self) -> None:
        with pytest.raises(KeyError, match="no clarification template"):
            export.build("somebody_else", "xlsx", "Tower", [])
        with pytest.raises(ValueError, match="xlsx or docx"):
            export.build("company_default", "pdf", "Tower", [])

    def test_nothing_in_the_clarification_code_can_send_anything(self) -> None:
        """The only way out is a download: no module here opens a connection, sends mail or
        queues a job that would."""
        root = Path(drafting.__file__).resolve().parents[1]
        files = [
            *(root / "clarifications").glob("*.py"),
            root / "services" / "clarifications.py",
            root / "api" / "clarifications.py",
            root / "agents" / "clarification_drafter.py",
        ]
        forbidden = {
            "smtplib",
            "email",
            "socket",
            "http",
            "urllib",
            "httpx",
            "requests",
            "aiohttp",
            "ftplib",
            "firebid.jobs",
        }
        for file in files:
            tree = ast.parse(file.read_text(encoding="utf-8"))
            for node in ast.walk(tree):
                names = (
                    [alias.name for alias in node.names]
                    if isinstance(node, ast.Import)
                    else [node.module or ""]
                    if isinstance(node, ast.ImportFrom)
                    else []
                )
                for name in names:
                    assert not any(
                        name == bad or name.startswith(bad + ".") for bad in forbidden
                    ), f"{file.name} imports {name}"


@pytest.mark.req("FR-RFI-01")
def test_the_model_reserves_construction_rfis_and_nothing_else() -> None:
    from firebid.db.models.clarifications import CLARIFICATION_KINDS

    assert CLARIFICATION_KINDS == ("tender_clarification", "construction_rfi")
