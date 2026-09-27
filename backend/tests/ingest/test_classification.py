"""Proposing a document's type from its digest, by rules (FR-DOC-02).

The digests here come from the sandbox's own extraction, run on real workbooks and documents
built in the test, so the rules are tested against what they will actually be given.
"""

from __future__ import annotations

import io

import docx
import openpyxl
import pytest

from firebid.ingest.classification import RULE_THRESHOLD, Digest, DocType, classify
from firebid.parsing.digest import digest

pytestmark = pytest.mark.req("FR-DOC-02")


def workbook(*rows: list[object], title: str = "BOQ") -> bytes:
    book = openpyxl.Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = title
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def document(*paragraphs: str, table_header: list[str] | None = None) -> bytes:
    made = docx.Document()
    for paragraph in paragraphs:
        made.add_paragraph(paragraph)
    if table_header:
        table = made.add_table(rows=1, cols=len(table_header))
        for cell, text in zip(table.rows[0].cells, table_header, strict=True):
            cell.text = text
    buffer = io.BytesIO()
    made.save(buffer)
    return buffer.getvalue()


def digest_of(payload: bytes, kind: str, filename: str) -> Digest:
    raw = digest(payload, kind)
    return Digest(
        kind=kind,
        filename=filename,
        text=raw["text"],
        sheet_names=tuple(raw["sheet_names"]),
        header_rows=tuple(tuple(row) for row in raw["header_rows"]),
        pages=raw["pages"],
    )


BOQ_ROWS: tuple[list[object], ...] = (
    ["BILL NO. 3 - FIRE PROTECTION SERVICES"],
    [],
    ["ITEM", "DESCRIPTION", "QTY", "UNIT", "RATE", "AMOUNT"],
    ["A", "Supply and install pendent sprinkler, K80, 68 deg C", 412, "no", None, None],
    ["B", "DN50 black steel pipe, schedule 40", 1260, "m", None, None],
)


class TestTheRules:
    def test_a_client_boq_is_known_by_its_columns(self) -> None:
        result = classify(digest_of(workbook(*BOQ_ROWS), "xlsx", "Tender BOQ.xlsx"))

        assert result.doc_type is DocType.BOQ
        assert result.settled
        assert "6 BOQ columns" in result.reasons[0]

    def test_a_valve_schedule_is_a_schedule_not_a_boq(self) -> None:
        payload = workbook(
            ["SUBSIDIARY VALVE SCHEDULE"],
            ["TAG", "SIZE", "LOCATION", "SERVES", "TYPE"],
            ["SV-L05-01", "DN150", "Riser A, L5", "Zone 5A", "Butterfly"],
            title="Valves",
        )

        result = classify(digest_of(payload, "xlsx", "valves.xlsx"))

        assert result.doc_type is DocType.SCHEDULE
        assert result.settled

    def test_a_specification_is_known_by_its_title_and_clauses(self) -> None:
        clauses = [
            f"{section}.{clause} Requirement text" for section in (1, 2) for clause in (1, 2, 3)
        ]
        payload = document("PARTICULAR SPECIFICATION FOR FIRE PROTECTION SERVICES", *clauses)

        result = classify(digest_of(payload, "docx", "spec.docx"))

        assert result.doc_type is DocType.SPECIFICATION
        assert result.settled

    def test_an_addendum_announces_itself(self) -> None:
        payload = document("TENDER ADDENDUM NO. 2", "The following drawings are replaced.")

        result = classify(digest_of(payload, "docx", "letter.docx"))

        assert result.doc_type is DocType.ADDENDUM
        assert result.settled

    def test_responses_to_tender_queries(self) -> None:
        payload = document("RESPONSES TO TENDERERS' QUERIES", "TQ 14: Confirmed.")

        assert classify(digest_of(payload, "docx", "tq.docx")).doc_type is (
            DocType.CLARIFICATION_RESPONSE
        )

    def test_conditions_of_contract(self) -> None:
        payload = document("CONDITIONS OF SUB-CONTRACT", "1. Definitions")

        assert classify(digest_of(payload, "docx", "cos.docx")).doc_type is (
            DocType.CONTRACT_CONDITIONS
        )

    def test_a_pdf_whose_pages_carry_title_blocks_is_a_drawing_set(self) -> None:
        result = classify(Digest(kind="pdf", filename="set.pdf", sheets=40, identified_sheets=38))

        assert result.doc_type is DocType.DRAWING
        assert result.settled

    def test_a_dxf_is_a_drawing(self) -> None:
        assert classify(Digest(kind="dxf", filename="a.dxf")).doc_type is DocType.DRAWING

    def test_something_unrecognised_is_other_and_unsettled(self) -> None:
        result = classify(digest_of(document("Site visit photographs"), "docx", "photos.docx"))

        assert result.doc_type is DocType.OTHER
        assert not result.settled, "the model, then a person, decides what it is"

    def test_two_strong_answers_are_a_doubt_not_a_tie_broken_silently(self) -> None:
        """An addendum that reissues the BOQ is both; someone should say which it is filed as."""
        payload = workbook(["ADDENDUM NO. 3 - REVISED BILL"], *BOQ_ROWS[2:])

        result = classify(digest_of(payload, "xlsx", "Addendum 3 BOQ.xlsx"))

        assert result.confidence < RULE_THRESHOLD
        assert len(result.reasons) == 2


class TestTheDigest:
    def test_a_workbook_gives_its_sheet_names_and_header_row(self) -> None:
        raw = digest(workbook(*BOQ_ROWS), "xlsx")

        assert raw["sheet_names"] == ["BOQ"]
        assert raw["header_rows"] == [["ITEM", "DESCRIPTION", "QTY", "UNIT", "RATE", "AMOUNT"]]

    def test_a_document_gives_its_paragraphs_and_table_headers(self) -> None:
        raw = digest(document("TITLE", table_header=["A", "B"]), "docx")

        assert "TITLE" in raw["text"]
        assert raw["header_rows"] == [["A", "B"]]

    def test_a_kind_with_nothing_to_digest_gives_nothing(self) -> None:
        assert digest(b"", "dxf")["text"] == ""
