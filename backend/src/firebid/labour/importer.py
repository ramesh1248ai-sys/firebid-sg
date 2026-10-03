"""Reading a productivity list workbook: every row checked, every problem by row and column.

A file with any error imports nothing. `read_json` is the sandbox's entry point, since an
uploaded workbook is opened only by the parser workers.
"""

from __future__ import annotations

import io
from decimal import Decimal, InvalidOperation
from typing import Any

from firebid.labour.productivity import SOURCES, Unsourced, check_source
from firebid.pricing.keys import ItemKey, unit_of

COLUMNS: dict[str, tuple[str, ...]] = {
    "type": ("type", "item type", "item"),
    "dn": ("dn", "size", "nominal size", "dn (mm)"),
    "joining": ("joining", "joint", "joining method"),
    "description": ("description",),
    "unit": ("unit", "uom"),
    "hours": ("man-hours per unit", "man-hours", "manhours", "mh per unit", "hours", "mh"),
    "trade": ("trade",),
    "source_type": ("source type", "source"),
    "source_reference": ("source reference", "reference", "source ref"),
}
REQUIRED = ("type", "description", "unit", "hours", "trade", "source_type", "source_reference")
SOURCE_WORDS = {
    "company standard": "company_standard",
    "company_standard": "company_standard",
    "standard": "company_standard",
    "historical project": "historical_project",
    "historical_project": "historical_project",
    "project": "historical_project",
    "estimator judgement": "estimator_judgement",
    "estimator judgment": "estimator_judgement",
    "estimator_judgement": "estimator_judgement",
    "judgement": "estimator_judgement",
}
HEADER_SEARCH_ROWS = 20


def _norm(value: Any) -> str:
    return " ".join(str(value if value is not None else "").split()).lower()


def _header(rows: list[tuple[Any, ...]]) -> tuple[int, dict[str, int]] | None:
    for index, row in enumerate(rows[:HEADER_SEARCH_ROWS]):
        cells = [_norm(cell) for cell in row]
        found: dict[str, int] = {}
        for name, words in COLUMNS.items():
            for position, cell in enumerate(cells):
                if cell in words and position not in found.values():
                    found[name] = position
                    break
        if all(name in found for name in REQUIRED):
            return index, found
    return None


def _problem(row: int, column: str, message: str) -> dict[str, Any]:
    return {"row": row, "column": column, "message": message}


def read_rows(rows: list[tuple[Any, ...]], sheet: str | None = None) -> dict[str, Any]:
    out: dict[str, Any] = {"sheet": sheet, "rows": [], "problems": []}
    found = _header(rows)
    if found is None:
        out["problems"].append(
            _problem(
                0,
                "header",
                "no header row with the columns "
                + ", ".join(COLUMNS[name][0] for name in REQUIRED),
            )
        )
        return out
    header, columns = found
    seen: dict[tuple[str, str, str, str], int] = {}
    for index in range(header + 1, len(rows)):
        cells = rows[index]
        number = index + 1  # as the spreadsheet shows it

        def cell(name: str, cells: tuple[Any, ...] = cells) -> Any:
            position = columns.get(name)
            return cells[position] if position is not None and position < len(cells) else None

        if all(_norm(value) == "" for value in cells):
            continue
        problems = [
            _problem(number, COLUMNS[name][0], "is empty")
            for name in REQUIRED
            if _norm(cell(name)) == ""
        ]
        hours: Decimal | None = None
        if _norm(cell("hours")):
            try:
                hours = Decimal(str(cell("hours")).replace(",", "").strip())
            except InvalidOperation:
                hours = None
            if hours is None or not hours.is_finite() or hours <= 0:
                problems.append(
                    _problem(
                        number,
                        COLUMNS["hours"][0],
                        f"{cell('hours')!r} is not a positive number of man-hours",
                    )
                )
                hours = None
        unit = unit_of(str(cell("unit") or ""))
        if _norm(cell("unit")) and unit is None:
            problems.append(_problem(number, "unit", f"{cell('unit')!r} is not a known unit"))
        source = SOURCE_WORDS.get(_norm(cell("source_type")))
        reference = " ".join(str(cell("source_reference") or "").split())
        if _norm(cell("source_type")) and source not in SOURCES:
            problems.append(
                _problem(
                    number,
                    "source type",
                    f"{cell('source_type')!r} is not company standard, historical project "
                    "or estimator judgement",
                )
            )
        elif source is not None:
            try:
                check_source(source, reference)
            except Unsourced as refusal:
                if not any(p["column"] == COLUMNS["source_reference"][0] for p in problems):
                    problems.append(_problem(number, "source reference", str(refusal)))
        key = ItemKey.of(type=cell("type"), dn=cell("dn"), joining=cell("joining"))
        if problems:
            out["problems"].extend(problems)
            continue
        if hours is None or unit is None or source is None:  # pragma: no cover - reported
            continue
        identity = (key.type, key.dn, key.joining, unit)
        if identity in seen:
            out["problems"].append(
                _problem(number, "type", f"the same item and unit as row {seen[identity]}")
            )
            continue
        seen[identity] = number
        out["rows"].append(
            {
                "row": number,
                "type": key.type,
                "dn": key.dn,
                "joining": key.joining,
                "description": " ".join(str(cell("description")).split()),
                "unit": unit,
                "hours": str(hours.quantize(Decimal("0.0001"))),
                "trade": _norm(cell("trade")).replace(" ", "_"),
                "source_type": source,
                "source_reference": reference,
            }
        )
    if not out["rows"] and not out["problems"]:
        out["problems"].append(
            _problem(header + 1, "rows", "the list has no entries under its header")
        )
    return out


def read_json(payload: bytes) -> dict[str, Any]:
    """Sandbox entry: the first sheet with the expected header, as plain data."""
    from openpyxl import load_workbook

    book = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    try:
        first: dict[str, Any] | None = None
        for sheet in book.worksheets:
            rows = [tuple(row) for row in sheet.iter_rows(values_only=True)]
            found = read_rows(rows, sheet.title)
            if _header(rows) is not None:
                return found
            first = first or found
        return first or {"sheet": None, "rows": [], "problems": [_problem(0, "header", "empty")]}
    finally:
        book.close()
