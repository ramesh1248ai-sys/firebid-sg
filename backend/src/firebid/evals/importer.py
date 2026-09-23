"""Reads a filled `golden_takeoff.xlsx` and turns it into a validated `TenderTruth`.

The whole value of this module is its error messages. The person who filled the workbook is
an estimator, not a developer, and they filled it after hours. A traceback tells them nothing;
"Counts row 14: sheet 'FP-L05-202' is not listed on the Sheets tab" tells them exactly what to
fix and where.

So: every problem is collected before anything is reported, each one names its tab, row and
column, and nothing is written until the whole workbook is clean. A partial import is worse
than none — it looks like it worked.
"""

from __future__ import annotations

from collections.abc import Iterator
from dataclasses import dataclass
from datetime import date, datetime
from enum import StrEnum
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet
from pydantic import ValidationError

from firebid.evals.schema import (
    BoqLineTruth,
    InputClass,
    ObjectCount,
    ObjectType,
    PipeLength,
    RevisionStatus,
    SheetTruth,
    TenderTruth,
)
from firebid.evals.template import BOQ, COUNTS, DUPLICATES, PIPES, SHEETS, TENDER

# Data starts below the header and the hint row.
FIRST_DATA_ROW = 3


@dataclass(frozen=True)
class Problem:
    """One thing wrong, located precisely enough to fix without hunting."""

    tab: str
    row: int | None
    column: str | None
    message: str

    def __str__(self) -> str:
        where = self.tab
        if self.row is not None:
            where += f" row {self.row}"
        if self.column:
            where += f", column '{self.column}'"
        return f"{where}: {self.message}"


class ImportFailed(Exception):
    """The workbook could not be imported. Carries every problem, not just the first."""

    def __init__(self, problems: list[Problem]) -> None:
        super().__init__(f"{len(problems)} problem(s) in the workbook")
        self.problems = problems

    def report(self) -> str:
        """A report an estimator can work through top to bottom."""
        lines = [f"{len(self.problems)} problem(s) to fix:", ""]
        by_tab: dict[str, list[Problem]] = {}
        for problem in self.problems:
            by_tab.setdefault(problem.tab, []).append(problem)
        for tab, problems in by_tab.items():
            lines.append(f"{tab}")
            for problem in sorted(problems, key=lambda p: (p.row or 0, p.column or "")):
                location = f"  row {problem.row}" if problem.row else "  "
                if problem.column:
                    location += f", column '{problem.column}'"
                lines.append(f"{location}: {problem.message}")
            lines.append("")
        return "\n".join(lines).rstrip() + "\n"


class _Reader:
    def __init__(self) -> None:
        self.problems: list[Problem] = []

    def fail(self, tab: str, row: int | None, column: str | None, message: str) -> None:
        self.problems.append(Problem(tab=tab, row=row, column=column, message=message))

    # ---- cell readers, each recording a problem rather than raising ------------------

    def text(
        self, tab: str, row: int, column: str, value: Any, *, required: bool = True
    ) -> str | None:
        if value is None or str(value).strip() == "":
            if required:
                self.fail(tab, row, column, "this is required and is empty")
            return None
        return str(value).strip()

    def whole_number(
        self, tab: str, row: int, column: str, value: Any, *, minimum: int = 0
    ) -> int | None:
        if value is None or str(value).strip() == "":
            self.fail(tab, row, column, "this is required and is empty")
            return None
        try:
            # A spreadsheet stores 12 as 12.0; that is not a fractional answer.
            number = float(str(value).strip().replace(",", ""))
        except ValueError:
            self.fail(tab, row, column, f"'{value}' is not a number")
            return None
        if number != int(number):
            self.fail(tab, row, column, f"'{value}' must be a whole number")
            return None
        if int(number) < minimum:
            self.fail(tab, row, column, f"'{value}' must be {minimum} or more")
            return None
        return int(number)

    def decimal(self, tab: str, row: int, column: str, value: Any) -> float | None:
        if value is None or str(value).strip() == "":
            self.fail(tab, row, column, "this is required and is empty")
            return None
        try:
            return float(str(value).strip().replace(",", ""))
        except ValueError:
            self.fail(tab, row, column, f"'{value}' is not a number")
            return None

    def choice[T: StrEnum](
        self,
        tab: str,
        row: int,
        column: str,
        value: Any,
        options: type[T],
        *,
        required: bool = True,
    ) -> T | None:
        raw = self.text(tab, row, column, value, required=required)
        if raw is None:
            return None
        try:
            return options(raw.lower())
        except ValueError:
            allowed = ", ".join(sorted(str(option) for option in options))
            self.fail(tab, row, column, f"'{raw}' is not one of: {allowed}")
            return None

    def yes_no(self, tab: str, row: int, column: str, value: Any) -> bool:
        raw = self.text(tab, row, column, value, required=False)
        if raw is None:
            return False
        if raw.lower() in ("yes", "y", "true", "1"):
            return True
        if raw.lower() in ("no", "n", "false", "0"):
            return False
        self.fail(tab, row, column, f"'{raw}' should be yes or no")
        return False

    def day(self, tab: str, row: int, column: str, value: Any) -> date | None:
        if value is None or str(value).strip() == "":
            return None
        if isinstance(value, datetime):
            return value.date()
        if isinstance(value, date):
            return value
        try:
            return date.fromisoformat(str(value).strip()[:10])
        except ValueError:
            self.fail(tab, row, column, f"'{value}' is not a date; use YYYY-MM-DD")
            return None


def _rows(sheet: Worksheet) -> Iterator[tuple[int, tuple[Any, ...]]]:
    """Data rows, skipping ones left entirely blank."""
    for index, values in enumerate(
        sheet.iter_rows(min_row=FIRST_DATA_ROW, values_only=True), start=FIRST_DATA_ROW
    ):
        if any(value is not None and str(value).strip() != "" for value in values):
            yield index, values


def import_workbook(path: Path) -> TenderTruth:
    """Read a filled template. Raises `ImportFailed` carrying every problem found."""
    try:
        book = load_workbook(path, read_only=True, data_only=True)
    except Exception as error:  # a corrupt or non-xlsx file
        raise ImportFailed(
            [Problem(tab="(file)", row=None, column=None, message=f"cannot be opened: {error}")]
        ) from error

    reader = _Reader()
    missing = [tab for tab in (TENDER, SHEETS, COUNTS) if tab not in book.sheetnames]
    if missing:
        raise ImportFailed(
            [
                Problem(
                    tab="(workbook)",
                    row=None,
                    column=None,
                    message=f"missing tab(s): {', '.join(missing)}. Is this the right template?",
                )
            ]
        )

    tender_rows = list(_rows(book[TENDER]))
    if not tender_rows:
        reader.fail(TENDER, None, None, "no tender row; fill in one row identifying the tender")
    elif len(tender_rows) > 1:
        reader.fail(
            TENDER, tender_rows[1][0], None, "more than one tender row; one workbook is one tender"
        )

    tender_id = consultant = None
    received_on = None
    tender_class = None
    notes = ""
    if tender_rows:
        row, values = tender_rows[0]
        tender_id = reader.text(TENDER, row, "tender_id", _at(values, 0))
        consultant = reader.text(TENDER, row, "consultant", _at(values, 1))
        received_on = reader.day(TENDER, row, "received_on", _at(values, 2))
        tender_class = reader.choice(TENDER, row, "input_class", _at(values, 3), InputClass)
        notes = reader.text(TENDER, row, "notes", _at(values, 4), required=False) or ""

    sheets = _read_sheets(book[SHEETS], reader, tender_class)
    known = {sheet["sheet_number"] for sheet in sheets}

    _read_counts(book[COUNTS], reader, sheets, known)
    if PIPES in book.sheetnames:
        _read_pipes(book[PIPES], reader, sheets, known)
    if DUPLICATES in book.sheetnames:
        _read_duplicates(book[DUPLICATES], reader, sheets, known)
    boq_lines = _read_boq(book[BOQ], reader) if BOQ in book.sheetnames else []

    book.close()

    if reader.problems:
        raise ImportFailed(reader.problems)

    try:
        return TenderTruth(
            tender_id=tender_id or "",
            consultant=consultant or "",
            received_on=received_on,
            input_class=tender_class or InputClass.MIXED,
            sheets=tuple(SheetTruth(**sheet) for sheet in sheets),
            boq_lines=tuple(boq_lines),
            notes=notes,
        )
    except ValidationError as error:
        # Whole-tender rules, such as "exactly one current revision per sheet number".
        raise ImportFailed(
            [
                Problem(tab="(tender)", row=None, column=None, message=message)
                for message in _messages(error)
            ]
        ) from error


def _messages(error: ValidationError) -> list[str]:
    found: list[str] = []
    for item in error.errors():
        text = str(item.get("msg", "")).removeprefix("Value error, ")
        found.extend(part.strip() for part in text.split(";") if part.strip())
    return found or [str(error)]


def _at(values: tuple[Any, ...], index: int) -> Any:
    return values[index] if index < len(values) else None


def _read_sheets(
    worksheet: Worksheet, reader: _Reader, tender_class: InputClass | None
) -> list[dict[str, Any]]:
    sheets: list[dict[str, Any]] = []
    seen: dict[tuple[str, str], int] = {}
    for row, values in _rows(worksheet):
        number = reader.text(SHEETS, row, "sheet_number", _at(values, 0))
        revision = reader.text(SHEETS, row, "revision", _at(values, 1))
        status = reader.choice(SHEETS, row, "status", _at(values, 2), RevisionStatus)
        sheet_class = reader.choice(
            SHEETS, row, "input_class", _at(values, 3), InputClass, required=False
        )
        not_to_scale = reader.yes_no(SHEETS, row, "not_to_scale", _at(values, 4))

        if number is None or revision is None:
            continue
        key = (number, revision)
        if key in seen:
            reader.fail(
                SHEETS,
                row,
                "sheet_number",
                f"'{number}' revision '{revision}' is already on row {seen[key]}",
            )
            continue
        seen[key] = row
        sheets.append(
            {
                "sheet_number": number,
                "revision": revision,
                "status": status or RevisionStatus.CURRENT,
                "input_class": sheet_class or tender_class or InputClass.MIXED,
                "not_to_scale": not_to_scale,
                "counts": [],
                "pipe_lengths": [],
                "duplicates_of": [],
            }
        )

    if not sheets:
        reader.fail(SHEETS, None, None, "no sheets listed; a tender needs at least one")
    return sheets


def _find(sheets: list[dict[str, Any]], number: str) -> dict[str, Any] | None:
    """The current revision of a sheet number, which is what counts are recorded against."""
    for sheet in sheets:
        if sheet["sheet_number"] == number and sheet["status"] is RevisionStatus.CURRENT:
            return sheet
    for sheet in sheets:
        if sheet["sheet_number"] == number:
            return sheet
    return None


def _read_counts(
    worksheet: Worksheet, reader: _Reader, sheets: list[dict[str, Any]], known: set[str]
) -> None:
    seen: dict[tuple[str, ObjectType], int] = {}
    for row, values in _rows(worksheet):
        number = reader.text(COUNTS, row, "sheet_number", _at(values, 0))
        object_type = reader.choice(COUNTS, row, "object_type", _at(values, 1), ObjectType)
        count = reader.whole_number(COUNTS, row, "count", _at(values, 2))
        if number is None or object_type is None or count is None:
            continue
        if number not in known:
            reader.fail(
                COUNTS, row, "sheet_number", f"'{number}' is not listed on the {SHEETS} tab"
            )
            continue
        key = (number, object_type)
        if key in seen:
            reader.fail(
                COUNTS,
                row,
                "object_type",
                f"'{object_type}' on '{number}' is already counted on row {seen[key]}",
            )
            continue
        seen[key] = row
        sheet = _find(sheets, number)
        if sheet is not None:
            sheet["counts"].append(ObjectCount(object_type=object_type, count=count))


def _read_pipes(
    worksheet: Worksheet, reader: _Reader, sheets: list[dict[str, Any]], known: set[str]
) -> None:
    seen: dict[tuple[str, int], int] = {}
    for row, values in _rows(worksheet):
        number = reader.text(PIPES, row, "sheet_number", _at(values, 0))
        diameter = reader.whole_number(PIPES, row, "nominal_diameter_mm", _at(values, 1), minimum=1)
        length = reader.whole_number(PIPES, row, "length_mm", _at(values, 2))
        if number is None or diameter is None or length is None:
            continue
        if number not in known:
            reader.fail(PIPES, row, "sheet_number", f"'{number}' is not listed on the {SHEETS} tab")
            continue
        sheet = _find(sheets, number)
        if sheet is not None and sheet["not_to_scale"]:
            reader.fail(
                PIPES,
                row,
                "sheet_number",
                f"'{number}' is marked not-to-scale, so its lengths cannot be measured",
            )
            continue
        key = (number, diameter)
        if key in seen:
            reader.fail(
                PIPES,
                row,
                "nominal_diameter_mm",
                f"DN{diameter} on '{number}' is already recorded on row {seen[key]}",
            )
            continue
        seen[key] = row
        if sheet is not None:
            sheet["pipe_lengths"].append(PipeLength(nominal_diameter_mm=diameter, length_mm=length))


def _read_duplicates(
    worksheet: Worksheet, reader: _Reader, sheets: list[dict[str, Any]], known: set[str]
) -> None:
    for row, values in _rows(worksheet):
        number = reader.text(DUPLICATES, row, "sheet_number", _at(values, 0))
        source = reader.text(DUPLICATES, row, "duplicates_of", _at(values, 1))
        if number is None or source is None:
            continue
        for label, value in (("sheet_number", number), ("duplicates_of", source)):
            if value not in known:
                reader.fail(DUPLICATES, row, label, f"'{value}' is not listed on the {SHEETS} tab")
        if number == source:
            reader.fail(DUPLICATES, row, "duplicates_of", "a sheet cannot duplicate itself")
            continue
        if number in known and source in known:
            sheet = _find(sheets, number)
            if sheet is not None:
                sheet["duplicates_of"].append(source)


def _read_boq(worksheet: Worksheet, reader: _Reader) -> list[BoqLineTruth]:
    lines: list[BoqLineTruth] = []
    seen: dict[str, int] = {}
    for row, values in _rows(worksheet):
        reference = reader.text(BOQ, row, "line_reference", _at(values, 0))
        description = reader.text(BOQ, row, "description", _at(values, 1), required=False) or ""
        unit = reader.text(BOQ, row, "unit", _at(values, 2))
        quantity = reader.decimal(BOQ, row, "quantity", _at(values, 3))
        maps_to = reader.choice(BOQ, row, "maps_to", _at(values, 4), ObjectType, required=False)
        if reference is None or unit is None or quantity is None:
            continue
        if reference in seen:
            reader.fail(
                BOQ, row, "line_reference", f"'{reference}' is already on row {seen[reference]}"
            )
            continue
        seen[reference] = row
        lines.append(
            BoqLineTruth(
                line_reference=reference,
                description=description,
                unit=unit,
                quantity=quantity,
                maps_to=maps_to,
            )
        )
    return lines
