"""Builds `golden_takeoff.xlsx`, the workbook estimators fill from a historical tender.

The template is generated rather than committed as a binary, so a change to it shows up as a
reviewable diff in this file instead of an opaque blob.

Everything here is shaped by one fact: the person filling this in is a busy estimator doing it
after the tender is already out. So the sheet does the remembering — dropdown lists instead of
free text for anything with a fixed vocabulary, one row per thing rather than a grid, and an
instructions tab that says what "verified" means and what to do when the honest answer is
"I don't know".
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

from firebid.evals.schema import InputClass, ObjectType, RevisionStatus

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(color="FFFFFF", bold=True)
NOTE_FONT = Font(italic=True, color="595959")

# Tabs, in the order an estimator works through them.
TENDER = "Tender"
SHEETS = "Sheets"
COUNTS = "Counts"
PIPES = "Pipe lengths"
DUPLICATES = "Duplicates"
BOQ = "BOQ lines"
INSTRUCTIONS = "Instructions"
LISTS = "Lists"

COMMON_DIAMETERS = (25, 32, 40, 50, 65, 80, 100, 150, 200, 250, 300)


def _header(sheet: Worksheet, columns: list[tuple[str, int, str]]) -> None:
    """Write a header row, size the columns, and freeze it so it stays visible."""
    for index, (title, width, note) in enumerate(columns, start=1):
        cell = sheet.cell(row=1, column=index, value=title)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        sheet.column_dimensions[get_column_letter(index)].width = width
        if note:
            hint = sheet.cell(row=2, column=index, value=note)
            hint.font = NOTE_FONT
            hint.alignment = Alignment(wrap_text=True, vertical="top")
    sheet.row_dimensions[1].height = 28
    sheet.row_dimensions[2].height = 30
    sheet.freeze_panes = "A3"


def _validate(sheet: Worksheet, column: str, formula: str, message: str) -> None:
    """A dropdown, refusing anything not on the list.

    `showErrorMessage` matters: a warning that can be clicked past is how "DN65 " with a
    trailing space gets into a dataset.
    """
    validation = DataValidation(
        type="list", formula1=formula, allow_blank=True, showErrorMessage=True
    )
    validation.error = message
    validation.errorTitle = "Not a listed value"
    sheet.add_data_validation(validation)
    validation.add(f"{column}3:{column}1000")


def _lists_tab(book: Workbook) -> None:
    """The vocabularies the dropdowns point at. Hidden: it is machinery, not input."""
    sheet: Worksheet = book.create_sheet(LISTS)
    columns = {
        "A": ("object_type", [str(value) for value in ObjectType]),
        "B": ("input_class", [str(value) for value in InputClass]),
        "C": ("status", [str(value) for value in RevisionStatus]),
        "D": ("diameter_mm", [str(value) for value in COMMON_DIAMETERS]),
        "E": ("yes_no", ["yes", "no"]),
    }
    for letter, (title, values) in columns.items():
        sheet[f"{letter}1"] = title
        sheet[f"{letter}1"].font = Font(bold=True)
        for offset, value in enumerate(values, start=2):
            sheet[f"{letter}{offset}"] = value
    sheet.sheet_state = "hidden"


def _named_range(letter: str, count: int) -> str:
    return f"={LISTS}!${letter}$2:${letter}${count + 1}"


def build_workbook() -> Workbook:
    book = Workbook()
    # A new workbook opens with one blank sheet; every tab here is created deliberately.
    default = book.active
    if default is not None:
        book.remove(default)

    instructions = book.create_sheet(INSTRUCTIONS)
    _instructions(instructions)

    tender = book.create_sheet(TENDER)
    _header(
        tender,
        [
            ("tender_id", 22, "Short code, letters and digits, e.g. MC-2024-014"),
            ("consultant", 26, "The consultant who issued the drawings"),
            ("received_on", 16, "YYYY-MM-DD, when the tender arrived"),
            ("input_class", 18, "How the set arrived overall"),
            ("notes", 46, "Anything unusual about this tender"),
        ],
    )

    sheets = book.create_sheet(SHEETS)
    _header(
        sheets,
        [
            ("sheet_number", 22, "Exactly as printed, e.g. FP-L05-201"),
            ("revision", 12, "As printed, e.g. R04"),
            ("status", 14, "current or superseded"),
            ("input_class", 16, "This sheet specifically"),
            ("not_to_scale", 14, "yes if the sheet is NTS; lengths are then not measured"),
        ],
    )

    counts = book.create_sheet(COUNTS)
    _header(
        counts,
        [
            ("sheet_number", 22, "Must appear on the Sheets tab"),
            ("object_type", 28, "Pick from the list"),
            ("count", 10, "How many you verified. 0 is a real answer"),
        ],
    )

    pipes = book.create_sheet(PIPES)
    _header(
        pipes,
        [
            ("sheet_number", 22, "Must appear on the Sheets tab"),
            ("nominal_diameter_mm", 20, "DN in millimetres, e.g. 150"),
            ("length_mm", 16, "Centreline length in millimetres, a whole number"),
        ],
    )

    duplicates = book.create_sheet(DUPLICATES)
    _header(
        duplicates,
        [
            ("sheet_number", 22, "The sheet that repeats content"),
            ("duplicates_of", 22, "The sheet it repeats, e.g. an enlarged plan of a GA"),
        ],
    )

    boq = book.create_sheet(BOQ)
    _header(
        boq,
        [
            ("line_reference", 16, "As in the client BOQ, e.g. 2.4.1"),
            ("description", 48, "As written in the BOQ"),
            ("unit", 10, "e.g. nr, m, item"),
            ("quantity", 14, "As you priced it"),
            ("maps_to", 28, "The object type, or leave blank for prelims and PC sums"),
        ],
    )

    _lists_tab(book)
    object_types = len(list(ObjectType))
    input_classes = len(list(InputClass))
    statuses = len(list(RevisionStatus))
    diameters = len(COMMON_DIAMETERS)

    _validate(tender, "D", _named_range("B", input_classes), "Pick an input class from the list.")
    _validate(sheets, "C", _named_range("C", statuses), "Status is current or superseded.")
    _validate(sheets, "D", _named_range("B", input_classes), "Pick an input class from the list.")
    _validate(sheets, "E", _named_range("E", 2), "Answer yes or no.")
    _validate(counts, "B", _named_range("A", object_types), "Pick an object type from the list.")
    _validate(pipes, "B", _named_range("D", diameters), "Pick a nominal diameter, in millimetres.")
    _validate(boq, "E", _named_range("A", object_types), "Pick an object type, or leave blank.")

    return book


def _instructions(sheet: Worksheet) -> None:
    sheet.column_dimensions["A"].width = 110
    lines: list[tuple[str, bool]] = [
        ("FireBid SG — golden takeoff", True),
        ("", False),
        (
            "This workbook records what a historical tender actually contained, as verified by "
            "an estimator. It is the yardstick the platform's accuracy is measured against, so "
            "it is worth more than a fast answer.",
            False,
        ),
        ("", False),
        ("What 'verified' means", True),
        (
            "A number you would defend in a handover. If you counted it, or checked a count "
            "someone else made, it is verified. If you are estimating from memory, it is not — "
            "leave the row out and note it on the Tender tab.",
            False,
        ),
        ("", False),
        ("When you are not sure", True),
        (
            "Leave the cell empty rather than guessing. An empty cell is read as 'not recorded' "
            "and is excluded from the score. A wrong number is read as truth and makes the "
            "platform look accurate when it is not — which is the one outcome this whole "
            "exercise exists to prevent.",
            False,
        ),
        ("", False),
        ("Order of work", True),
        ("1. Tender tab: one row, identifying the tender.", False),
        (
            "2. Sheets tab: one row per sheet AND revision. Include superseded revisions — how "
            "the platform handles them is itself being measured. Exactly one revision per "
            "sheet number is 'current'.",
            False,
        ),
        ("3. Counts tab: one row per sheet and object type. Zero is a real answer.", False),
        (
            "4. Pipe lengths tab: centreline length per nominal diameter, per sheet, in "
            "millimetres. Skip sheets marked not-to-scale.",
            False,
        ),
        (
            "5. Duplicates tab: where a sheet repeats content from another, such as an enlarged "
            "plan of part of a general arrangement. These matter: double-counting them is a "
            "real error the platform must avoid.",
            False,
        ),
        ("6. BOQ lines tab: the client's BOQ as you priced it.", False),
        ("", False),
        ("Sending it back", True),
        (
            "Send the workbook and the drawing set to the nominated data owner, not by email to "
            "an individual. The drawings stay in the eval bucket and never enter the code "
            "repository. Client-confidential material is handled under the tender's own "
            "confidentiality terms; if a tender's terms forbid retaining the documents, say so "
            "on the Tender tab and send the workbook alone.",
            False,
        ),
        ("", False),
        ("Checking it", True),
        (
            "Run: firebid-eval import <workbook.xlsx>. It reports problems by tab, row and "
            "column, and writes nothing until every problem is fixed.",
            False,
        ),
    ]
    for index, (text, bold) in enumerate(lines, start=1):
        cell = sheet.cell(row=index, column=1, value=text)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        if bold:
            cell.font = Font(bold=True, size=12 if index == 1 else 11)
        if not bold and text:
            sheet.row_dimensions[index].height = 30


def write_template(path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    build_workbook().save(path)
    return path
