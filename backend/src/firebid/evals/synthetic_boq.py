"""A client's bill of quantities for the synthetic installation (P1-09).

Laid out as a Singapore quantity surveyor lays one out:

* a **summary** sheet collecting each bill's total, with a chart of them;
* **Bill No. 1 - Fire Sprinkler Installation**: a title block with the client's logo,
  merged headings, lettered sections, numbered items, units and quantities, an empty rate
  column for the tenderer, amounts as `=Qty*Rate` formulas (and a few typed in as plain
  values, as some consultants do), section sub-totals and a bill total;
* a **provisional sums** section of lump sums;
* a **hidden** sheet of lists behind the unit column's data validation;
* comments on rate cells, defined names, print titles and a print area.

Everything a naive re-save loses is here, so the priced-export round trip can prove it keeps
the client's file intact.

Its quantities are the client's view of the installation in `synthetic_network`, so they
differ from what takeoff measures in known ways: DN50 branch is 70 m (measured 72 m), DN150
main 8.0 m (8.05 m), a flow switch nobody drew, and no tees at all.
"""

from __future__ import annotations

import io
from dataclasses import dataclass
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
from openpyxl.drawing.image import Image as XlImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

BILL = "Bill 1 - Sprinklers"
SUMMARY = "Summary"
LISTS = "Lists"
HEADER_ROW = 7


@dataclass(frozen=True)
class Line:
    item: str
    description: str
    unit: str
    quantity: Decimal | None
    maps_to: str | None  # the synthetic takeoff's item it is, or None (a PS, a lump sum)
    amount: Decimal | None = None  # a typed amount (a lump sum), not a formula
    formula: bool = True  # the amount is =Qty*Rate


SECTIONS: list[tuple[str, str, list[Line]]] = [
    (
        "A",
        "SPRINKLER HEADS",
        [
            Line(
                "A1",
                "Pendent sprinkler head, quick response, K80, 68°C, chrome",
                "nr",
                Decimal(16),
                "Sprinkler, pendent",
            ),
            Line(
                "A2",
                "Upright sprinkler head, quick response, K80, 68°C",
                "nr",
                Decimal(4),
                "Sprinkler, upright",
            ),
            Line(
                "A3",
                "Horizontal sidewall sprinkler head, K80, 68°C, white",
                "nr",
                Decimal(4),
                "Sprinkler, sidewall",
            ),
        ],
    ),
    (
        "B",
        "PIPEWORK (black steel to BS EN 10255 medium grade, including hangers)",
        [
            Line(
                "B1",
                "150 mm diameter pipe, grooved joints",
                "m",
                Decimal("8.0"),
                "Pipe, DN150, main",
            ),
            Line(
                "B2",
                "100 mm diameter pipe, grooved joints",
                "m",
                Decimal("8.25"),
                "Pipe, DN100, main",
            ),
            Line(
                "B3", "50 mm diameter pipe, screwed joints", "m", Decimal(70), "Pipe, DN50, branch"
            ),
            Line(
                "B4",
                "25 mm diameter sprinkler drop, screwed joints",
                "m",
                Decimal(12),
                "Sprinkler drop, DN25 (vertical, not drawn)",
            ),
            Line(
                "B5",
                "150 mm diameter riser",
                "m",
                Decimal(4),
                "Riser, DN150 (vertical, not drawn)",
                formula=False,
                amount=Decimal(0),
            ),
        ],
    ),
    (
        "C",
        "VALVES AND ANCILLARIES",
        [
            Line(
                "C1",
                "150 mm gate valve, flanged, with monitored supervisory switch",
                "nr",
                Decimal(1),
                "Gate valve, DN150",
            ),
            Line("C2", "150 mm check valve, grooved", "nr", Decimal(1), "Check valve, DN150"),
            Line("C3", "Flow switch, 150 mm, with retard", "nr", Decimal(1), None),
            Line(
                "C4",
                "150 x 100 mm concentric reducer",
                "nr",
                Decimal(1),
                "Fitting, DN150xDN100 (fitting reducer)",
            ),
        ],
    ),
    (
        "D",
        "PROVISIONAL SUMS",
        [
            Line(
                "D1",
                "Allow for testing and commissioning",
                "sum",
                None,
                None,
                amount=Decimal(5000),
                formula=False,
            ),
            Line(
                "D2",
                "Allow for liaison with SCDF and attendance at inspections",
                "sum",
                None,
                None,
                amount=Decimal(3000),
                formula=False,
            ),
        ],
    ),
]


@dataclass
class ClientBoqFixture:
    payload: bytes
    # "Bill 1 - Sprinklers!12" -> what the line is (the truth for mapping accuracy)
    truth: dict[str, str | None]
    # item -> (row, rate cell, amount cell)
    cells: dict[str, tuple[int, str, str]]


def _logo() -> bytes:
    from PIL import Image, ImageDraw

    image = Image.new("RGB", (160, 48), "white")
    draw = ImageDraw.Draw(image)
    draw.rectangle((2, 2, 157, 45), outline=(180, 20, 20), width=3)
    draw.text((14, 16), "CLIENT QS", fill=(180, 20, 20))
    out = io.BytesIO()
    image.save(out, format="PNG")
    return out.getvalue()


def client_boq() -> ClientBoqFixture:
    book = Workbook()
    summary = book.create_sheet(SUMMARY, 0)
    book.remove(book["Sheet"])  # the blank one every new workbook starts with
    bill = book.create_sheet(BILL)
    lists = book.create_sheet(LISTS)
    lists.sheet_state = "hidden"
    for index, unit in enumerate(("nr", "m", "sum", "set", "lot"), start=1):
        lists.cell(row=index, column=1, value=unit)

    thin = Side(style="thin")
    boxed = Border(left=thin, right=thin, top=thin, bottom=thin)
    bold = Font(bold=True)
    shaded = PatternFill("solid", fgColor="DDDDDD")

    # Title block.
    bill.merge_cells("A1:F1")
    bill["A1"] = "PROPOSED COMMERCIAL DEVELOPMENT AT MARINA BAY"
    bill["A1"].font = Font(bold=True, size=14)
    bill["A1"].alignment = Alignment(horizontal="center")
    bill.merge_cells("A2:F2")
    bill["A2"] = "BILL NO. 1 - FIRE SPRINKLER INSTALLATION"
    bill["A2"].alignment = Alignment(horizontal="center")
    bill.merge_cells("A4:F4")
    bill["A4"] = "Quantities are provisional and subject to remeasurement."
    logo = XlImage(io.BytesIO(_logo()))
    bill.add_image(logo, "H1")

    headings = ("Item", "Description", "Unit", "Qty", "Rate", "Amount")
    for column, heading in enumerate(headings, start=1):
        cell = bill.cell(row=HEADER_ROW, column=column, value=heading)
        cell.font = bold
        cell.fill = shaded
        cell.border = boxed
    bill.column_dimensions["A"].width = 7
    bill.column_dimensions["B"].width = 60
    bill.column_dimensions["E"].width = 12
    bill.column_dimensions["F"].width = 14

    units = DataValidation(type="list", formula1=f"={LISTS}!$A$1:$A$5", allow_blank=True)
    bill.add_data_validation(units)

    truth: dict[str, str | None] = {}
    cells: dict[str, tuple[int, str, str]] = {}
    subtotals: list[str] = []
    row = HEADER_ROW + 2
    for letter, title, lines in SECTIONS:
        bill.cell(row=row, column=1, value=letter).font = bold
        bill.merge_cells(start_row=row, start_column=2, end_row=row, end_column=6)
        bill.cell(row=row, column=2, value=title).font = bold
        row += 1
        first = row
        for line in lines:
            bill.cell(row=row, column=1, value=line.item)
            bill.cell(row=row, column=2, value=line.description).alignment = Alignment(
                wrap_text=True
            )
            bill.cell(row=row, column=3, value=line.unit)
            units.add(f"C{row}")
            if line.quantity is not None:
                bill.cell(row=row, column=4, value=float(line.quantity)).number_format = "0.00"
            rate = bill.cell(row=row, column=5)
            rate.number_format = "#,##0.00"
            rate.border = boxed
            amount = bill.cell(row=row, column=6)
            amount.number_format = "#,##0.00"
            if line.formula and line.quantity is not None:
                amount.value = f"=D{row}*E{row}"
            elif line.amount is not None:
                amount.value = float(line.amount)
            truth[f"{BILL}!{row}"] = line.maps_to
            cells[line.item] = (row, f"E{row}", f"F{row}")
            row += 1
        bill.cell(row=row, column=2, value=f"Total carried to collection ({letter})").font = bold
        bill.cell(row=row, column=6, value=f"=SUM(F{first}:F{row - 1})").font = bold
        subtotals.append(f"F{row}")
        row += 2
    bill.cell(row=row, column=2, value="TOTAL OF BILL NO. 1").font = bold
    bill.cell(row=row, column=6, value="=" + "+".join(subtotals)).font = bold
    total_row = row

    bill["E10"].comment = Comment("Rate to include all fittings and accessories.", "QS")
    bill[cells["B1"][1]].comment = Comment("Rate per metre, measured net.", "QS")

    bill.print_title_rows = f"{HEADER_ROW}:{HEADER_ROW}"
    bill.print_area = f"A1:F{total_row}"
    bill.page_setup.orientation = "portrait"
    bill.page_setup.fitToWidth = 1

    book.defined_names["BillOneTotal"] = DefinedName(
        "BillOneTotal", attr_text=f"'{BILL}'!$F${total_row}"
    )

    # The summary: each section's sub-total, and a chart of them.
    summary["A1"] = "SUMMARY OF TENDER"
    summary["A1"].font = Font(bold=True, size=14)
    summary["A3"], summary["B3"] = "Section", "Amount (SGD)"
    for index, ((letter, title, _), subtotal) in enumerate(
        zip(SECTIONS, subtotals, strict=True), start=4
    ):
        summary.cell(row=index, column=1, value=f"{letter}  {title[:40]}")
        summary.cell(row=index, column=2, value=f"='{BILL}'!{subtotal}")
    last = 3 + len(SECTIONS)
    summary.cell(row=last + 1, column=1, value="TOTAL").font = bold
    summary.cell(row=last + 1, column=2, value="=BillOneTotal").font = bold
    chart = BarChart()
    chart.title = "Tender by section"
    chart.add_data(Reference(summary, min_col=2, min_row=3, max_row=last), titles_from_data=True)
    chart.set_categories(Reference(summary, min_col=1, min_row=4, max_row=last))
    summary.add_chart(chart, "D3")

    out = io.BytesIO()
    book.save(out)
    return ClientBoqFixture(out.getvalue(), truth, cells)
