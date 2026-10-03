"""Synthetic supplier quotations, with their known answers (P2-04).

One quotation in each form a supplier sends one, each with every field the platform
captures, so an extraction can be checked field by field:

* **`pdf_quote`:** a USD quotation from an overseas valve supplier, FOB, as a PDF with a
  table of lines and a list of exclusions.
* **`xlsx_quote`:** an SGD quotation from a local pipe supplier, as a workbook, valid for
  30 days from its date.
* **`email_quote`:** an email from a sprinkler supplier whose body carries the terms and
  whose attached workbook carries the lines.

`ERP_EXPORT` is the company ERP's export as the file adapter reads it: an item master,
purchase orders and historical project costs.
"""

from __future__ import annotations

import io
from dataclasses import dataclass, field
from datetime import date
from decimal import Decimal
from email.message import EmailMessage

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@dataclass(frozen=True)
class ExpectedLine:
    description: str
    brand: str
    model: str
    unit: str
    unit_price: Decimal
    moq: str
    lead_time: str


@dataclass(frozen=True)
class Expected:
    supplier: str
    quote_number: str
    quote_date: date
    valid_until: date
    currency: str
    delivery_terms: str
    incoterm: str | None
    lines: tuple[ExpectedLine, ...]
    exclusions: tuple[str, ...] = field(default_factory=tuple)


PDF = Expected(
    supplier="Pacific Valve Co. Ltd",
    quote_number="PV-Q-2026-1042",
    quote_date=date(2026, 9, 15),
    valid_until=date(2026, 12, 14),
    currency="USD",
    delivery_terms="FOB Shanghai",
    incoterm="FOB",
    lines=(
        ExpectedLine(
            "Gate valve OS&Y flanged DN150",
            "Pacific",
            "GV-150F",
            "no",
            Decimal("100.00"),
            "10",
            "8 weeks",
        ),
        ExpectedLine(
            "Check valve flanged DN150",
            "Pacific",
            "CV-150F",
            "no",
            Decimal("86.50"),
            "10",
            "8 weeks",
        ),
        ExpectedLine(
            "Butterfly valve grooved DN100",
            "Pacific",
            "BV-100G",
            "no",
            Decimal("42.75"),
            "20",
            "6 weeks",
        ),
    ),
    exclusions=("Local delivery and unloading", "Import duties and GST"),
)

WORKBOOK = Expected(
    supplier="Lion City Steel Pte Ltd",
    quote_number="LCS/26/0877",
    quote_date=date(2026, 9, 20),
    valid_until=date(2026, 10, 20),
    currency="SGD",
    delivery_terms="Delivered to site, Singapore",
    incoterm=None,
    lines=(
        ExpectedLine(
            "Black steel pipe BS EN 10255 Heavy DN50",
            "LCS",
            "BSP-50H",
            "m",
            Decimal("18.40"),
            "60",
            "2 weeks",
        ),
        ExpectedLine(
            "Black steel pipe ASTM A53 Sch 40 DN100",
            "LCS",
            "BSP-100S40",
            "m",
            Decimal("41.20"),
            "60",
            "2 weeks",
        ),
    ),
)

EMAIL = Expected(
    supplier="Apex Fire Products Pte Ltd",
    quote_number="AFP-7731",
    quote_date=date(2026, 9, 25),
    valid_until=date(2026, 11, 24),
    currency="SGD",
    delivery_terms="DDP Singapore",
    incoterm="DDP",
    lines=(
        ExpectedLine(
            "Pendent sprinkler K80 68C quick response chrome",
            "Apex",
            "SP-P80Q",
            "no",
            Decimal("9.80"),
            "500",
            "3 weeks",
        ),
        ExpectedLine(
            "Sidewall sprinkler K80 68C quick response white",
            "Apex",
            "SP-S80Q",
            "no",
            Decimal("12.40"),
            "100",
            "3 weeks",
        ),
    ),
    exclusions=("Installation",),
)

COLUMNS = ("Description", "Brand", "Model", "Unit", "Unit price", "MOQ", "Lead time")


def _row(line: ExpectedLine) -> tuple[str, ...]:
    return (
        line.description,
        line.brand,
        line.model,
        line.unit,
        str(line.unit_price),
        line.moq,
        line.lead_time,
    )


def pdf_quote(expected: Expected = PDF) -> bytes:
    """The quotation as a one-page PDF: labelled terms, a table, and exclusions."""
    import matplotlib

    matplotlib.use("Agg")
    from matplotlib import pyplot as plt
    from matplotlib.backends.backend_pdf import PdfPages

    figure = plt.figure(figsize=(11.69, 8.27))
    y = 0.92
    figure.text(0.06, y, "QUOTATION", fontsize=16, weight="bold")
    terms = (
        ("Supplier:", expected.supplier),
        ("Quotation No:", expected.quote_number),
        ("Date:", expected.quote_date.strftime("%d %B %Y")),
        ("Valid until:", expected.valid_until.strftime("%d %B %Y")),
        ("Currency:", expected.currency),
        ("Delivery terms:", expected.delivery_terms),
    )
    for label, value in terms:
        y -= 0.04
        figure.text(0.06, y, label, fontsize=10)
        figure.text(0.26, y, value, fontsize=10)
    xs = (0.06, 0.40, 0.50, 0.62, 0.69, 0.80, 0.87)
    y -= 0.07
    for x, heading in zip(xs, COLUMNS, strict=True):
        figure.text(x, y, heading, fontsize=10, weight="bold")
    for line in expected.lines:
        y -= 0.04
        for x, cell in zip(xs, _row(line), strict=True):
            figure.text(x, y, cell, fontsize=10)
    y -= 0.06
    figure.text(0.06, y, "Exclusions:", fontsize=10)
    for exclusion in expected.exclusions:
        y -= 0.035
        figure.text(0.08, y, f"- {exclusion}", fontsize=10)
    buffer = io.BytesIO()
    with PdfPages(buffer) as pages:
        pages.savefig(figure)
    plt.close(figure)
    return buffer.getvalue()


def lines_workbook(
    expected: Expected, *, with_terms: bool = True, validity: str | None = None
) -> bytes:
    """A quotation workbook: the terms as labelled rows (unless left to the email), then
    the table of lines. `validity` writes the validity in the supplier's own words."""
    from openpyxl import Workbook

    book = Workbook()
    sheet = book.active
    assert sheet is not None  # noqa: S101 - a new workbook has one
    sheet.title = "Quotation"
    if with_terms:
        sheet.append(("Supplier", expected.supplier))
        sheet.append(("Quotation No", expected.quote_number))
        sheet.append(("Date", expected.quote_date))
        sheet.append(("Validity", validity or expected.valid_until))
        sheet.append(("Currency", expected.currency))
        sheet.append(("Delivery terms", expected.delivery_terms))
        sheet.append(())
    sheet.append(COLUMNS)
    for line in expected.lines:
        sheet.append((*_row(line)[:4], float(line.unit_price), *_row(line)[5:]))
    sheet.append(("Total", "", "", "", "as per order"))
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def xlsx_quote(expected: Expected = WORKBOOK) -> bytes:
    """The local supplier's workbook, whose validity is "30 days" from its date."""
    return lines_workbook(expected, validity="30 days")


def email_quote(expected: Expected = EMAIL) -> bytes:
    """An email (.eml): the terms in its body, the lines in an attached workbook."""
    message = EmailMessage()
    message["From"] = f'"{expected.supplier}" <sales@apexfire.example>'
    message["To"] = "estimating@contractor.example"
    message["Subject"] = f"Quotation {expected.quote_number}"
    message["Date"] = "Fri, 25 Sep 2026 10:15:00 +0800"
    message.set_content(
        "Dear Sir,\n\n"
        "Please find our quotation attached.\n\n"
        f"Quotation No: {expected.quote_number}\n"
        f"Date: {expected.quote_date.strftime('%d/%m/%Y')}\n"
        f"Valid until: {expected.valid_until.strftime('%d/%m/%Y')}\n"
        f"Currency: {expected.currency}\n"
        f"Delivery terms: {expected.delivery_terms}\n\n"
        "Exclusions:\n" + "".join(f"- {item}\n" for item in expected.exclusions) + "\n"
        "Regards,\nSales\n"
    )
    main, _, sub = XLSX.partition("/")
    message.add_attachment(
        lines_workbook(expected, with_terms=False),
        maintype=main,
        subtype=sub,
        filename=f"{expected.quote_number}.xlsx",
    )
    return message.as_bytes()


# --- The ERP's export (FR-CST-08) -------------------------------------------------------------

GATE_150 = "gate_valve|150||||"
PIPE_50 = "pipe|50|black_steel|||"

ERP_ITEMS = (
    ("VLV-GV-150", "Gate valve OS&Y DN150", "no", GATE_150),
    ("PIP-BS-050", "Black steel pipe DN50", "m", PIPE_50),
)
ERP_PURCHASE_ORDERS = (
    (
        "PO-25-0412",
        date(2025, 4, 12),
        GATE_150,
        "Gate valve OS&Y DN150",
        "no",
        "118.00",
        "Pacific Valve",
    ),
    (
        "PO-25-0903",
        date(2025, 9, 3),
        GATE_150,
        "Gate valve OS&Y DN150",
        "no",
        "122.50",
        "Pacific Valve",
    ),
    (
        "PO-26-0118",
        date(2026, 1, 18),
        PIPE_50,
        "Black steel pipe DN50",
        "m",
        "17.90",
        "Lion City Steel",
    ),
)
ERP_HISTORICAL = (
    ("Tampines Hub", date(2025, 11, 30), GATE_150, "no", "120.00"),
    ("Jurong Tower", date(2026, 3, 31), PIPE_50, "m", "18.10"),
)


def erp_export(*, bad_row: bool = False) -> bytes:
    """The ERP's export workbook. `bad_row` adds a purchase order with no price."""
    from openpyxl import Workbook

    book = Workbook()
    items = book.active
    assert items is not None  # noqa: S101 - a new workbook has one
    items.title = "Items"
    items.append(("code", "description", "unit", "item_key"))
    for item in ERP_ITEMS:
        items.append(item)
    orders = book.create_sheet("PurchaseOrders")
    orders.append(
        ("po_number", "date", "item_key", "description", "unit", "unit_price_sgd", "supplier")
    )
    for order in ERP_PURCHASE_ORDERS:
        orders.append(order)
    if bad_row:
        orders.append(("PO-26-0999", date(2026, 2, 1), GATE_150, "Gate valve", "no", "tbc", ""))
    costs = book.create_sheet("HistoricalCosts")
    costs.append(("project", "date", "item_key", "unit", "unit_price_sgd"))
    for cost in ERP_HISTORICAL:
        costs.append(cost)
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()
