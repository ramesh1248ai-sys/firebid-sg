"""A synthetic fire protection specification, with its known answers (P1-06).

Written the way a Singapore M&E consultant writes a particular specification: numbered
sections and clauses, the fire protection systems among other services, pipe material and
joining by size range, sprinkler heads by type, approved makes. Its answers are exact, so
every attribute the extractor reads can be checked.

One clause contradicts the drawings on purpose: the car park sprinkler pipework is to be
galvanised, where the drawing notes say black steel. P1-06 records it as a conditional
attribute; P2-03's drawing-against-specification check uses it.
"""

from __future__ import annotations

import io
from dataclasses import dataclass

# (number, heading, body) in reading order. A number with no body is a heading.
CLAUSES: tuple[tuple[str, str, str], ...] = (
    ("1", "GENERAL", ""),
    ("1.1", "Scope", "This section covers the fire protection services to the development."),
    ("1.2", "Standards", "The installation shall comply with SS CP 52 and SS EN 12845."),
    ("2", "AUTOMATIC SPRINKLER SYSTEM", ""),
    ("2.1", "Pipework", ""),
    (
        "2.1.1",
        "",
        "Sprinkler pipework up to and including DN 50 shall be black steel to BS EN 10255 "
        "Heavy grade. Pipework DN 65 and above shall be black steel to ASTM A53 Grade B, "
        "Schedule 40.",
    ),
    (
        "2.1.2",
        "",
        "Pipes up to and including 50 mm shall be joined by screwed fittings. Pipes 65 mm "
        "and above shall be joined by roll-grooved mechanical couplings.",
    ),
    (
        "2.1.3",
        "",
        "All sprinkler pipework within the basement car park shall be hot-dip galvanised.",
    ),
    ("2.2", "Sprinklers", ""),
    (
        "2.2.1",
        "",
        "Sprinkler heads shall be quick response pendent type with a K-factor of K80, rated "
        "at 68°C, chrome finish, unless otherwise indicated.",
    ),
    (
        "2.2.2",
        "",
        "Sidewall sprinklers where shown shall be quick response, K80, 68°C, white finish.",
    ),
    ("2.2.3", "", "Approved makes: Tyco, Viking or Reliable."),
    ("3", "HOSE REEL SYSTEM", ""),
    (
        "3.1",
        "Pipework",
        "Hose reel pipework shall be galvanised steel to BS EN 10255 Medium grade, joined "
        "by screwed fittings throughout.",
    ),
    ("4", "FIRE HYDRANT SYSTEM", ""),
    (
        "4.1",
        "Pipework",
        "Hydrant mains DN 100 and above shall be ductile iron to BS EN 545, with flanged joints.",
    ),
    ("5", "LOW VOLTAGE ELECTRICAL INSTALLATION", ""),
    (
        "5.1",
        "Cables",
        "Cables shall be copper conductors with XLPE insulation, installed in galvanised trunking.",
    ),
)


# Supports (P2-01), as a clause of the sprinkler section a specification may or may not
# have: hanger spacing by size, and, on a project that calls for it, seismic restraint.
SUPPORT_CLAUSES: tuple[tuple[str, str, str], ...] = (
    ("2.3", "Supports", ""),
    (
        "2.3.1",
        "",
        "Hangers for pipes up to and including DN 50 shall be spaced at not more than 3.0 m "
        "centres. Hangers for pipes DN 65 and above shall be fixed at intervals not exceeding "
        "4000 mm.",
    ),
)
SEISMIC_CLAUSES: tuple[tuple[str, str, str], ...] = (
    (
        "2.3.2",
        "",
        "Seismic bracing shall be provided to all sprinkler pipework DN 65 and above.",
    ),
)


def with_supports(seismic: bool = False) -> tuple[tuple[str, str, str], ...]:
    """The specification with its supports clauses, placed at the end of section 2."""
    at = next(i for i, clause in enumerate(CLAUSES) if clause[0] == "3")
    extra = SUPPORT_CLAUSES + (SEISMIC_CLAUSES if seismic else ())
    return (*CLAUSES[:at], *extra, *CLAUSES[at:])


# The rest of a specification (P2-03): what the contractor must do beyond install, and what
# is and is not in the contract. Placed after section 5, under headings the section rules
# read as fire protection.
OBLIGATION_CLAUSES: tuple[tuple[str, str, str], ...] = (
    (
        "2.4",
        "Valves",
        "Gate valves and non-return valves shall be provided as shown on the drawings.",
    ),
    ("6", "FIRE PROTECTION SERVICES - TESTING, COMMISSIONING AND HANDOVER", ""),
    (
        "6.1",
        "Testing",
        "All pipework shall be hydrostatically tested at 14 bar for 2 hours. Pipework shall "
        "be flushed before the sprinklers are fitted.",
    ),
    (
        "6.2",
        "Painting and identification",
        "All exposed pipework shall be painted with two coats of signal red finish. Pipework "
        "shall be identified with colour bands at every floor.",
    ),
    (
        "6.3",
        "Commissioning",
        "The installation shall be commissioned in the presence of the Engineer.",
    ),
    (
        "6.4",
        "Warranty and maintenance",
        "The Contractor shall warrant the installation for 12 months from practical "
        "completion. The defects liability period shall be 12 months. The Contractor shall "
        "carry out maintenance at monthly intervals for 12 months.",
    ),
    (
        "6.5",
        "Spares and training",
        "The Contractor shall hand over spare sprinklers amounting to 2 sets of each type. "
        "The Contractor shall provide training to the Employer's staff for 2 days.",
    ),
    (
        "6.6",
        "Submittals",
        "The Contractor shall submit shop drawings and hydraulic calculations within 4 weeks "
        "of award.",
    ),
    (
        "6.7",
        "Authority",
        "The Contractor shall attend all SCDF inspections and assist the Qualified Person in "
        "obtaining the Fire Safety Certificate.",
    ),
    ("6.8", "Makes", "Only makes on the Employer's AVL shall be used."),
    ("7", "FIRE PROTECTION SERVICES - SCOPE AND INTERFACES", ""),
    (
        "7.1",
        "",
        "Power supply to the fire pump control panels shall be provided by the electrical "
        "contractor.",
    ),
    ("7.2", "", "The Contractor shall provide the water supply connection from the PUB main."),
    ("7.3", "", "Builder's works, openings and plinths are excluded from this contract."),
    ("7.4", "", "Ceiling access panels are shown on the architect's drawings."),
    (
        "7.5",
        "",
        "Cabling between flow switches and the fire alarm panel shall be by the fire alarm "
        "contractor.",
    ),
    ("7.6", "", "A flow test header shall be provided in the pump room."),
    ("7.7", "", "The Contractor shall provide fire stopping to pipe penetrations where required."),
    ("7.8", "", "Pressure gauges shall be of the make scheduled or equal."),
)


def extended() -> tuple[tuple[str, str, str], ...]:
    """The specification with its obligations and interfaces: clause 2.4 in section 2, the
    two new sections at the end."""
    at = next(i for i, clause in enumerate(CLAUSES) if clause[0] == "3")
    return (*CLAUSES[:at], OBLIGATION_CLAUSES[0], *CLAUSES[at:], *OBLIGATION_CLAUSES[1:])


# P2-07: what the specification puts on the contractor by way of design, and the conditions
# the work is done in. Appended to the extended specification for the risk fixture only.
RISK_CLAUSES: tuple[tuple[str, str, str], ...] = (
    ("8", "FIRE PROTECTION SERVICES - DESIGN AND SITE CONDITIONS", ""),
    (
        "8.1",
        "Design",
        "The sprinkler installation shall be procured on a design and build basis. The "
        "Contractor shall be responsible for the detailed design of the installation.",
    ),
    (
        "8.2",
        "Calculations",
        "Hydraulic calculations shall be prepared by the Contractor for every design area.",
    ),
    (
        "8.3",
        "Qualified Person",
        "The Contractor shall engage a Qualified Person to endorse the design and the "
        "submissions to the authorities.",
    ),
    (
        "8.4",
        "Working hours",
        "Works in the retail podium shall be carried out at night between 2200 and 0600 hours.",
    ),
    (
        "8.5",
        "Occupation",
        "The existing building will remain occupied throughout the works.",
    ),
    (
        "8.6",
        "Existing systems",
        "Any shutdown of the existing sprinkler system shall be agreed seven days in advance.",
    ),
)

# The risks those clauses, and clause 6.6, seed: (kind, the clauses cited).
EXPECTED_DESIGN_RISKS: tuple[tuple[str, tuple[str, ...]], ...] = (
    ("design_and_build", ("8.1",)),
    ("shop_drawings", ("6.6",)),
    ("hydraulic_calculations", ("6.6", "8.2")),
    ("qp_engagement", ("8.3",)),
)
EXPECTED_WORDING_RISKS: tuple[tuple[str, str], ...] = (
    ("night_work", "8.4"),
    ("occupied_building", "8.5"),
    ("shutdown", "8.6"),
)


def with_risks() -> tuple[tuple[str, str, str], ...]:
    """The extended specification with section 8."""
    return (*extended(), *RISK_CLAUSES)


# Every obligation of the extended specification: (category, clause, quantities).
EXPECTED_OBLIGATIONS: tuple[tuple[str, str, dict[str, object]], ...] = (
    ("testing", "6.1", {"pressure": 14, "pressure_unit": "bar", "duration_hours": 2}),
    ("flushing", "6.1", {}),
    ("painting", "6.2", {"coats": 2}),
    ("identification", "6.2", {}),
    ("commissioning", "6.3", {}),
    ("warranty", "6.4", {"period_months": 12}),
    ("defects_liability", "6.4", {"period_months": 12}),
    ("maintenance", "6.4", {"period_months": 12}),
    ("spares", "6.5", {"count": 2, "count_of": "sets"}),
    ("training", "6.5", {"period_days": 2}),
    ("submittals", "6.6", {"period_weeks": 4}),
    ("authority", "6.7", {}),
    ("approved_makes", "6.8", {}),
)

# What the drawings say against it (`synthetic_qto.car_park_plan`), and every issue seeded
# between the two: (rule, the clause it cites, what it is about).
CAR_PARK_NOTES = (
    "ALL SPRINKLER PIPEWORK TO BE BLACK STEEL",
    "PIPES DN65 AND ABOVE: WELDED JOINTS",
)
SEEDED_ISSUES: tuple[tuple[str, str | None, str], ...] = (
    ("conflict:pipe_material", "2.1.3", "pipe_material"),  # the P1-06 contradiction
    ("conflict:joining_method", "2.1.2", "joining_method"),
    ("missing_from_drawings", "7.6", "test_header"),
    ("missing_from_specification", None, "sprinkler_upright"),
    ("ambiguous_clause", "7.7", "where required"),
    ("ambiguous_clause", "7.8", "or equal"),
)

# The interface rows of the scope matrix, the same for every system: (key, status, clause).
EXPECTED_INTERFACES: tuple[tuple[str, str, str | None], ...] = (
    ("power_supply", "by_others", "7.1"),
    ("water_supply", "included", "7.2"),
    ("builders_works", "excluded", "7.3"),
    ("ceiling_access", "unclear", "7.4"),
    ("painting", "included", "6.2"),
    ("fire_alarm_interface", "by_others", "7.5"),
    ("fire_stopping", "included", "7.7"),
    ("drainage", "unclear", None),
)


@dataclass(frozen=True)
class Expected:
    system: str
    attribute: str
    value: str
    clause: str
    dn_min: int | None = None
    dn_max: int | None = None
    condition: str | None = None


EXPECTED: tuple[Expected, ...] = (
    Expected("sprinkler", "pipe_material", "black_steel", "2.1.1", None, 50),
    Expected("sprinkler", "pipe_standard", "BS EN 10255", "2.1.1", None, 50),
    Expected("sprinkler", "pipe_class", "Heavy", "2.1.1", None, 50),
    Expected("sprinkler", "pipe_material", "black_steel", "2.1.1", 65, None),
    Expected("sprinkler", "pipe_standard", "ASTM A53 Grade B", "2.1.1", 65, None),
    Expected("sprinkler", "pipe_class", "Schedule 40", "2.1.1", 65, None),
    Expected("sprinkler", "joining_method", "threaded", "2.1.2", None, 50),
    Expected("sprinkler", "joining_method", "grooved", "2.1.2", 65, None),
    Expected(
        "sprinkler", "pipe_material", "galvanised_steel", "2.1.3", None, None, "basement car park"
    ),
    Expected("sprinkler", "sprinkler_type", "pendent", "2.2.1"),
    Expected("sprinkler", "response", "quick", "2.2.1"),
    Expected("sprinkler", "k_factor", "80", "2.2.1"),
    Expected("sprinkler", "temperature_rating_c", "68", "2.2.1"),
    Expected("sprinkler", "finish", "chrome", "2.2.1"),
    Expected("sprinkler", "sprinkler_type", "sidewall", "2.2.2"),
    Expected("sprinkler", "response", "quick", "2.2.2"),
    Expected("sprinkler", "k_factor", "80", "2.2.2"),
    Expected("sprinkler", "temperature_rating_c", "68", "2.2.2"),
    Expected("sprinkler", "finish", "white", "2.2.2"),
    Expected("sprinkler", "approved_make", "Tyco", "2.2.3"),
    Expected("sprinkler", "approved_make", "Viking", "2.2.3"),
    Expected("sprinkler", "approved_make", "Reliable", "2.2.3"),
    Expected("hose_reel", "pipe_material", "galvanised_steel", "3.1"),
    Expected("hose_reel", "pipe_standard", "BS EN 10255", "3.1"),
    Expected("hose_reel", "pipe_class", "Medium", "3.1"),
    Expected("hose_reel", "joining_method", "threaded", "3.1"),
    Expected("hydrant", "pipe_material", "ductile_iron", "4.1", 100, None),
    Expected("hydrant", "pipe_standard", "BS EN 545", "4.1", 100, None),
    Expected("hydrant", "joining_method", "flanged", "4.1", 100, None),
)


EXPECTED_SUPPORTS: tuple[Expected, ...] = (
    Expected("sprinkler", "hanger_spacing_mm", "3000", "2.3.1", None, 50),
    Expected("sprinkler", "hanger_spacing_mm", "4000", "2.3.1", 65, None),
)
EXPECTED_SEISMIC: tuple[Expected, ...] = (
    Expected("sprinkler", "seismic_restraint", "required", "2.3.2", 65, None),
)


def specification_docx(
    title: str = "PARTICULAR SPECIFICATION FOR FIRE PROTECTION SERVICES",
    revision: str = "Rev B",
    clauses: tuple[tuple[str, str, str], ...] = CLAUSES,
) -> bytes:
    """The specification as a Word document: heading styles for sections, numbered text."""
    import docx

    document = docx.Document()
    document.core_properties.title = title
    document.add_paragraph(title, style="Title")
    document.add_paragraph(f"Revision {revision}")
    for number, heading, body in clauses:
        depth = number.count(".")
        if not body:
            document.add_heading(f"{number} {heading}", level=min(depth + 1, 3))
            continue
        if heading:
            document.add_heading(f"{number} {heading}", level=min(depth + 1, 3))
            document.add_paragraph(body)
        else:
            document.add_paragraph(f"{number} {body}")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def specification_pdf(
    revision: str = "Rev B", clauses: tuple[tuple[str, str, str], ...] = CLAUSES
) -> bytes:
    """The same specification as a PDF, laid out with larger bold headings."""
    import matplotlib

    matplotlib.use("Agg")
    from matplotlib import pyplot as plt
    from matplotlib.backends.backend_pdf import PdfPages

    lines: list[tuple[str, float, bool]] = [
        ("PARTICULAR SPECIFICATION FOR FIRE PROTECTION SERVICES", 14, True),
        (f"Revision {revision}", 10, False),
    ]
    for number, heading, body in clauses:
        if heading:
            lines.append((f"{number} {heading}", 12 if "." not in number else 11, True))
        if body:
            text = f"{number} {body}" if not heading else body
            lines.extend((part, 10, False) for part in _wrap(text, 92))
    buffer = io.BytesIO()
    with PdfPages(buffer) as pages:
        for start in range(0, len(lines), 40):
            figure = plt.figure(figsize=(8.27, 11.69))
            y = 0.95
            for text, size, bold in lines[start : start + 40]:
                figure.text(0.08, y, text, fontsize=size, weight="bold" if bold else "normal")
                y -= 0.022
            pages.savefig(figure)
            plt.close(figure)
    return buffer.getvalue()


def _wrap(text: str, width: int) -> list[str]:
    words, lines, line = text.split(), [], ""
    for word in words:
        if len(line) + len(word) + 1 > width:
            lines.append(line)
            line = word
        else:
            line = f"{line} {word}".strip()
    if line:
        lines.append(line)
    return lines
