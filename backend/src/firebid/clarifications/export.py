"""The clarification register as a file a person downloads (FR-RFI-07).

A template is a list of column headings and the field under each (`clarifications.yaml`):
the company's default form, or a client's own log. The same rows go into an xlsx workbook or
a docx document. Building the file is all this does: it is handed to the person who asked
for it, and nothing here, or anywhere in the platform, sends it on (autonomy level L3 is
not permitted).
"""

from __future__ import annotations

import io
from dataclasses import dataclass
from typing import Any

from firebid.clarifications.drafting import settings

FIELDS = (
    "number",
    "subject",
    "project",
    "level_grid",
    "sheets",
    "problem",
    "evidence",
    "options",
    "cost_impact",
    "programme_impact",
    "reviewer",
    "status",
    "due",
    "issued",
    "response",
)
FORMATS = {
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}


@dataclass(frozen=True)
class Template:
    key: str
    label: str
    title: str
    columns: tuple[tuple[str, str], ...]  # (heading, field)


def templates(config: dict[str, Any] | None = None) -> list[Template]:
    found = (config if config is not None else settings()).get("templates") or []
    out = []
    for item in found:
        columns = tuple((str(c["heading"]), str(c["field"])) for c in item.get("columns") or [])
        unknown = [name for _, name in columns if name not in FIELDS]
        if unknown:
            raise ValueError(f"template {item.get('key')!r} names unknown fields: {unknown}")
        out.append(
            Template(
                key=str(item["key"]),
                label=str(item.get("label") or item["key"]),
                title=str(item.get("title") or "Tender clarifications"),
                columns=columns,
            )
        )
    return out


def template(key: str, config: dict[str, Any] | None = None) -> Template:
    for found in templates(config):
        if found.key == key:
            return found
    raise KeyError(f"no clarification template {key!r}")


def cells(chosen: Template, rows: list[dict[str, str]]) -> list[list[str]]:
    """The table a template makes of the rows: its headings, then a row a clarification."""
    return [
        [heading for heading, _ in chosen.columns],
        *([row.get(name, "") for _, name in chosen.columns] for row in rows),
    ]


def workbook(chosen: Template, project: str, rows: list[dict[str, str]]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    book = Workbook()
    sheet = book.active
    assert sheet is not None  # noqa: S101 - a new workbook has one
    sheet.title = "Clarifications"
    sheet.append([chosen.title])
    sheet.append([project])
    sheet.append([])
    table = cells(chosen, rows)
    for line in table:
        sheet.append(line)
    sheet["A1"].font = Font(bold=True, size=14)
    for cell in sheet[4]:
        cell.font = Font(bold=True)
    for index, (_, name) in enumerate(chosen.columns, start=1):
        letter = get_column_letter(index)
        sheet.column_dimensions[letter].width = (
            60 if name in ("problem", "evidence", "options") else 22
        )
    for row in sheet.iter_rows(min_row=5):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def document(chosen: Template, project: str, rows: list[dict[str, str]]) -> bytes:
    from docx import Document

    file = Document()
    file.add_heading(chosen.title, level=1)
    file.add_paragraph(project)
    table = cells(chosen, rows)
    grid = file.add_table(rows=len(table), cols=len(chosen.columns))
    grid.style = "Table Grid"
    for r, line in enumerate(table):
        for c, words in enumerate(line):
            cell = grid.cell(r, c)
            cell.text = words
            if r == 0:
                for run in cell.paragraphs[0].runs:
                    run.bold = True
    buffer = io.BytesIO()
    file.save(buffer)
    return buffer.getvalue()


def build(key: str, kind: str, project: str, rows: list[dict[str, str]]) -> bytes:
    if kind not in FORMATS:
        raise ValueError("a clarification register is exported as xlsx or docx")
    chosen = template(key)
    return workbook(chosen, project, rows) if kind == "xlsx" else document(chosen, project, rows)
