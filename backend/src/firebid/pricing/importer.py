"""Reading a rate list workbook: every row checked, every problem reported by row and column.

The first sheet with the expected header is read. A file with any error imports nothing:
the report says what to fix. `read_json` is the sandbox's entry point, since an uploaded
workbook is opened only by the parser workers.
"""

from __future__ import annotations

import io
import re
from dataclasses import asdict, dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from typing import Any

from firebid.pricing.keys import ItemKey, unit_of
from firebid.pricing.rates import SOURCES

# Column: the header words it may appear under.
COLUMNS: dict[str, tuple[str, ...]] = {
    "type": ("type", "item type", "item"),
    "dn": ("dn", "size", "nominal size", "dn (mm)"),
    "material": ("material",),
    "schedule": ("schedule", "class", "pipe class"),
    "joining": ("joining", "joint", "joining method"),
    "brand": ("brand", "make", "manufacturer"),
    "description": ("description",),
    "unit": ("unit", "uom"),
    "rate": ("rate", "unit rate", "rate (sgd)", "unit rate (sgd)"),
    "source_type": ("source type", "source"),
    "source_reference": ("source reference", "reference", "source ref"),
    "effective_from": ("effective from", "effective date", "date"),
    "valid_until": ("valid until", "validity", "validity end", "expiry"),
}
REQUIRED = (
    "type",
    "description",
    "unit",
    "rate",
    "source_type",
    "source_reference",
    "effective_from",
)
SOURCE_WORDS = {
    "company standard": "company_standard",
    "company_standard": "company_standard",
    "standard": "company_standard",
    "po": "purchase_order",
    "purchase order": "purchase_order",
    "purchase_order": "purchase_order",
    "quotation": "quotation",
    "quote": "quotation",
}
HEADER_SEARCH_ROWS = 20


@dataclass
class Problem:
    row: int
    column: str
    message: str


@dataclass
class RateRow:
    row: int
    key: str
    parts: dict[str, str]
    description: str
    unit: str
    rate: str  # a Decimal, as text, so it survives the trip out of the sandbox exactly
    source_type: str
    source_reference: str
    effective_from: str
    valid_until: str | None


@dataclass
class RateList:
    sheet: str | None = None
    rows: list[RateRow] = field(default_factory=list)
    problems: list[Problem] = field(default_factory=list)


def _norm(value: Any) -> str:
    return " ".join(str(value or "").split()).lower()


def _header(rows: list[tuple[Any, ...]]) -> tuple[int, dict[str, int]] | None:
    for index, row in enumerate(rows[:HEADER_SEARCH_ROWS]):
        cells = [_norm(cell) for cell in row]
        found: dict[str, int] = {}
        for name, words in COLUMNS.items():
            for position, cell in enumerate(cells):
                if cell in words and position not in found.values():
                    found[name] = position
                    break
        if all(name in found for name in REQUIRED):
            return index, found
    return None


def _date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value or "").strip()
    for pattern in ("%Y-%m-%d", "%d/%m/%Y", "%d.%m.%Y", "%d-%m-%Y", "%d %b %Y", "%d %B %Y"):
        try:
            return datetime.strptime(text, pattern).date()
        except ValueError:
            continue
    return None


def _decimal(value: Any) -> Decimal | None:
    if isinstance(value, bool):
        return None
    text = re.sub(r"[,\s]|S\$|\$|SGD", "", str(value if value is not None else ""), flags=re.I)
    try:
        number = Decimal(text)
    except InvalidOperation:
        return None
    return number if number.is_finite() else None


def read_rows(rows: list[tuple[Any, ...]], sheet: str | None = None) -> RateList:
    out = RateList(sheet=sheet)
    found = _header(rows)
    if found is None:
        out.problems.append(
            Problem(
                0,
                "header",
                "no header row with the columns "
                + ", ".join(COLUMNS[name][0] for name in REQUIRED),
            )
        )
        return out
    header, columns = found
    seen: dict[tuple[str, str, str, str], int] = {}
    for index in range(header + 1, len(rows)):
        cells = rows[index]
        number = index + 1  # as the spreadsheet shows it

        def cell(name: str, cells: tuple[Any, ...] = cells) -> Any:
            position = columns.get(name)
            return cells[position] if position is not None and position < len(cells) else None

        if all(_norm(value) == "" for value in cells):
            continue
        problems: list[Problem] = []
        for name in REQUIRED:
            if _norm(cell(name)) == "":
                problems.append(Problem(number, COLUMNS[name][0], "is empty"))
        rate = _decimal(cell("rate"))
        if _norm(cell("rate")) and (rate is None or rate <= 0):
            problems.append(Problem(number, "rate", f"{cell('rate')!r} is not a positive amount"))
        elif rate is not None and rate != rate.quantize(Decimal("0.01")):
            problems.append(Problem(number, "rate", f"{rate} has more than two decimal places"))
        unit = unit_of(str(cell("unit") or ""))
        if _norm(cell("unit")) and unit is None:
            problems.append(Problem(number, "unit", f"{cell('unit')!r} is not a known unit"))
        source = SOURCE_WORDS.get(_norm(cell("source_type")))
        if _norm(cell("source_type")) and source not in SOURCES:
            problems.append(
                Problem(
                    number,
                    "source type",
                    f"{cell('source_type')!r} is not company standard, PO or quotation",
                )
            )
        effective = _date(cell("effective_from"))
        if _norm(cell("effective_from")) and effective is None:
            problems.append(
                Problem(number, "effective from", f"{cell('effective_from')!r} is not a date")
            )
        until = _date(cell("valid_until")) if _norm(cell("valid_until")) else None
        if _norm(cell("valid_until")) and until is None:
            problems.append(
                Problem(number, "valid until", f"{cell('valid_until')!r} is not a date")
            )
        if effective and until and until < effective:
            problems.append(Problem(number, "valid until", "is before the effective date"))
        key = ItemKey.of(
            type=cell("type"),
            dn=cell("dn"),
            material=cell("material"),
            schedule=cell("schedule"),
            joining=cell("joining"),
            brand=cell("brand"),
        )
        if problems:
            out.problems.extend(problems)
            continue
        if rate is None or effective is None or source is None or not unit:  # pragma: no cover
            continue  # every one of these was reported above
        identity = (key.text(), unit, source, " ".join(str(cell("source_reference")).split()))
        if identity in seen:
            out.problems.append(
                Problem(
                    number,
                    "type",
                    f"the same item, unit and source as row {seen[identity]}",
                )
            )
            continue
        seen[identity] = number
        out.rows.append(
            RateRow(
                row=number,
                key=key.text(),
                parts=key.parts(),
                description=" ".join(str(cell("description")).split()),
                unit=unit,
                rate=str(rate.quantize(Decimal("0.01"))),
                source_type=source,
                source_reference=" ".join(str(cell("source_reference")).split()),
                effective_from=effective.isoformat(),
                valid_until=until.isoformat() if until else None,
            )
        )
    if not out.rows and not out.problems:
        out.problems.append(Problem(header + 1, "rows", "the list has no rates under its header"))
    return out


def read_workbook(payload: bytes) -> RateList:
    from openpyxl import load_workbook

    book = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    try:
        first: RateList | None = None
        for sheet in book.worksheets:
            rows = [tuple(row) for row in sheet.iter_rows(values_only=True)]
            found = read_rows(rows, sheet.title)
            if _header(rows) is not None:
                return found
            first = first or found
        return first or RateList(problems=[Problem(0, "header", "the workbook has no sheets")])
    finally:
        book.close()


def read_json(payload: bytes) -> dict[str, Any]:
    """The sandbox's entry point: the rate list as plain data."""
    found = read_workbook(payload)
    return {
        "sheet": found.sheet,
        "rows": [asdict(row) for row in found.rows],
        "problems": [asdict(problem) for problem in found.problems],
    }
