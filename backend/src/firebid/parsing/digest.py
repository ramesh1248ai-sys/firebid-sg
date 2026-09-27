"""A short digest of a document to classify it by. Runs in the sandbox, returns plain data.

Only the opening of each file is read: a classification needs the title and the first few
headings, not three hundred pages of specification. That keeps the job fast and the digest
small enough to store beside the document and to send, if needed, to the model.
"""

from __future__ import annotations

import io
from typing import Any

# How much of each file is read.
PDF_PAGES = 3
DOCX_PARAGRAPHS = 300
XLSX_SHEETS = 8
XLSX_ROWS = 30
TEXT_LIMIT = 8_000


def digest(payload: bytes, kind: str) -> dict[str, Any]:
    if kind == "pdf":
        return _pdf(payload)
    if kind == "docx":
        return _docx(payload)
    if kind == "xlsx":
        return _xlsx(payload)
    return {"text": "", "pages": 0, "sheet_names": [], "header_rows": []}


def _pdf(payload: bytes) -> dict[str, Any]:
    from firebid.parsing.pdf import _open

    document = _open(payload)
    try:
        parts = []
        for index in range(min(len(document), PDF_PAGES)):
            parts.append(document[index].get_textpage().get_text_range())
        text = "\n".join(parts)
        return {
            "text": text[:TEXT_LIMIT],
            "pages": len(document),
            "sheet_names": [],
            "header_rows": [],
        }
    finally:
        document.close()


def _docx(payload: bytes) -> dict[str, Any]:
    import docx

    document = docx.Document(io.BytesIO(payload))
    lines = [paragraph.text for paragraph in document.paragraphs[:DOCX_PARAGRAPHS]]
    header_rows = []
    for table in document.tables[:10]:
        if table.rows:
            header_rows.append([cell.text.strip() for cell in table.rows[0].cells])
    title = document.core_properties.title or ""
    text = "\n".join([title, *lines]).strip()
    return {"text": text[:TEXT_LIMIT], "pages": 0, "sheet_names": [], "header_rows": header_rows}


def _xlsx(payload: bytes) -> dict[str, Any]:
    import openpyxl

    workbook = openpyxl.load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    try:
        lines: list[str] = []
        header_rows: list[list[str]] = []
        names = list(workbook.sheetnames)
        for name in names[:XLSX_SHEETS]:
            sheet = workbook[name]
            header_found = False
            for row in sheet.iter_rows(max_row=XLSX_ROWS, values_only=True):
                cells = [str(value).strip() for value in row if value not in (None, "")]
                if not cells:
                    continue
                lines.append(" | ".join(cells))
                # The header is the first row with several text cells: titles above a BOQ
                # are one cell wide, and its column headings are not.
                text_cells = [cell for cell in cells if not _numeric(cell)]
                if not header_found and len(text_cells) >= 3:
                    header_rows.append(text_cells)
                    header_found = True
        return {
            "text": "\n".join(lines)[:TEXT_LIMIT],
            "pages": 0,
            "sheet_names": names,
            "header_rows": header_rows,
        }
    finally:
        workbook.close()


def _numeric(value: str) -> bool:
    try:
        float(value.replace(",", ""))
    except ValueError:
        return False
    return True
