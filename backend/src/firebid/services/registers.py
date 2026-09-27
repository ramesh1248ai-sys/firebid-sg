"""The drawing and specification registers, and what people do to them (FR-DOC-03).

The registers are views over revisions: nothing here decides which revision is Current, that
is the revision service's job through the state machine. What this module adds is the
people's side: confirming or correcting a reading, withdrawing a page that is not a drawing,
confirming a document's type, exporting the registers, and the Estimator's "Register
confirmed", which is stage S1's output.
"""

from __future__ import annotations

import hashlib
import io
import json
import uuid
from dataclasses import asdict, dataclass
from datetime import UTC, date, datetime
from typing import Any, cast

import structlog
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from firebid.db.audit import record_event
from firebid.db.models.core import Bid
from firebid.db.models.documents import (
    Addendum,
    Document,
    DocumentRevision,
    RegisterConfirmation,
    Sheet,
    SheetRevision,
    TitleBlockLayout,
)
from firebid.db.models.workflow import HumanTask
from firebid.domain.actors import Actor, AuditContext
from firebid.domain.state_machines import SheetRevisionState
from firebid.drawings.title_block import Box, Field, FieldReading, Layout, TitleBlockReading, learn
from firebid.ingest.classification import RULE_THRESHOLD, DocType
from firebid.services.transitions import apply_transition

log = structlog.get_logger("firebid.registers")
State = SheetRevisionState


# --- The registers ---------------------------------------------------------------------------


@dataclass(frozen=True)
class DrawingRow:
    revision_id: uuid.UUID
    sheet_id: uuid.UUID
    document_id: uuid.UUID
    filename: str
    sheet_number: str | None
    title: str | None
    revision: str | None
    revision_date: date | None
    discipline: str | None
    level: str | None
    zone: str | None
    scale: str | None
    state: str
    confidence: float | None
    read_by: str | None
    conflict_reason: str | None
    addendum: str | None
    content_class: str | None
    quality_band: str | None
    manual_takeoff_recommended: bool


@dataclass(frozen=True)
class DocumentRow:
    revision_id: uuid.UUID
    document_id: uuid.UUID
    filename: str
    doc_type: str | None
    type_confidence: float | None
    type_decided_by: str | None
    doc_key: str | None
    title: str | None
    revision: str | None
    revision_date: date | None
    state: str
    conflict_reason: str | None
    addendum: str | None


# Current first, then what still needs a person, then the history.
STATE_ORDER = {
    "current": 0,
    "conflict": 1,
    "received": 2,
    "registered": 3,
    "superseded": 4,
    "withdrawn": 5,
}


def drawing_register(
    session: Session,
    bid_id: uuid.UUID,
    *,
    discipline: str | None = None,
    level: str | None = None,
    state: str | None = None,
) -> list[DrawingRow]:
    query = (
        select(SheetRevision, Sheet, Document, Addendum.number)
        .join(Sheet, Sheet.id == SheetRevision.sheet_id)
        .join(Document, Document.id == Sheet.document_id)
        .outerjoin(Addendum, Addendum.id == SheetRevision.addendum_id)
        .where(SheetRevision.bid_id == bid_id)
    )
    if discipline:
        query = query.where(SheetRevision.discipline == discipline)
    if level:
        query = query.where(SheetRevision.level == level)
    if state:
        query = query.where(SheetRevision.state == state)
    rows = [
        DrawingRow(
            revision_id=revision.id,
            sheet_id=sheet.id,
            document_id=document.id,
            filename=document.filename,
            sheet_number=revision.sheet_number,
            title=revision.title,
            revision=revision.revision_label,
            revision_date=revision.revision_date,
            discipline=revision.discipline,
            level=revision.level,
            zone=revision.zone,
            scale=revision.scale_text,
            state=revision.state,
            confidence=revision.source_confidence,
            read_by=revision.extraction_method,
            conflict_reason=revision.conflict_reason,
            addendum=addendum,
            content_class=sheet.content_class,
            quality_band=sheet.quality_band,
            manual_takeoff_recommended=sheet.manual_takeoff_recommended,
        )
        for revision, sheet, document, addendum in session.execute(query).all()
    ]
    rows.sort(
        key=lambda row: (row.sheet_number or "~", STATE_ORDER.get(row.state, 9), row.revision or "")
    )
    return rows


def document_register(
    session: Session,
    bid_id: uuid.UUID,
    *,
    doc_type: str | None = None,
    state: str | None = None,
) -> list[DocumentRow]:
    query = (
        select(DocumentRevision, Document, Addendum.number)
        .join(Document, Document.id == DocumentRevision.document_id)
        .outerjoin(Addendum, Addendum.id == DocumentRevision.addendum_id)
        .where(DocumentRevision.bid_id == bid_id)
    )
    if doc_type:
        query = query.where(Document.doc_type == doc_type)
    if state:
        query = query.where(DocumentRevision.state == state)
    rows = [
        DocumentRow(
            revision_id=revision.id,
            document_id=document.id,
            filename=document.filename,
            doc_type=document.doc_type,
            type_confidence=document.doc_type_confidence,
            type_decided_by=str((document.classification or {}).get("method") or "") or None,
            doc_key=revision.doc_key,
            title=revision.title,
            revision=revision.revision_label,
            revision_date=revision.revision_date,
            state=revision.state,
            conflict_reason=revision.conflict_reason,
            addendum=addendum,
        )
        for revision, document, addendum in session.execute(query).all()
    ]
    rows.sort(
        key=lambda row: (row.doc_type or "~", row.doc_key or "~", STATE_ORDER.get(row.state, 9))
    )
    return rows


# --- Export ----------------------------------------------------------------------------------

DRAWING_COLUMNS: tuple[tuple[str, str, int], ...] = (
    ("Drawing number", "sheet_number", 18),
    ("Title", "title", 40),
    ("Revision", "revision", 10),
    ("Revision date", "revision_date", 14),
    ("Status", "state", 12),
    ("Discipline", "discipline", 18),
    ("Level", "level", 8),
    ("Zone", "zone", 8),
    ("Scale", "scale", 10),
    ("Addendum", "addendum", 10),
    ("Content", "content_class", 10),
    ("Expected accuracy", "quality_band", 12),
    ("Manual takeoff recommended", "manual_takeoff_recommended", 14),
    ("Read by", "read_by", 12),
    ("Confidence", "confidence", 11),
    ("Conflict", "conflict_reason", 50),
    ("File", "filename", 36),
)
DOCUMENT_COLUMNS: tuple[tuple[str, str, int], ...] = (
    ("Document", "doc_key", 30),
    ("Title", "title", 44),
    ("Type", "doc_type", 22),
    ("Revision", "revision", 10),
    ("Revision date", "revision_date", 14),
    ("Status", "state", 12),
    ("Addendum", "addendum", 10),
    ("Type decided by", "type_decided_by", 14),
    ("Type confidence", "type_confidence", 12),
    ("Conflict", "conflict_reason", 50),
    ("File", "filename", 36),
)


REGISTER_TITLES = {"drawings": "Drawing register", "documents": "Specification register"}


def export_xlsx(
    kind: str, rows: list[DrawingRow] | list[DocumentRow], *, bid_reference: str, by: str
) -> bytes:
    """The register as a workbook an estimator can open, filter and send (FR-DOC-03).

    Built fresh, so nothing of a client's workbook is at stake. The header is row 1 and the
    panes are frozen under it; dates are dates and confidences numbers, so Excel sorts and
    filters them properly. A second sheet says what this is and when it was taken.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    columns = DRAWING_COLUMNS if kind == "drawings" else DOCUMENT_COLUMNS
    title = REGISTER_TITLES[kind]

    book = Workbook()
    sheet = book.active
    if sheet is None:  # pragma: no cover - a new workbook always has one
        raise RuntimeError("openpyxl made a workbook with no sheet")
    sheet.title = title[:31]
    sheet.append([heading for heading, _, _ in columns])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in rows:
        values = asdict(row)
        sheet.append([_cell(values[field]) for _, field, _ in columns])
    for index, (_, _, width) in enumerate(columns, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for cell in sheet["D"][1:]:
        cell.number_format = "dd.mm.yyyy"
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    about = book.create_sheet("About")
    about.append(["Register", title])
    about.append(["Bid", bid_reference])
    about.append(["Exported", datetime.now(UTC).strftime("%d.%m.%Y %H:%M UTC")])
    about.append(["Exported by", by])
    about.append(["Rows", len(rows)])
    about.column_dimensions["A"].width = 14
    about.column_dimensions["B"].width = 40

    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def _cell(value: Any) -> Any:
    if isinstance(value, uuid.UUID):
        return str(value)
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, float):
        return round(value, 2)
    return value


# --- What people do --------------------------------------------------------------------------


def _audit(
    session: Session,
    bid_id: uuid.UUID,
    actor: Actor,
    action: str,
    entity: Any,
    before: dict[str, Any],
    after: dict[str, Any],
    reason: str | None = None,
) -> None:
    organisation_id = session.execute(
        select(Bid.organisation_id).where(Bid.id == bid_id)
    ).scalar_one()
    record_event(
        session,
        context=AuditContext(organisation_id=organisation_id, bid_id=bid_id),
        actor=actor,
        action=action,
        entity_type=type(entity).__tablename__,
        entity_id=entity.id,
        before=before,
        after=after,
        reason=reason,
    )


def confirm_reading(
    session: Session,
    revision: SheetRevision,
    *,
    actor: Actor,
    sheet_number: str,
    revision_label: str,
    title: str | None = None,
    consultant: str | None = None,
) -> SheetRevision:
    """A person says what a sheet is. Their reading is the decision (guardrail 2).

    Given a consultant's name, the layout of this title block is remembered for them, so the
    next sheets from the same consultant are read by position with no model and no person.
    """
    from firebid.services.revisions import settle

    if revision.state != str(State.RECEIVED):
        raise ValueError(
            "only an unregistered reading is confirmed; a registered one is resolved or withdrawn"
        )
    before = {
        "sheet_number": revision.sheet_number,
        "revision": revision.revision_label,
        "title": revision.title,
    }
    revision.sheet_number = sheet_number.strip().upper()
    revision.revision_label = revision_label.strip().upper()
    if title:
        revision.title = title.strip()[:300]
    revision.extraction_method = "person"
    revision.source_confidence = 1.0
    reading = dict(revision.reading or {})
    reading["confirmed_by"] = actor.label
    revision.reading = reading
    _audit(
        session,
        revision.bid_id,
        actor,
        "title block: confirmed",
        revision,
        before,
        {
            "sheet_number": revision.sheet_number,
            "revision": revision.revision_label,
            "title": revision.title,
        },
    )
    if consultant:
        remember_layout(session, revision, consultant.strip(), actor)
    _close_review_tasks(session, revision, actor)
    settle(session, revision, actor)
    session.flush()
    return revision


def remember_layout(
    session: Session, revision: SheetRevision, consultant: str, actor: Actor
) -> TitleBlockLayout | None:
    """Learn where this consultant puts each field, from the stored, now confirmed, reading."""
    reading = cast(dict[str, Any], revision.reading or {})
    page = reading.get("page")
    marks = reading.get("fingerprint")
    fields = cast(dict[str, dict[str, Any]], reading.get("fields") or {})
    region = reading.get("region")
    if not page or not marks or not region:
        return None
    found: dict[Field, FieldReading] = {}
    for name in (Field.SHEET_NUMBER, Field.REVISION, Field.TITLE, Field.SCALE, Field.REVISION_DATE):
        entry = fields.get(str(name)) or {}
        box = entry.get("box")
        if box:
            found[name] = FieldReading(str(entry.get("value") or ""), 1.0, "person", Box(*box))
    if Field.SHEET_NUMBER not in found:
        return None
    page_box = Box(*page)
    layout: Layout = learn(TitleBlockReading(region=Box(*region), fields=found), page_box)
    organisation_id = session.execute(
        select(Bid.organisation_id).where(Bid.id == revision.bid_id)
    ).scalar_one()
    remembered = TitleBlockLayout(
        organisation_id=organisation_id,
        consultant=consultant[:200],
        fingerprint=marks,
        layout=layout.to_json(),
        created_by_id=actor.id,
    )
    session.add(remembered)
    session.flush()
    log.info("title_block_layout_learnt", consultant=consultant)
    return remembered


def withdraw(
    session: Session, revision: SheetRevision | DocumentRevision, *, actor: Actor, reason: str
) -> None:
    """Take a sheet or document out of the register: a cover page, a duplicate, a mistake."""
    apply_transition(session, revision, target=State.WITHDRAWN, actor=actor, reason=reason)
    _close_review_tasks(session, revision, actor)
    session.flush()


def confirm_document_type(
    session: Session, document: Document, doc_type: DocType, *, actor: Actor
) -> Document:
    from firebid.services.classification import confirm, register_document

    before = {"doc_type": document.doc_type}
    confirm(document, doc_type, actor.label)
    _audit(
        session,
        document.bid_id,
        actor,
        "document type: confirmed",
        document,
        before,
        {"doc_type": document.doc_type},
    )
    register_document(session, document)
    for task in _open_tasks(session, document.bid_id):
        if task.payload.get("document_id") == str(document.id) or (
            task.payload.get("review_task_for") == str(document.id)
        ):
            _done(task, actor)
    session.flush()
    return document


def identify_document(
    session: Session,
    revision: DocumentRevision,
    *,
    actor: Actor,
    doc_key: str,
    revision_label: str | None,
) -> DocumentRevision:
    from firebid.ingest.document_identity import normalise_title
    from firebid.services.revisions import settle_document

    if revision.state != str(State.RECEIVED):
        raise ValueError("only an unregistered document is identified")
    before = {"doc_key": revision.doc_key, "revision": revision.revision_label}
    revision.doc_key = normalise_title(doc_key) or doc_key.strip().upper()
    revision.revision_label = revision_label.strip().upper() if revision_label else None
    # A person's identification is the decision; the sources are not voted against it.
    revision.sources = {"person": revision.revision_label}
    _audit(
        session,
        revision.bid_id,
        actor,
        "document: identified",
        revision,
        before,
        {"doc_key": revision.doc_key, "revision": revision.revision_label},
    )
    _close_review_tasks(session, revision, actor)
    settle_document(session, revision, actor)
    session.flush()
    return revision


def _open_tasks(session: Session, bid_id: uuid.UUID) -> list[HumanTask]:
    return list(
        session.execute(
            select(HumanTask).where(HumanTask.bid_id == bid_id, HumanTask.state == "open")
        ).scalars()
    )


def _done(task: HumanTask, actor: Actor) -> None:
    task.state = "done"
    task.completed_at = datetime.now(UTC)
    task.completed_by_id = actor.id


def _close_review_tasks(
    session: Session, revision: SheetRevision | DocumentRevision, actor: Actor
) -> None:
    key = f"{type(revision).__tablename__}_id"
    for task in _open_tasks(session, revision.bid_id):
        if task.payload.get(key) == str(revision.id):
            _done(task, actor)


# --- Register confirmed: stage S1's output ---------------------------------------------------


@dataclass(frozen=True)
class Blockers:
    conflicts: int
    unidentified: int
    unsure_types: int

    @property
    def any(self) -> bool:
        return bool(self.conflicts or self.unidentified or self.unsure_types)

    def describe(self) -> str:
        parts = []
        if self.conflicts:
            parts.append(f"{self.conflicts} in Conflict")
        if self.unidentified:
            parts.append(f"{self.unidentified} not yet identified")
        if self.unsure_types:
            parts.append(f"{self.unsure_types} with an unconfirmed document type")
        return ", ".join(parts)


class RegisterNotReady(Exception):
    def __init__(self, blockers: Blockers) -> None:
        super().__init__(f"The registers are not ready to confirm: {blockers.describe()}.")
        self.blockers = blockers


def blockers(session: Session, bid_id: uuid.UUID) -> Blockers:
    """What stops the registers being confirmed: anything a person has still to decide."""

    def count(model: type[SheetRevision] | type[DocumentRevision], state: str) -> int:
        return int(
            session.execute(
                select(func.count())
                .select_from(model)
                .where(model.bid_id == bid_id, model.state == state)
            ).scalar_one()
        )

    unsure = session.execute(
        select(Document).where(
            Document.bid_id == bid_id,
            Document.doc_type.is_not(None),
            Document.doc_type_confidence < RULE_THRESHOLD,
        )
    ).scalars()
    unsure_types = sum(
        1 for document in unsure if (document.classification or {}).get("method") != "person"
    )
    return Blockers(
        conflicts=count(SheetRevision, "conflict") + count(DocumentRevision, "conflict"),
        unidentified=count(SheetRevision, "received") + count(DocumentRevision, "received"),
        unsure_types=unsure_types,
    )


def confirm_registers(
    session: Session, bid: Bid, *, actor: Actor, role: str, comment: str | None = None
) -> RegisterConfirmation:
    """The Estimator confirms the registers are complete and right (stage S1's output).

    Refused while anything is still a question: a register with a drawing in Conflict does
    not yet have exactly one Current revision per sheet, whatever it looks like. The hash is
    of what the registers held, so a later change shows against what was confirmed.
    """
    if actor.id is None:
        raise ValueError("a register is confirmed by a person")
    found = blockers(session, bid.id)
    if found.any:
        raise RegisterNotReady(found)

    drawings = [row for row in drawing_register(session, bid.id) if row.state == "current"]
    documents = [row for row in document_register(session, bid.id) if row.state == "current"]
    snapshot = json.dumps(
        {
            "drawings": [[row.sheet_number, row.revision, str(row.sheet_id)] for row in drawings],
            "documents": [[row.doc_key, row.revision, str(row.document_id)] for row in documents],
        },
        sort_keys=True,
    )
    confirmation = RegisterConfirmation(
        bid_id=bid.id,
        confirmed_by_id=actor.id,
        confirmed_role=role,
        confirmed_at=datetime.now(UTC),
        snapshot_hash=hashlib.sha256(snapshot.encode()).hexdigest(),
        drawings=len(drawings),
        documents=len(documents),
        comment=comment,
    )
    session.add(confirmation)
    session.flush()
    stage_before = bid.stage
    if bid.stage == "S1":
        # The registers are S1's output, so the work moves on to takeoff. From any other
        # stage it does not: S0 still needs its G0 decision, and a later stage is not undone.
        bid.stage = "S2"
    _audit(
        session,
        bid.id,
        actor,
        "register: confirmed",
        confirmation,
        {"stage": stage_before},
        {
            "stage": bid.stage,
            "drawings": len(drawings),
            "documents": len(documents),
            "snapshot_hash": confirmation.snapshot_hash,
        },
        reason=comment,
    )
    log.info("register_confirmed", bid_id=str(bid.id), drawings=len(drawings))
    return confirmation


def latest_confirmation(session: Session, bid_id: uuid.UUID) -> RegisterConfirmation | None:
    return session.execute(
        select(RegisterConfirmation)
        .where(RegisterConfirmation.bid_id == bid_id)
        .order_by(RegisterConfirmation.confirmed_at.desc())
        .limit(1)
    ).scalar_one_or_none()
