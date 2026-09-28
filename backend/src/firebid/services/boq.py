"""Bills of quantities for a bid (P1-09): the company's, the client's, and between them.

* **Templates** (FR-ADM-03) are the organisation's, versioned, seeded from
  `config/boq_templates.yaml` and marked "to be confirmed".
* **The company BOQ** (FR-BOQ-01) is built from verified QTO items by the current template.
  Each build is a new BOQ version; the last one is current. Every line keeps the items it
  came from (`boq_line_source`), and its stable key, so a client line mapped to it stays
  mapped when the BOQ is built again.
* **The client's BOQ** (FR-BOQ-02) is read in the parser sandbox from its original bytes,
  which are never written. An ambiguous layout waits for the model's column proposal and a
  person's confirmation.
* **Mapping** (FR-BOQ-02): the rules propose what they are sure of, the model the rest, and
  a person confirms or corrects each.
* **Reconciliation** (FR-BOQ-03): client against measured quantity, with the variance
  flagged over the threshold as a clarification candidate (P2-06).
* **Exports** (FR-BOQ-04): the client's workbook priced by patching only its rate cells;
  the company BOQ and the reconciliation as new workbooks.
* **Trace** (FR-BOQ-05): a line with no QTO items must be marked provisional or lump sum,
  or it blocks G1 and G2.
* **Conventions** (FR-BOQ-06): per tender, worded for the qualifications.
"""

from __future__ import annotations

import hashlib
import io
import json
import uuid
from dataclasses import dataclass
from datetime import UTC, datetime
from decimal import Decimal
from typing import Any

import structlog
from sqlalchemy import delete, func, select
from sqlalchemy.orm import Session

from firebid.boq import generate as generation
from firebid.boq import matching, reconcile
from firebid.boq.xlsx_patch import CellEdit, patch_workbook
from firebid.db.audit import record_event
from firebid.db.models.commercial import (
    Boq,
    BoqLine,
    BoqLineSource,
    BoqTemplate,
    ClientBoq,
    ClientBoqLine,
    ClientBoqMapping,
    MeasurementConvention,
)
from firebid.db.models.core import Bid
from firebid.db.models.documents import Document, DocumentRevision
from firebid.db.models.takeoff import Evidence, QtoItem
from firebid.db.models.workflow import Approval
from firebid.domain.actors import Actor, AuditContext
from firebid.domain.values import Money
from firebid.storage.object_store import ObjectStore

log = structlog.get_logger("firebid.boq")

VERIFIED = ("verified", "baselined")
BOQ_KINDS = ("xlsx",)


class BoqError(ValueError):
    """A BOQ request that cannot be done, with the reason a person can act on."""


def _context(session: Session, bid_id: uuid.UUID) -> AuditContext:
    organisation = session.execute(select(Bid.organisation_id).where(Bid.id == bid_id))
    return AuditContext(organisation_id=organisation.scalar_one(), bid_id=bid_id)


def _organisation_of(session: Session, bid_id: uuid.UUID) -> uuid.UUID:
    return session.execute(select(Bid.organisation_id).where(Bid.id == bid_id)).scalar_one()


# --- Templates (FR-ADM-03) ------------------------------------------------------------------


def template_rows(session: Session, organisation_id: uuid.UUID) -> dict[str, BoqTemplate]:
    """Each template's version in force; the seed's the first time."""
    rows = list(
        session.execute(
            select(BoqTemplate).where(
                BoqTemplate.organisation_id == organisation_id, BoqTemplate.retired_at.is_(None)
            )
        ).scalars()
    )
    if not rows:
        for seed in generation.seed_templates():
            row = BoqTemplate(
                organisation_id=organisation_id,
                key=seed.key,
                version=1,
                title=seed.title,
                definition=seed.definition,
                status=seed.status,
                change_note="seeded default (config/boq_templates.yaml)",
            )
            session.add(row)
            rows.append(row)
        session.flush()
    return {row.key: row for row in rows}


def template(session: Session, organisation_id: uuid.UUID, key: str | None = None) -> BoqTemplate:
    rows = template_rows(session, organisation_id)
    if key is None:
        return rows.get("company_standard") or next(iter(rows.values()))
    if key not in rows:
        raise BoqError(f"no BOQ template {key!r}")
    return rows[key]


def template_history(session: Session, organisation_id: uuid.UUID, key: str) -> list[BoqTemplate]:
    template_rows(session, organisation_id)
    return list(
        session.execute(
            select(BoqTemplate)
            .where(BoqTemplate.organisation_id == organisation_id, BoqTemplate.key == key)
            .order_by(BoqTemplate.version)
        ).scalars()
    )


def edit_template(
    session: Session,
    organisation_id: uuid.UUID,
    key: str,
    *,
    definition: dict[str, Any],
    actor: Actor,
    title: str | None = None,
    status: str = "confirmed",
    note: str | None = None,
) -> BoqTemplate:
    """A new version of a template. The old one is retired, never changed."""
    current = template(session, organisation_id, key)
    for field_name in ("groups", "descriptions", "sections"):
        if field_name not in definition:
            raise BoqError(f"a template needs {field_name!r}")
    current.retired_at = datetime.now(UTC)
    session.flush()
    row = BoqTemplate(
        organisation_id=organisation_id,
        key=key,
        version=current.version + 1,
        title=title or current.title,
        definition=definition,
        status=status,
        change_note=note,
        created_by_id=actor.id,
    )
    session.add(row)
    session.flush()
    record_event(
        session,
        context=AuditContext(organisation_id=organisation_id),
        actor=actor,
        action="BOQ template: new version",
        entity_type=BoqTemplate.__tablename__,
        entity_id=row.id,
        before={"key": key, "version": current.version},
        after={"key": key, "version": row.version, "status": status},
        reason=note,
    )
    return row


# --- The company BOQ (FR-BOQ-01) ------------------------------------------------------------


def line_key(draft: generation.LineDraft) -> str:
    material = json.dumps([draft.section, draft.group, draft.level, draft.description, draft.unit])
    return hashlib.sha256(material.encode()).hexdigest()[:32]


def _items(session: Session, bid_id: uuid.UUID) -> list[generation.Item]:
    rows = session.execute(
        select(QtoItem)
        .where(QtoItem.bid_id == bid_id, QtoItem.state.in_(VERIFIED))
        .order_by(QtoItem.human_id)
    ).scalars()
    out = []
    for row in rows:
        stated: dict[str, Any] = dict(row.attributes or {})
        out.append(
            generation.Item(
                id=str(row.id),
                human_id=row.human_id,
                item_type=row.item_type,
                classification=row.classification or row.item_type,
                description=row.description,
                attributes={
                    k: str(v.get("value") if isinstance(v, dict) else v) for k, v in stated.items()
                },
                unit=row.unit,
                net_quantity=row.net_quantity,
                allowance_percent=row.allowance_percent,
                level=row.level,
            )
        )
    return out


def current_boq(session: Session, bid_id: uuid.UUID) -> Boq | None:
    return (
        session.execute(
            select(Boq)
            .where(Boq.bid_id == bid_id, Boq.kind == "company", Boq.is_current.is_(True))
            .order_by(Boq.version.desc())
        )
        .scalars()
        .first()
    )


def lines_of(session: Session, boq: Boq) -> list[BoqLine]:
    return list(
        session.execute(
            select(BoqLine).where(BoqLine.boq_id == boq.id).order_by(BoqLine.sort_order)
        ).scalars()
    )


def build_company_boq(
    session: Session, bid_id: uuid.UUID, actor: Actor, template_key: str | None = None
) -> Boq:
    """A new version of the company BOQ from the verified takeoff; the old one is kept."""
    chosen = template(session, _organisation_of(session, bid_id), template_key)
    items = _items(session, bid_id)
    if not items:
        raise BoqError("nothing is verified yet: verify the takeoff first")
    drafts = generation.generate(
        items,
        generation.Template(
            chosen.key, chosen.version, chosen.title, dict(chosen.definition), chosen.status
        ),
    )
    previous = current_boq(session, bid_id)
    # Lines a person added by hand (provisional sums, lump sums) carry over.
    carried = (
        [line for line in lines_of(session, previous) if line.is_provisional or line.is_lump_sum]
        if previous
        else []
    )
    if previous is not None:
        previous.is_current = False
    boq = Boq(
        bid_id=bid_id,
        kind="company",
        template_key=chosen.key,
        template_version=chosen.version,
        version=(previous.version + 1) if previous else 1,
        is_current=True,
        created_by_id=actor.id,
    )
    session.add(boq)
    session.flush()
    for order, draft in enumerate(drafts):
        line = BoqLine(
            bid_id=bid_id,
            boq_id=boq.id,
            section=draft.section,
            item_no=draft.item_no,
            description=draft.description,
            unit=draft.unit,
            quantity=draft.quantity,
            sort_order=order,
            line_key=line_key(draft),
            group_heading=draft.group,
            level=draft.level,
            allowance_percent=draft.allowance_percent,
        )
        session.add(line)
        session.flush()
        for item_id in draft.item_ids:
            session.add(
                BoqLineSource(boq_line_id=line.id, qto_item_id=uuid.UUID(item_id), bid_id=bid_id)
            )
    for order, old in enumerate(carried, start=len(drafts)):
        session.add(
            BoqLine(
                bid_id=bid_id,
                boq_id=boq.id,
                section=old.section,
                item_no=old.item_no,
                description=old.description,
                unit=old.unit,
                quantity=old.quantity,
                sort_order=order,
                is_provisional=old.is_provisional,
                is_lump_sum=old.is_lump_sum,
                marker_note=old.marker_note,
                line_key=old.line_key,
                group_heading=old.group_heading,
                amount=old.amount,
            )
        )
    session.flush()
    _relink(session, bid_id, boq)
    record_event(
        session,
        context=_context(session, bid_id),
        actor=actor,
        action="BOQ: built from verified takeoff",
        entity_type=Boq.__tablename__,
        entity_id=boq.id,
        after={
            "version": boq.version,
            "template": f"{chosen.key}@{chosen.version}",
            "lines": len(drafts),
            "items": len(items),
        },
    )
    log.info("boq_built", bid_id=str(bid_id), version=boq.version, lines=len(drafts))
    return boq


def add_marked_line(
    session: Session,
    bid_id: uuid.UUID,
    actor: Actor,
    *,
    description: str,
    unit: str,
    quantity: Decimal,
    marker: str,
    note: str,
    amount: Decimal | None = None,
) -> BoqLine:
    """A provisional sum or lump sum: a line with no measured quantity behind it, marked so."""
    if marker not in ("provisional", "lump_sum"):
        raise BoqError("a line without a trace is 'provisional' or 'lump_sum'")
    if not note.strip():
        raise BoqError("say why the line has no measured quantity")
    boq = current_boq(session, bid_id)
    if boq is None:
        raise BoqError("build the BOQ first")
    count = session.execute(
        select(func.count()).select_from(BoqLine).where(BoqLine.boq_id == boq.id)
    ).scalar_one()
    line = BoqLine(
        bid_id=bid_id,
        boq_id=boq.id,
        section="PROVISIONAL SUMS" if marker == "provisional" else "LUMP SUMS",
        item_no=f"{'PS' if marker == 'provisional' else 'LS'}{count + 1}",
        description=description,
        unit=unit,
        quantity=quantity,
        sort_order=count,
        is_provisional=marker == "provisional",
        is_lump_sum=marker == "lump_sum",
        marker_note=note,
        amount=Money.of(amount) if amount is not None else None,
        line_key=hashlib.sha256(f"{marker}:{description}".encode()).hexdigest()[:32],
    )
    session.add(line)
    session.flush()
    record_event(
        session,
        context=_context(session, bid_id),
        actor=actor,
        action=f"BOQ: {marker.replace('_', ' ')} added",
        entity_type=BoqLine.__tablename__,
        entity_id=line.id,
        after={"description": description, "unit": unit, "quantity": str(quantity)},
        reason=note,
    )
    return line


def mark_line(
    session: Session, line: BoqLine, marker: str | None, actor: Actor, note: str | None
) -> BoqLine:
    """Mark (or unmark) a line as provisional or a lump sum."""
    if marker not in (None, "provisional", "lump_sum"):
        raise BoqError("mark a line 'provisional' or 'lump_sum', or clear it")
    if marker is not None and not (note or "").strip():
        raise BoqError("say why the line has no measured quantity")
    before = {"provisional": line.is_provisional, "lump_sum": line.is_lump_sum}
    line.is_provisional = marker == "provisional"
    line.is_lump_sum = marker == "lump_sum"
    line.marker_note = note if marker else None
    session.flush()
    record_event(
        session,
        context=_context(session, line.bid_id),
        actor=actor,
        action="BOQ line: marker",
        entity_type=BoqLine.__tablename__,
        entity_id=line.id,
        before=before,
        after={"provisional": line.is_provisional, "lump_sum": line.is_lump_sum},
        reason=note,
    )
    return line


# --- Trace (FR-BOQ-05) ----------------------------------------------------------------------


def untraced_lines(session: Session, bid_id: uuid.UUID) -> list[dict[str, Any]]:
    """Lines of the current BOQ with no QTO item behind them and no marker saying why."""
    boq = current_boq(session, bid_id)
    if boq is None:
        return []
    traced = {
        row
        for row in session.execute(
            select(BoqLineSource.boq_line_id)
            .join(Evidence, Evidence.qto_item_id == BoqLineSource.qto_item_id)
            .where(BoqLineSource.bid_id == bid_id)
        ).scalars()
    }
    return [
        {"id": str(line.id), "item_no": line.item_no, "description": line.description}
        for line in lines_of(session, boq)
        if line.id not in traced and not (line.is_provisional or line.is_lump_sum)
    ]


# --- Conventions (FR-BOQ-06) ----------------------------------------------------------------


def conventions_of(session: Session, bid_id: uuid.UUID) -> MeasurementConvention | None:
    return (
        session.execute(
            select(MeasurementConvention)
            .where(MeasurementConvention.bid_id == bid_id)
            .order_by(MeasurementConvention.version.desc())
        )
        .scalars()
        .first()
    )


def set_conventions(
    session: Session, bid_id: uuid.UUID, chosen: dict[str, str], actor: Actor
) -> MeasurementConvention:
    try:
        full = reconcile.validate(chosen)
    except ValueError as refusal:
        raise BoqError(str(refusal)) from refusal
    previous = conventions_of(session, bid_id)
    row = MeasurementConvention(
        bid_id=bid_id,
        settings=full,
        version=(previous.version + 1) if previous else 1,
        created_by_id=actor.id,
    )
    session.add(row)
    session.flush()
    record_event(
        session,
        context=_context(session, bid_id),
        actor=actor,
        action="measurement conventions: set",
        entity_type=MeasurementConvention.__tablename__,
        entity_id=row.id,
        before=dict(previous.settings) if previous else None,
        after=full,
    )
    return row


def qualification_text(session: Session, bid_id: uuid.UUID) -> str:
    row = conventions_of(session, bid_id)
    return reconcile.qualification_text(dict(row.settings) if row else {})


# --- The client's BOQ (FR-BOQ-02) -----------------------------------------------------------


def queue_reading(session: Session, document: Document) -> None:
    """After registration: read a client's BOQ workbook in the sandbox pool."""
    if document.doc_type != "boq" or document.kind not in BOQ_KINDS:
        return
    from firebid.jobs.enqueue import enqueue
    from firebid.jobs.tasks import read_client_boq_job

    enqueue(
        session,
        read_client_boq_job,
        document_id=str(document.id),
        user_id=str(document.created_by_id) if document.created_by_id else "",
    )


def _revision(session: Session, document: Document) -> DocumentRevision | None:
    return session.execute(
        select(DocumentRevision).where(DocumentRevision.document_id == document.id)
    ).scalar_one_or_none()


def read_client_boq(
    session: Session,
    store: ObjectStore,
    document: Document,
    *,
    confirmed: dict[str, Any] | None = None,
    user_id: str = "",
) -> list[ClientBoq]:
    """The client's bill, one row per sheet. Reading again replaces what an earlier read found."""
    from firebid.boq.reader import read_json
    from firebid.sandbox.runner import run_sandboxed

    found: dict[str, Any] = run_sandboxed(read_json, store.get(document.storage_key), confirmed)
    revision = _revision(session, document)
    existing = list(
        session.execute(select(ClientBoq).where(ClientBoq.document_id == document.id)).scalars()
    )
    for row in existing:
        session.execute(delete(ClientBoqLine).where(ClientBoqLine.client_boq_id == row.id))
        session.delete(row)
    session.flush()
    if "ambiguous" in found:
        row = ClientBoq(
            bid_id=document.bid_id,
            document_id=document.id,
            document_revision_id=revision.id if revision else None,
            status="needs_columns",
            reason=str(found["ambiguous"]),
            proposal={"preview": found["preview"]},
        )
        session.add(row)
        session.flush()
        _queue_columns(session, row, user_id)
        return [row]
    rows = []
    for sheet in found["sheets"]:
        row = ClientBoq(
            bid_id=document.bid_id,
            document_id=document.id,
            document_revision_id=revision.id if revision else None,
            sheet_name=sheet["name"],
            header_row=sheet["header_row"],
            column_map=sheet["columns"],
            status="read",
        )
        session.add(row)
        session.flush()
        for line in sheet["lines"]:
            session.add(
                ClientBoqLine(
                    bid_id=document.bid_id,
                    client_boq_id=row.id,
                    row_index=line["row"],
                    section=line["section"],
                    item_no=line["item_no"],
                    description=line["description"],
                    unit=line["unit"],
                    quantity=Decimal(str(line["quantity"]))
                    if line["quantity"] is not None
                    else None,
                    amount_is_formula=line["amount_is_formula"],
                    kind=line["kind"],
                    rate_cell=line["rate_cell"],
                    amount_cell=line["amount_cell"],
                )
            )
        rows.append(row)
    session.flush()
    log.info("client_boq_read", document_id=str(document.id), sheets=len(rows))
    return rows


def _queue_columns(session: Session, row: ClientBoq, user_id: str) -> None:
    from firebid.jobs.enqueue import enqueue
    from firebid.jobs.tasks import propose_boq_columns_job

    enqueue(session, propose_boq_columns_job, client_boq_id=str(row.id), user_id=user_id)


def propose_columns_with_model(session: Session, row: ClientBoq, router: Any) -> ClientBoq:
    """Ask the model which columns hold which field, sheet by sheet; a person confirms."""
    from firebid.agents.base import AgentInput
    from firebid.agents.boq_reader import BoqColumnReader, ColumnProposal, SheetPreview
    from firebid.agents.runtime import Escalated, run_agent_with_result

    stored: dict[str, Any] = dict(row.proposal or {})
    preview: dict[str, Any] = dict(stored.get("preview") or {})
    proposals: dict[str, Any] = {}
    for sheet, rows in preview.items():
        try:
            _, result = run_agent_with_result(
                session,
                BoqColumnReader(router),
                AgentInput(
                    bid_id=row.bid_id,
                    idempotency_key=f"boq-columns:{row.id}:{sheet}",
                    payload=SheetPreview(sheet=sheet, rows=rows),
                ),
            )
        except Escalated:
            continue
        if result is not None and isinstance(result.output, ColumnProposal):
            proposals[sheet] = result.output.model_dump()
    row.proposal = {"preview": preview, "columns": proposals}
    session.flush()
    return row


def confirm_columns(
    session: Session,
    store: ObjectStore,
    row: ClientBoq,
    columns: dict[str, dict[str, Any]],
    actor: Actor,
) -> list[ClientBoq]:
    """A person's header row and columns per sheet; the bill is read again with them."""
    document = session.get(Document, row.document_id)
    if document is None:
        raise BoqError("the workbook is gone")
    confirmed: dict[str, Any] = {}
    for sheet, choice in columns.items():
        fields = {k: int(v) for k, v in dict(choice.get("columns") or {}).items()}
        if not {"description", "quantity"} <= fields.keys():
            raise BoqError(f"{sheet}: say which columns hold the description and the quantity")
        confirmed[sheet] = [int(choice["header_row"]), fields]
    record_event(
        session,
        context=_context(session, row.bid_id),
        actor=actor,
        action="client BOQ: columns confirmed",
        entity_type=ClientBoq.__tablename__,
        entity_id=row.id,
        after={"columns": {s: list(c) for s, c in confirmed.items()}},
    )
    return read_client_boq(session, store, document, confirmed=confirmed)


def client_lines(session: Session, bid_id: uuid.UUID) -> list[tuple[ClientBoq, ClientBoqLine]]:
    return [
        (sheet, line)
        for sheet, line in session.execute(
            select(ClientBoq, ClientBoqLine)
            .join(ClientBoqLine, ClientBoqLine.client_boq_id == ClientBoq.id)
            .where(ClientBoq.bid_id == bid_id, ClientBoq.status == "read")
            .order_by(ClientBoq.sheet_name, ClientBoqLine.row_index)
        ).all()
    ]


def reference(sheet: ClientBoq, line: ClientBoqLine) -> str:
    return f"{sheet.sheet_name}!{line.row_index}"


# --- Mapping (FR-BOQ-02) --------------------------------------------------------------------


def mappings(session: Session, bid_id: uuid.UUID) -> dict[uuid.UUID, ClientBoqMapping]:
    return {
        row.client_boq_line_id: row
        for row in session.execute(
            select(ClientBoqMapping).where(ClientBoqMapping.bid_id == bid_id)
        ).scalars()
    }


def _relink(session: Session, bid_id: uuid.UUID, boq: Boq) -> None:
    """Point every mapping at the current BOQ's line with its key, and refresh its numbers."""
    by_key = {line.line_key: line for line in lines_of(session, boq) if line.line_key}
    lines = dict(
        session.execute(
            select(ClientBoqLine.id, ClientBoqLine.quantity).where(ClientBoqLine.bid_id == bid_id)
        )
        .tuples()
        .all()
    )
    for mapping in mappings(session, bid_id).values():
        target = by_key.get(mapping.boq_line_key) if mapping.boq_line_key else None
        mapping.boq_line_id = target.id if target else None
        _refresh(mapping, lines.get(mapping.client_boq_line_id), target)
    session.flush()


def _refresh(mapping: ClientBoqMapping, client: Decimal | None, target: BoqLine | None) -> None:
    measured = target.quantity if target else None
    mapping.measured_quantity = measured
    if target is None:
        mapping.variance_percent = None
        mapping.flagged = False
        return
    found = reconcile.variance(client, measured)
    mapping.variance_percent = float(found.percent) if found.percent is not None else None
    mapping.flagged = found.flagged


def propose_mappings(session: Session, bid_id: uuid.UUID, router: Any | None = None) -> int:
    """Rules first, then the model for what they left; a person's decisions are kept."""
    boq = current_boq(session, bid_id)
    if boq is None:
        raise BoqError("build the BOQ first")
    measured_lines = [
        line for line in lines_of(session, boq) if not (line.is_provisional or line.is_lump_sum)
    ]
    by_key = {line.line_key: line for line in measured_lines}
    existing = mappings(session, bid_id)
    clients = [
        (sheet, line)
        for sheet, line in client_lines(session, bid_id)
        if line.id not in existing or existing[line.id].state == "proposed"
    ]
    proposals, left = matching.match(
        [
            matching.Client(str(line.id), line.description or "", line.unit, line.kind)
            for _, line in clients
        ],
        [
            matching.Measured(line.line_key or "", line.description, line.unit)
            for line in measured_lines
        ],
    )
    answered: dict[str, tuple[str | None, float, str, str, dict[str, Any]]] = {
        p.ref: (p.maps_to, p.confidence, p.reason, "rule", {"rules": "boq.matching"})
        for p in proposals
    }
    if left and router is not None:
        answered.update(_ask_model(session, bid_id, clients, left, measured_lines, router))
    lookup = {str(line.id): (sheet, line) for sheet, line in clients}
    for ref, (key, confidence, reason, method, provenance) in answered.items():
        _, line = lookup[ref]
        row = existing.get(line.id)
        if row is None:
            row = ClientBoqMapping(bid_id=bid_id, client_boq_line_id=line.id)
            session.add(row)
        row.boq_line_key = key
        row.confidence = confidence
        row.reason = reason
        row.method = method
        row.state = "proposed"
        row.provenance = provenance
        target = by_key.get(key) if key else None
        row.boq_line_id = target.id if target else None
        _refresh(row, line.quantity, target)
    session.flush()
    return len(answered)


def _ask_model(
    session: Session,
    bid_id: uuid.UUID,
    clients: list[tuple[ClientBoq, ClientBoqLine]],
    left: list[matching.Client],
    measured: list[BoqLine],
    router: Any,
) -> dict[str, tuple[str | None, float, str, str, dict[str, Any]]]:
    from firebid.agents.base import AgentInput
    from firebid.agents.boq_reader import (
        BoqMapper,
        ClientLineIn,
        MappingAnswer,
        MappingBatch,
        MeasuredLineIn,
    )
    from firebid.agents.runtime import Escalated, run_agent_with_result

    by_id = {str(line.id): line for _, line in clients}
    measured_in = [
        MeasuredLineIn(
            key=line.line_key or "",
            description=line.description,
            unit=line.unit,
            quantity=str(line.quantity),
        )
        for line in measured
    ]
    out: dict[str, tuple[str | None, float, str, str, dict[str, Any]]] = {}
    batch_size = 40
    for start in range(0, len(left), batch_size):
        batch = left[start : start + batch_size]
        request = AgentInput(
            bid_id=bid_id,
            idempotency_key="boq-map:"
            + hashlib.sha256(
                json.dumps([c.ref for c in batch] + [m.key for m in measured_in]).encode()
            ).hexdigest()[:24],
            payload=MappingBatch(
                client=[
                    ClientLineIn(
                        ref=c.ref,
                        section=by_id[c.ref].section,
                        description=c.description,
                        unit=c.unit,
                        quantity=str(by_id[c.ref].quantity)
                        if by_id[c.ref].quantity is not None
                        else None,
                    )
                    for c in batch
                ],
                measured=measured_in,
            ),
        )
        try:
            run, result = run_agent_with_result(session, BoqMapper(router), request)
        except Escalated:
            continue
        if result is None or not isinstance(result.output, MappingAnswer):
            continue
        for mapping in result.output.mappings:
            out[mapping.line] = (
                mapping.maps_to,
                mapping.confidence,
                mapping.reason,
                "model",
                {
                    "agent_run_id": str(run.id),
                    "model": run.model,
                    "prompt_version": run.prompt_version,
                },
            )
    return out


def decide_mapping(
    session: Session,
    row: ClientBoqMapping,
    actor: Actor,
    *,
    decision: str,
    boq_line_key: str | None = None,
    note: str | None = None,
) -> ClientBoqMapping:
    """`confirm` the proposal, `correct` it to another line (or to nothing), or `reject` it."""
    if decision not in ("confirm", "correct", "reject"):
        raise BoqError("decide 'confirm', 'correct' or 'reject'")
    before = {"state": row.state, "maps_to": row.boq_line_key, "method": row.method}
    boq = current_boq(session, row.bid_id)
    if decision == "correct":
        if boq_line_key is not None and boq is not None:
            keys = {line.line_key for line in lines_of(session, boq)}
            if boq_line_key not in keys:
                raise BoqError("no such line in the current BOQ")
        row.boq_line_key = boq_line_key
        row.method = "person"
        row.confidence = 1.0
        row.reason = note or "corrected by a person"
    row.state = "rejected" if decision == "reject" else "confirmed"
    row.confirmed_by_id = actor.id
    row.confirmed_at = datetime.now(UTC)
    target = (
        next((line for line in lines_of(session, boq) if line.line_key == row.boq_line_key), None)
        if boq is not None and row.boq_line_key
        else None
    )
    row.boq_line_id = target.id if target else None
    client = session.get(ClientBoqLine, row.client_boq_line_id)
    _refresh(row, client.quantity if client else None, target)
    session.flush()
    record_event(
        session,
        context=_context(session, row.bid_id),
        actor=actor,
        action=f"client BOQ mapping: {decision}",
        entity_type=ClientBoqMapping.__tablename__,
        entity_id=row.id,
        before=before,
        after={"state": row.state, "maps_to": row.boq_line_key, "method": row.method},
        reason=note,
    )
    return row


# --- Reconciliation (FR-BOQ-03) -------------------------------------------------------------


@dataclass
class ReconciliationRow:
    kind: str  # mapped | client_only | measured_only
    client_ref: str | None
    client_item: str | None
    client_description: str | None
    client_unit: str | None
    client_quantity: Decimal | None
    line_item: str | None
    line_description: str | None
    unit: str | None
    measured_quantity: Decimal | None
    variance: Decimal | None
    variance_percent: Decimal | None
    flagged: bool
    state: str | None
    qto_items: list[str]
    evidence_links: list[str]


def reconciliation(session: Session, bid_id: uuid.UUID) -> list[ReconciliationRow]:
    """Every client line against what was measured, and what was measured the client left out."""
    boq = current_boq(session, bid_id)
    lines = lines_of(session, boq) if boq else []
    by_key = {line.line_key: line for line in lines}
    sources: dict[uuid.UUID, list[str]] = {}
    humans = {
        row.id: row.human_id
        for row in session.execute(select(QtoItem).where(QtoItem.bid_id == bid_id)).scalars()
    }
    for source in session.execute(
        select(BoqLineSource).where(BoqLineSource.bid_id == bid_id)
    ).scalars():
        sources.setdefault(source.boq_line_id, []).append(str(source.qto_item_id))
    existing = mappings(session, bid_id)
    rows: list[ReconciliationRow] = []
    mapped_keys: set[str] = set()

    def trace(line: BoqLine | None) -> tuple[list[str], list[str]]:
        if line is None:
            return [], []
        ids = sources.get(line.id, [])
        return (
            sorted(humans.get(uuid.UUID(i), i) for i in ids),
            [f"/bids/{bid_id}/qto/items/{i}/evidence" for i in ids],
        )

    for sheet, client in client_lines(session, bid_id):
        mapping = existing.get(client.id)
        target = by_key.get(mapping.boq_line_key) if mapping and mapping.boq_line_key else None
        if mapping and mapping.state == "rejected":
            target = None
        if target is not None and target.line_key:
            mapped_keys.add(target.line_key)
        found = reconcile.variance(client.quantity, target.quantity) if target is not None else None
        items, links = trace(target)
        rows.append(
            ReconciliationRow(
                kind="mapped" if target is not None else "client_only",
                client_ref=reference(sheet, client),
                client_item=client.item_no,
                client_description=client.description,
                client_unit=client.unit,
                client_quantity=client.quantity,
                line_item=target.item_no if target else None,
                line_description=target.description if target else None,
                unit=target.unit if target else None,
                measured_quantity=target.quantity if target else None,
                variance=found.difference if found else None,
                variance_percent=found.percent if found else None,
                # A client's line with nothing measured behind it is a question for the
                # client, unless it is a sum they allow with no quantity to measure.
                flagged=bool(found.flagged) if found else client.kind == "line",
                state=mapping.state if mapping else None,
                qto_items=items,
                evidence_links=links,
            )
        )
    for line in lines:
        if line.line_key in mapped_keys or line.is_provisional or line.is_lump_sum:
            continue
        items, links = trace(line)
        rows.append(
            ReconciliationRow(
                kind="measured_only",
                client_ref=None,
                client_item=None,
                client_description=None,
                client_unit=None,
                client_quantity=None,
                line_item=line.item_no,
                line_description=line.description,
                unit=line.unit,
                measured_quantity=line.quantity,
                variance=None,
                variance_percent=None,
                flagged=True,
                state=None,
                qto_items=items,
                evidence_links=links,
            )
        )
    return rows


def clarification_candidates(session: Session, bid_id: uuid.UUID) -> list[ReconciliationRow]:
    """What P2-06 raises clarifications about: variances over the threshold, and items one
    side has that the other does not."""
    return [row for row in reconciliation(session, bid_id) if row.flagged]


# --- Exports (FR-BOQ-04) --------------------------------------------------------------------


def priced_client_workbook(
    session: Session, store: ObjectStore, bid_id: uuid.UUID, document_id: uuid.UUID
) -> bytes:
    """The client's own workbook with rates written into their cells, nothing else changed.

    A rate comes from the mapped company line's `unit_rate` (P1-10 prices it). A plain-value
    amount cell is written as quantity x rate; a formula amount is left for Excel.
    """
    document = session.get(Document, document_id)
    if document is None or document.bid_id != bid_id:
        raise BoqError("no such workbook on this bid")
    boq = current_boq(session, bid_id)
    by_key = {line.line_key: line for line in lines_of(session, boq)} if boq else {}
    existing = mappings(session, bid_id)
    edits: list[CellEdit] = []
    for sheet, client in client_lines(session, bid_id):
        if sheet.document_id != document.id:
            continue
        mapping = existing.get(client.id)
        if mapping is None or mapping.state != "confirmed" or not mapping.boq_line_key:
            continue
        line = by_key.get(mapping.boq_line_key)
        if line is None or line.unit_rate is None or not client.rate_cell:
            continue
        rate = line.unit_rate.amount
        edits.append(CellEdit(sheet.sheet_name or "", client.rate_cell, rate))
        if client.amount_cell and not client.amount_is_formula and client.quantity is not None:
            edits.append(
                CellEdit(
                    sheet.sheet_name or "",
                    client.amount_cell,
                    (rate * client.quantity).quantize(Decimal("0.01")),
                )
            )
    original = store.get(document.storage_key)
    return patch_workbook(original, edits) if edits else original


def company_workbook(session: Session, bid_id: uuid.UUID) -> bytes:
    """The company BOQ as a new workbook (openpyxl may write this one: it is ours)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    boq = current_boq(session, bid_id)
    if boq is None:
        raise BoqError("build the BOQ first")
    book = Workbook()
    sheet = book.active
    if sheet is None:  # pragma: no cover - a new workbook always has one
        sheet = book.create_sheet()
    sheet.title = "BOQ"
    sheet.append(
        ["Item", "Description", "Unit", "Net qty", "Allowance %", "Rate", "Amount", "Trace"]
    )
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    section = None
    for line in lines_of(session, boq):
        if line.section != section:
            section = line.section
            sheet.append([None, section])
            sheet.cell(row=sheet.max_row, column=2).font = Font(bold=True)
        marker = (
            "provisional sum" if line.is_provisional else "lump sum" if line.is_lump_sum else ""
        )
        sheet.append(
            [
                line.item_no,
                line.description + (f" ({line.level})" if line.level else ""),
                line.unit,
                float(line.quantity),
                float(line.allowance_percent) if line.allowance_percent is not None else None,
                float(line.unit_rate.amount) if line.unit_rate else None,
                float(line.amount.amount) if line.amount else None,
                marker or "QTO",
            ]
        )
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def reconciliation_workbook(session: Session, bid_id: uuid.UUID) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    book = Workbook()
    sheet = book.active
    if sheet is None:  # pragma: no cover
        sheet = book.create_sheet()
    sheet.title = "Reconciliation"
    sheet.append(
        [
            "Client ref",
            "Client item",
            "Client description",
            "Client unit",
            "Client qty",
            "Our item",
            "Our description",
            "Unit",
            "Measured qty",
            "Variance",
            "Variance %",
            "Clarify",
            "QTO items",
            "Evidence",
        ]
    )
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    flag = PatternFill("solid", fgColor="FDE68A")
    for row in reconciliation(session, bid_id):
        sheet.append(
            [
                row.client_ref,
                row.client_item,
                row.client_description,
                row.client_unit,
                float(row.client_quantity) if row.client_quantity is not None else None,
                row.line_item,
                row.line_description,
                row.unit,
                float(row.measured_quantity) if row.measured_quantity is not None else None,
                float(row.variance) if row.variance is not None else None,
                float(row.variance_percent) if row.variance_percent is not None else None,
                "yes" if row.flagged else "",
                ", ".join(row.qto_items),
                "\n".join(row.evidence_links),
            ]
        )
        if row.flagged:
            for cell in sheet[sheet.max_row]:
                cell.fill = flag
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()



# --- G2 (the BOQ behind the estimate) -------------------------------------------------------


@dataclass
class G2Blockers:
    g1_approved: bool
    boq_built: bool
    untraced_lines: list[dict[str, Any]]

    @property
    def clear(self) -> bool:
        return self.g1_approved and self.boq_built and not self.untraced_lines

    def describe(self) -> str:
        parts = []
        if not self.g1_approved:
            parts.append("G1 is not approved")
        if not self.boq_built:
            parts.append("no BOQ is built")
        if self.untraced_lines:
            parts.append(
                f"{len(self.untraced_lines)} BOQ line(s) with no QTO trace, "
                "not marked provisional or lump sum"
            )
        return "; ".join(parts)


def g2_blockers(session: Session, bid_id: uuid.UUID) -> G2Blockers:
    """What stands between the bid and G2 from the BOQ's side. Pricing (P1-10 on) adds its
    own checks; this step's is the trace (FR-BOQ-05)."""
    g1 = session.execute(
        select(func.count())
        .select_from(Approval)
        .where(Approval.bid_id == bid_id, Approval.gate == "G1", Approval.decision == "approved")
    ).scalar_one()
    return G2Blockers(
        g1_approved=g1 > 0,
        boq_built=current_boq(session, bid_id) is not None,
        untraced_lines=untraced_lines(session, bid_id),
    )


def boq_snapshot_hash(lines: list[BoqLine]) -> str:
    material = sorted(
        f"{line.line_key}|{line.quantity}|{line.unit}|{line.is_provisional}|{line.is_lump_sum}"
        for line in lines
    )
    return hashlib.sha256("\n".join(material).encode()).hexdigest()


def approve_g2(
    session: Session, bid: Bid, actor: Actor, role: str, comment: str | None = None
) -> Approval:
    """Record G2 on the BOQ as it stands, or raise `BoqError` saying what blocks it."""
    blockers = g2_blockers(session, bid.id)
    if not blockers.clear:
        raise BoqError("G2 is blocked: " + blockers.describe())
    if actor.id is None:
        raise BoqError("a gate is approved by a named person")
    boq = current_boq(session, bid.id)
    if boq is None:  # pragma: no cover - the blockers say one is built
        raise BoqError("no BOQ is built")
    approval = Approval(
        bid_id=bid.id,
        gate="G2",
        decision="approved",
        approver_id=actor.id,
        approver_role=role,
        decided_at=datetime.now(UTC),
        comment=comment,
        snapshot_hash=boq_snapshot_hash(lines_of(session, boq)),
    )
    session.add(approval)
    session.flush()
    record_event(
        session,
        context=AuditContext(organisation_id=bid.organisation_id, bid_id=bid.id),
        actor=actor,
        action="gate G2: approve",
        entity_type=Approval.__tablename__,
        entity_id=approval.id,
        after={
            "gate": "G2",
            "decision": "approved",
            "boq_version": boq.version,
            "snapshot_hash": approval.snapshot_hash,
        },
        reason=comment,
    )
    return approval
