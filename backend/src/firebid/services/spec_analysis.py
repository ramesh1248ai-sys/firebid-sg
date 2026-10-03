"""Full specification analysis for a bid, through the database (P2-03).

* **Obligations** (FR-SPEC-02) are read from a specification's stored clauses by rule, each
  with its clause, its words and whether the citation holds. They are proposals: a person
  confirms or rejects each. Reading a revision twice changes nothing.
* **Issues** (FR-SPEC-03) are the cross-check of every Current specification against the
  notes on the Current sheets and the takeoff (`specs.crosscheck`). An issue is kept by what
  it is, so running the analysis again leaves a person's dismissal in place; an issue no
  longer found is marked resolved. Open issues are the clarification candidates (P2-06).
* **The scope matrix** (FR-SPEC-04) is a row per obligation and interface per system. A
  row a person has changed is theirs: the analysis updates the clause it cites, never the
  status they set. The matrix is confirmed by a named person and exports to a workbook.

Nothing here opens a tender file: it works from the stored clauses, attributes, sheet text
and takeoff.
"""

from __future__ import annotations

import hashlib
import io
import uuid
from dataclasses import dataclass
from datetime import UTC, datetime
from typing import Any

import structlog
from sqlalchemy import select
from sqlalchemy.orm import Session

from firebid.db.audit import record_event
from firebid.db.models.core import Bid
from firebid.db.models.documents import DocumentRevision
from firebid.db.models.specs import ScopeRow, SpecClause, SpecIssue, SpecObligation
from firebid.domain.actors import Actor, AuditContext
from firebid.services import qto
from firebid.services import specs as spec_service
from firebid.specs import crosscheck, obligations, scope_matrix
from firebid.specs.clauses import Clause
from firebid.specs.crosscheck import Note, SheetRef, SpecRef, SpecValue

log = structlog.get_logger("firebid.spec_analysis")

SEVERITIES = {"high": 0, "medium": 1, "low": 2}


class AnalysisError(ValueError):
    """A request that cannot be done, with the reason a person can act on."""


def _context(session: Session, bid_id: uuid.UUID) -> AuditContext:
    organisation_id = session.execute(
        select(Bid.organisation_id).where(Bid.id == bid_id)
    ).scalar_one()
    return AuditContext(organisation_id=organisation_id, bid_id=bid_id)


def current_specifications(session: Session, bid_id: uuid.UUID) -> list[DocumentRevision]:
    """The Current specification revisions that have been read into clauses."""
    read = set(
        session.execute(
            select(SpecClause.document_revision_id).where(SpecClause.bid_id == bid_id).distinct()
        ).scalars()
    )
    return [
        revision
        for revision in session.execute(
            select(DocumentRevision)
            .where(DocumentRevision.bid_id == bid_id, DocumentRevision.state == "current")
            .order_by(DocumentRevision.doc_key)
        ).scalars()
        if revision.id in read
    ]


def _clause_ids(session: Session, revision_id: uuid.UUID) -> dict[str, uuid.UUID]:
    return {
        number: clause_id
        for number, clause_id in session.execute(
            select(SpecClause.number, SpecClause.id).where(
                SpecClause.document_revision_id == revision_id
            )
        )
    }


# --- Obligations (FR-SPEC-02) ----------------------------------------------------------------


def _obligation_key(item: obligations.Obligation) -> str:
    material = f"{item.clause}|{item.category}|{' '.join(item.quote.split())}"
    return hashlib.sha256(material.encode()).hexdigest()[:32]


def store_obligations(
    session: Session,
    revision: DocumentRevision,
    clauses: list[Clause],
    found: list[obligations.Obligation],
    provenance: dict[str, Any],
) -> list[SpecObligation]:
    """Store obligations not already held for the revision, each with its citation checked."""
    by_number = {clause.number: clause for clause in clauses}
    ids = _clause_ids(session, revision.id)
    held = set(
        session.execute(
            select(SpecObligation.key).where(SpecObligation.document_revision_id == revision.id)
        ).scalars()
    )
    stored = []
    for item in found:
        key = _obligation_key(item)
        if key in held:
            continue
        held.add(key)
        holds, reason = obligations.quote_holds(by_number.get(item.clause), item.quote)
        row = SpecObligation(
            bid_id=revision.bid_id,
            document_revision_id=revision.id,
            key=key,
            clause_id=ids.get(item.clause),
            clause_number=item.clause,
            system=item.system,
            category=item.category,
            summary=item.summary,
            quantities=dict(item.quantities),
            quote=item.quote,
            method=item.method,
            # A citation that does not hold leaves the obligation for a person, marked low.
            confidence=round(item.confidence if holds else min(item.confidence, 0.2), 4),
            citation_ok=holds,
            citation_reason=reason,
            state="proposed",
            provenance=dict(provenance),
        )
        session.add(row)
        stored.append(row)
    session.flush()
    return stored


def read_obligations(session: Session, revision: DocumentRevision) -> list[SpecObligation]:
    """The rules' obligations for one specification revision, from its stored clauses."""
    clauses = spec_service.clauses_of(session, revision.id)
    placed = spec_service.systems_of(session, revision.id)
    found = obligations.extract(clauses, placed)
    stored = store_obligations(
        session, revision, clauses, found, {"rule_version": obligations.RULES_VERSION}
    )
    if stored:
        log.info("obligations_read", revision_id=str(revision.id), obligations=len(stored))
    return stored


def read_obligations_with_model(
    session: Session, revision_id: uuid.UUID, router: Any
) -> list[SpecObligation]:
    """Ask the model about fire protection clauses that oblige something the rules put in
    no category. Its answers are checked against their clauses like the rules' are."""
    from firebid.agents.base import AgentInput
    from firebid.agents.runtime import Escalated, run_agent_with_result
    from firebid.agents.spec_reader import (
        ObligationInput,
        SpecObligationExtractor,
        SpecObligations,
    )

    revision = session.get(DocumentRevision, revision_id)
    if revision is None:
        return []
    clauses = spec_service.clauses_of(session, revision_id)
    placed = spec_service.systems_of(session, revision_id)
    known = [
        obligations.Obligation(row.system, row.category, row.summary, row.clause_number, row.quote)
        for row in session.execute(
            select(SpecObligation).where(SpecObligation.document_revision_id == revision_id)
        ).scalars()
    ]
    left = obligations.unread(clauses, placed, known)
    if not left:
        return []
    numbers = [clause.number for clause in left]
    request = AgentInput(
        bid_id=revision.bid_id,
        # The clauses asked about, as a digest: a specification has too many for a key.
        idempotency_key=(
            f"spec_obligations:{revision_id}:"
            + hashlib.sha256(",".join(numbers).encode()).hexdigest()[:24]
        ),
        payload=ObligationInput(
            clauses=[(c.number, c.heading, c.text) for c in clauses], clause_numbers=numbers
        ),
    )
    try:
        run, result = run_agent_with_result(session, SpecObligationExtractor(router), request)
    except Escalated:
        return []
    if result is None or not isinstance(result.output, SpecObligations):
        return []
    found = [
        obligations.Obligation(
            system=placed.get(item.clause, "fire_protection"),
            category=item.category,
            summary=item.summary,
            clause=item.clause,
            quote=item.quote,
            quantities=dict(item.quantities),
            method="model",
            confidence=item.confidence,
        )
        for item in result.output.obligations
    ]
    return store_obligations(
        session,
        revision,
        clauses,
        found,
        {
            "provider": run.provider,
            "model": run.model,
            "prompt_version": run.prompt_version,
            "agent_run_id": str(run.id),
        },
    )


def list_obligations(session: Session, bid_id: uuid.UUID) -> list[SpecObligation]:
    """The obligations of the bid's Current specifications, in reading order."""
    current = {revision.id for revision in current_specifications(session, bid_id)}
    rows = [
        row
        for row in session.execute(
            select(SpecObligation).where(SpecObligation.bid_id == bid_id)
        ).scalars()
        if row.document_revision_id in current
    ]
    return sorted(rows, key=lambda r: (_order(r.clause_number), r.category))


def _order(number: str) -> tuple[int, ...]:
    return tuple(int(part) if part.isdigit() else 0 for part in number.split("."))


def decide_obligation(
    session: Session, row: SpecObligation, decision: str, actor: Actor, note: str | None = None
) -> SpecObligation:
    """A person confirms or rejects an obligation, or puts it back to proposed."""
    states = {"confirm": "verified", "reject": "rejected", "reopen": "proposed"}
    if decision not in states:
        raise AnalysisError("an obligation is confirmed, rejected or reopened")
    before = row.state
    row.state = states[decision]
    row.decided_by = actor.label if decision != "reopen" else None
    row.decided_by_id = actor.id if decision != "reopen" else None
    row.decided_at = datetime.now(UTC) if decision != "reopen" else None
    row.note = note
    session.flush()
    record_event(
        session,
        context=_context(session, row.bid_id),
        actor=actor,
        action=f"specification obligation: {decision}",
        entity_type=SpecObligation.__tablename__,
        entity_id=row.id,
        before={"state": before},
        after={"state": row.state, "category": row.category, "clause": row.clause_number},
        reason=note,
    )
    return row


# --- Issues (FR-SPEC-03) ---------------------------------------------------------------------


def _inputs(
    session: Session, bid_id: uuid.UUID
) -> tuple[list[Note], dict[str, list[SheetRef]], list[SheetRef]]:
    """What the drawings say: the notes on Current sheets, what the takeoff counts and on
    which sheets, and every Current sheet looked at."""
    sheets = qto.current_sheets(session, bid_id)
    refs = {
        sheet_id: SheetRef(
            sheet_id=str(sheet_id),
            sheet_number=info.number,
            revision=info.revision.revision_label or "",
            title=info.revision.title or "",
            level=info.level,
        )
        for sheet_id, info in sheets.items()
    }
    notes = [
        Note(refs[sheet_id], span["text"], span["minx"], span["miny"])
        for sheet_id, spans in qto.sheet_texts(session, bid_id, sheets).items()
        for span in spans
        # A note is a sentence, not a size label or a tag: nothing shorter states a material.
        if len(span["text"]) >= 12
    ]
    by_id = {str(sheet_id): ref for sheet_id, ref in refs.items()}
    drawn: dict[str, list[SheetRef]] = {}
    for item in qto.live_items(session, bid_id):
        if item.state == "rejected":
            continue
        derivation: dict[str, Any] = dict(item.derivation or {})
        for source in derivation.get("sources") or []:
            ref = by_id.get(str(source.get("sheet_id")))
            if ref is not None and ref not in drawn.setdefault(item.item_type, []):
                drawn[item.item_type].append(ref)
    ordered = sorted(refs.values(), key=lambda ref: ref.sheet_number)
    return notes, drawn, ordered


def _values(session: Session, bid_id: uuid.UUID, revision_id: uuid.UUID) -> list[SpecValue]:
    """The revision's attributes as they stand, a person's rejections left out."""
    return [
        SpecValue(
            system=row.system,
            attribute=row.attribute,
            value=row.value,
            clause=row.clause_number,
            quote=row.quote,
            dn_min=row.dn_min,
            dn_max=row.dn_max,
            condition=row.condition,
        )
        for row in spec_service.current_attributes(session, bid_id)
        if row.document_revision_id == revision_id and row.state != "rejected"
    ]


def find_issues(session: Session, bid_id: uuid.UUID) -> list[crosscheck.Issue]:
    """Every issue between the bid's Current specifications and its drawings."""
    notes, drawn, sheets = _inputs(session, bid_id)
    per_revision = []
    for revision in current_specifications(session, bid_id):
        reference = SpecRef(
            document_id=str(revision.document_id),
            revision_id=str(revision.id),
            revision=revision.revision_label or "",
            title=revision.title or revision.doc_key or "",
        )
        found = crosscheck.check(
            spec_service.clauses_of(session, revision.id),
            spec_service.systems_of(session, revision.id),
            _values(session, bid_id, revision.id),
            notes,
            drawn,
            sheets,
            reference,
        )
        # The clause's own row, so a reader opens it in one click.
        ids = _clause_ids(session, revision.id)
        for issue in found:
            clause = issue.spec.get("clause")
            if clause and clause in ids:
                issue.spec["clause_id"] = str(ids[clause])
        per_revision.append(found)
    # An item one specification never mentions is missing only if none of them mentions it.
    silent = [
        {i.detail["what"] for i in found if i.rule == "missing_from_specification"}
        for found in per_revision
    ]
    nowhere = set.intersection(*silent) if silent else set()
    issues: dict[str, crosscheck.Issue] = {}
    for found in per_revision:
        for issue in found:
            if issue.rule == "missing_from_specification" and issue.detail["what"] not in nowhere:
                continue
            issues.setdefault(issue.key, issue)
    return sorted(issues.values(), key=lambda i: (SEVERITIES[i.severity], i.category, i.title))


@dataclass
class Outcome:
    obligations: int = 0
    issues_open: int = 0
    issues_new: int = 0
    issues_resolved: int = 0
    rows: int = 0


def analyse(session: Session, bid_id: uuid.UUID) -> Outcome:
    """Read obligations, cross-check and build the scope matrix. The same inputs change
    nothing, and nothing a person decided is undone."""
    outcome = Outcome()
    for revision in current_specifications(session, bid_id):
        read_obligations(session, revision)
    outcome.obligations = len(list_obligations(session, bid_id))

    now = datetime.now(UTC)
    held = {
        row.key: row
        for row in session.execute(select(SpecIssue).where(SpecIssue.bid_id == bid_id)).scalars()
    }
    seen = set()
    for issue in find_issues(session, bid_id):
        seen.add(issue.key)
        row = held.get(issue.key)
        if row is None:
            session.add(
                SpecIssue(
                    bid_id=bid_id,
                    key=issue.key,
                    category=issue.category,
                    rule=issue.rule,
                    severity=issue.severity,
                    title=issue.title,
                    system=issue.system,
                    spec_ref=issue.spec,
                    drawing_ref=issue.drawing,
                    detail=issue.detail,
                    rules_version=crosscheck.RULES_VERSION,
                    state="open",
                    last_found_at=now,
                )
            )
            outcome.issues_new += 1
            continue
        row.last_found_at = now
        if row.state == "resolved":
            row.state = "open"  # it is back
        # The citations follow the documents: a new revision of a sheet is cited as it is.
        if dict(row.spec_ref) != issue.spec:
            row.spec_ref = issue.spec
        if dict(row.drawing_ref) != issue.drawing:
            row.drawing_ref = issue.drawing
    for key, row in held.items():
        if key not in seen and row.state == "open":
            row.state = "resolved"
            outcome.issues_resolved += 1
    session.flush()
    outcome.issues_open = len(list_issues(session, bid_id, "open"))
    outcome.rows = len(_build_matrix(session, bid_id))
    log.info(
        "specification_analysed",
        bid_id=str(bid_id),
        obligations=outcome.obligations,
        issues_open=outcome.issues_open,
        rows=outcome.rows,
    )
    return outcome


def list_issues(session: Session, bid_id: uuid.UUID, state: str | None = None) -> list[SpecIssue]:
    query = select(SpecIssue).where(SpecIssue.bid_id == bid_id)
    if state:
        query = query.where(SpecIssue.state == state)
    return sorted(
        session.execute(query).scalars(),
        key=lambda r: (SEVERITIES.get(r.severity, 9), r.category, r.title),
    )


def clarification_candidates(session: Session, bid_id: uuid.UUID) -> list[SpecIssue]:
    """The issues to raise with the consultant: every one still open (for P2-06)."""
    return list_issues(session, bid_id, "open")


def decide_issue(
    session: Session, row: SpecIssue, decision: str, actor: Actor, note: str | None = None
) -> SpecIssue:
    """A person dismisses an issue (it is not one, with the reason), or reopens it."""
    if decision not in ("dismiss", "reopen"):
        raise AnalysisError("an issue is dismissed or reopened")
    if decision == "dismiss" and not (note or "").strip():
        raise AnalysisError("say why this is not an issue")
    before = row.state
    row.state = "dismissed" if decision == "dismiss" else "open"
    row.decided_by = actor.label if decision == "dismiss" else None
    row.decided_by_id = actor.id if decision == "dismiss" else None
    row.decided_at = datetime.now(UTC) if decision == "dismiss" else None
    row.note = note
    session.flush()
    record_event(
        session,
        context=_context(session, row.bid_id),
        actor=actor,
        action=f"specification issue: {decision}",
        entity_type=SpecIssue.__tablename__,
        entity_id=row.id,
        before={"state": before},
        after={"state": row.state, "rule": row.rule, "title": row.title},
        reason=note,
    )
    return row


# --- The scope matrix (FR-SPEC-04) -----------------------------------------------------------


def _build_matrix(session: Session, bid_id: uuid.UUID) -> list[ScopeRow]:
    held = {
        (row.system, row.kind, row.key): row
        for row in session.execute(select(ScopeRow).where(ScopeRow.bid_id == bid_id)).scalars()
    }
    for revision in current_specifications(session, bid_id):
        clauses = spec_service.clauses_of(session, revision.id)
        placed = spec_service.systems_of(session, revision.id)
        ids = _clause_ids(session, revision.id)
        stated = [
            obligations.Obligation(
                row.system, row.category, row.summary, row.clause_number, row.quote
            )
            for row in session.execute(
                select(SpecObligation).where(
                    SpecObligation.document_revision_id == revision.id,
                    SpecObligation.state != "rejected",
                )
            ).scalars()
        ]
        for found in scope_matrix.build(clauses, placed, stated):
            key = (found.system, found.kind, found.key)
            row = held.get(key)
            if row is None:
                row = ScopeRow(
                    bid_id=bid_id,
                    system=found.system,
                    kind=found.kind,
                    key=found.key,
                    label=found.label,
                    status=found.status,
                    proposed_status=found.status,
                    document_revision_id=revision.id,
                    clause_id=ids.get(found.clause) if found.clause else None,
                    clause_number=found.clause,
                    quote=found.quote,
                    reason=found.reason,
                    source="rule",
                )
                session.add(row)
                held[key] = row
                continue
            # A second specification that mentions it speaks where the first was silent.
            if row.clause_number is not None and found.clause is None:
                continue
            row.proposed_status = found.status
            row.document_revision_id = revision.id
            row.clause_id = ids.get(found.clause) if found.clause else None
            row.clause_number = found.clause
            row.quote = found.quote
            row.reason = found.reason
            if row.source == "rule" and row.status != found.status:
                row.status = found.status
                row.confirmed_by = row.confirmed_by_id = row.confirmed_at = None
    session.flush()
    return matrix(session, bid_id)


def matrix(session: Session, bid_id: uuid.UUID) -> list[ScopeRow]:
    """The matrix as it stands: by system, obligations first, then the interfaces in the
    order the company lists them."""
    rows = session.execute(select(ScopeRow).where(ScopeRow.bid_id == bid_id)).scalars()
    order = {key: index for index, key in enumerate(obligations.CATEGORY_KEYS)}
    order.update(
        {item.key: len(order) + index for index, item in enumerate(scope_matrix.interfaces())}
    )
    return sorted(
        rows,
        key=lambda r: (r.system, r.kind != "obligation", order.get(r.key, len(order)), r.label),
    )


def edit_row(
    session: Session, row: ScopeRow, status: str, actor: Actor, note: str | None = None
) -> ScopeRow:
    """A person sets a row's status. The row is theirs from then on, and unconfirmed."""
    if status not in scope_matrix.STATUSES:
        raise AnalysisError(f"a status is one of {', '.join(scope_matrix.STATUSES)}")
    before = {"status": row.status, "source": row.source}
    row.status = status
    row.source = "person"
    row.note = note
    row.edited_by = actor.label
    row.edited_at = datetime.now(UTC)
    row.confirmed_by = row.confirmed_by_id = row.confirmed_at = None
    session.flush()
    record_event(
        session,
        context=_context(session, row.bid_id),
        actor=actor,
        action="scope matrix: row edited",
        entity_type=ScopeRow.__tablename__,
        entity_id=row.id,
        before=before,
        after={"status": status, "system": row.system, "key": row.key},
        reason=note,
    )
    return row


def confirm_matrix(session: Session, bid_id: uuid.UUID, actor: Actor) -> list[ScopeRow]:
    """A named person confirms the matrix as it stands."""
    if actor.id is None:
        raise AnalysisError("a matrix is confirmed by a named person")
    rows = matrix(session, bid_id)
    if not rows:
        raise AnalysisError("there is no scope matrix yet: run the analysis first")
    now = datetime.now(UTC)
    for row in rows:
        row.confirmed_by, row.confirmed_by_id, row.confirmed_at = actor.label, actor.id, now
    session.flush()
    record_event(
        session,
        context=_context(session, bid_id),
        actor=actor,
        action="scope matrix: confirmed",
        entity_type=ScopeRow.__tablename__,
        entity_id=rows[0].id,
        after={
            "rows": len(rows),
            "unclear": sum(1 for row in rows if row.status == "unclear"),
            "edited": sum(1 for row in rows if row.source == "person"),
        },
    )
    return rows


STATUS_WORDS = {
    "included": "Included",
    "excluded": "Excluded",
    "by_others": "By others",
    "unclear": "Unclear",
}


def export_matrix(session: Session, bid: Bid) -> bytes:
    """The matrix as a workbook: a row per obligation and interface, with its clause."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    rows = matrix(session, bid.id)
    revisions = {
        revision.id: revision
        for revision in session.execute(
            select(DocumentRevision).where(DocumentRevision.bid_id == bid.id)
        ).scalars()
    }
    book = Workbook()
    if book.active is not None:  # the empty sheet a new workbook starts with
        book.remove(book.active)
    sheet = book.create_sheet("Scope matrix")
    headings = (
        "System",
        "Kind",
        "Item",
        "Status",
        "Clause",
        "Specification",
        "Revision",
        "Clause text",
        "Basis",
        "Set by",
        "Note",
        "Confirmed by",
        "Confirmed on",
    )
    sheet.append(headings)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in rows:
        revision = revisions.get(row.document_revision_id) if row.document_revision_id else None
        sheet.append(
            (
                row.system.replace("_", " "),
                row.kind,
                row.label,
                STATUS_WORDS[row.status],
                row.clause_number or "",
                (revision.title or revision.doc_key or "") if revision else "",
                (revision.revision_label or "") if revision else "",
                row.quote or "",
                row.reason,
                row.edited_by or "rule",
                row.note or "",
                row.confirmed_by or "",
                row.confirmed_at.date().isoformat() if row.confirmed_at else "",
            )
        )
    sheet.freeze_panes = "A2"
    widths = (14, 12, 44, 12, 10, 40, 10, 80, 36, 18, 30, 18, 14)
    for letter, width in zip("ABCDEFGHIJKLM", widths, strict=True):
        sheet.column_dimensions[letter].width = width
    about = book.create_sheet("About")
    about.append(("Bid", bid.human_id))
    about.append(("Client", bid.client_name))
    about.append(("Exported", datetime.now(UTC).isoformat(timespec="seconds")))
    about.append(("Rows", len(rows)))
    about.append(
        (
            "Note",
            "Statuses read from the specification are proposals until the matrix is confirmed.",
        )
    )
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()
