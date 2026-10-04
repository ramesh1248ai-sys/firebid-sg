"""The internal review pack for G2 and G3 (FR-PKG-01; P2-08).

One pack, assembled from what the platform already holds, so its figures are the
estimate's own: the cost build-up by component, the bill and labour by system, the lines
that drive the cost and the margin, the risk allowances with their treatments, the largest
variances against the client's bill, what is still open, what is unpriced, and how much of
the takeoff a person has verified. Each section says where in the platform it comes from.

The pack is rendered as a workbook and as a PDF. Both are handed to the person who asks.
"""

from __future__ import annotations

import io
from dataclasses import dataclass, field
from datetime import UTC, date, datetime
from decimal import Decimal
from typing import Any

from sqlalchemy.orm import Session

from firebid.db.models.core import Bid

TOP = 10
# What a cell says when there is no figure for it.
NONE = "-"
OPEN_CLARIFICATIONS = ("draft", "internal_review", "approved_to_issue", "issued", "responded")


@dataclass
class Section:
    key: str
    title: str
    link: str  # the page of the platform the section comes from
    columns: list[str]
    rows: list[list[str]] = field(default_factory=list)
    note: str = ""

    def as_json(self) -> dict[str, Any]:
        return {
            "key": self.key,
            "title": self.title,
            "link": self.link,
            "columns": self.columns,
            "rows": self.rows,
            "note": self.note,
        }


@dataclass
class Pack:
    bid_id: str
    human_id: str
    client: str
    reference: str
    generated_at: datetime
    priced_on: date
    sections: list[Section]
    # The headline figures, as text, for the summary and for what an approval rests on.
    figures: dict[str, str]

    def section(self, key: str) -> Section:
        return next(section for section in self.sections if section.key == key)


def _money(value: Decimal | None) -> str:
    return f"{value:,.2f}" if value is not None else NONE


def _quantity(value: Decimal | None) -> str:
    return format(value.normalize(), "f") if value is not None else NONE


def build(session: Session, bid: Bid, today: date | None = None) -> Pack:
    from firebid.services import (
        boq,
        clarifications,
        costing,
        labour,
        pricing,
        qto,
        review,
        risk,
        spec_analysis,
    )

    base = f"/bids/{bid.id}"
    built = costing.build_up(session, bid, today)
    current = boq.current_boq(session, bid.id)
    lines = boq.lines_of(session, current) if current is not None else []
    totals = pricing.boq_totals(session, current) if current is not None else None
    estimate = labour.estimate(session, bid, today)

    by_component = Section(
        "estimate_by_component",
        "Estimate by component",
        f"{base}/boq",
        ["Component", "Basis", "Source", "Amount (SGD)"],
        [
            [
                line.label,
                line.basis.replace("_", " "),
                line.source,
                _money(line.amount.amount if line.amount else None),
            ]
            for line in built.lines
        ]
        + [
            ["Direct cost", "", "", _money(built.direct.amount)],
            ["Total (excluding GST)", "", "", _money(built.total.amount)],
            [f"GST at {built.gst.rate.percent}%", "", "", _money(built.gst.gst.amount)],
            ["Total (including GST)", "", "", _money(built.gst.inclusive.amount)],
        ],
        note=f"Priced on {built.priced_on.isoformat()}."
        + (f" Not set: {', '.join(built.not_set).replace('_', ' ')}." if built.not_set else ""),
    )

    labour_by = {total.key: total for total in estimate.by_section}
    systems = sorted(set(totals.sections if totals else {}) | set(labour_by))
    by_system = Section(
        "estimate_by_system",
        "Estimate by system",
        f"{base}/boq",
        ["System", "Priced bill (SGD)", "Labour hours", "Labour cost (SGD)"],
        [
            [
                name or "(no section)",
                _money(totals.sections[name].amount)
                if totals and name in totals.sections
                else NONE,
                _quantity(labour_by[name].hours) if name in labour_by else NONE,
                _money(labour_by[name].cost) if name in labour_by else NONE,
            ]
            for name in systems
        ],
        note="" if current is not None else "No bill of quantities is built.",
    )

    priced = sorted(
        (line for line in lines if line.amount is not None),
        key=lambda line: line.amount.amount if line.amount else Decimal(0),
        reverse=True,
    )
    bill_total = sum((line.amount.amount for line in priced if line.amount), Decimal(0))
    margin = next((line for line in built.lines if line.component == "margin"), None)
    margin_amount = margin.amount.amount if margin is not None and margin.amount else None
    margin_percent = (
        (margin_amount / built.total.amount * 100).quantize(Decimal("0.1"))
        if margin_amount is not None and built.total.amount
        else None
    )
    drivers = Section(
        "cost_drivers",
        "Key cost drivers and margin",
        f"{base}/boq",
        ["Item", "Description", "Quantity", "Rate (SGD)", "Amount (SGD)", "Share of bill"],
        [
            [
                line.item_no or "",
                line.description,
                f"{_quantity(line.quantity)} {line.unit}",
                _money(line.unit_rate.amount if line.unit_rate else None),
                _money(line.amount.amount if line.amount else None),
                f"{(line.amount.amount / bill_total * 100).quantize(Decimal('0.1'))}%"
                if line.amount and bill_total
                else NONE,
            ]
            for line in priced[:TOP]
        ],
        note=(
            f"Margin: SGD {_money(margin_amount)} ({margin_percent}% of the total excluding GST)."
            if margin_amount is not None
            else "Margin is not set."
        ),
    )

    found = risk.risks(session, bid.id)
    allowance = sum(
        (row.cost_allowance.amount for row in found if row.cost_allowance is not None), Decimal(0)
    )
    risks = Section(
        "risk_allowances",
        "Risk allowances and treatments",
        f"{base}/risk",
        [
            "Risk",
            "Category",
            "Treatment",
            "Owner",
            "Status",
            "Allowance (SGD)",
            "Man-hours",
            "Impact",
        ],
        [
            [
                row.title,
                row.category.replace("_", " "),
                row.treatment or "none",
                row.owner or "",
                row.status,
                _money(row.cost_allowance.amount if row.cost_allowance is not None else None),
                _quantity(row.programme_hours),
                row.impact_state,
            ]
            for row in found
        ],
        note=f"Allowances total SGD {_money(allowance)}."
        + (
            f" {sum(1 for r in found if r.treatment is None)} risk(s) have no treatment."
            if any(r.treatment is None for r in found)
            else ""
        ),
    )

    flagged = [
        row
        for row in (boq.reconciliation(session, bid.id) if current is not None else [])
        if row.flagged
    ]
    flagged.sort(key=lambda row: abs(row.variance_percent or Decimal(10**6)), reverse=True)
    variances = Section(
        "variances",
        "Top quantity variances against the client's bill",
        f"{base}/boq",
        ["Client ref", "Description", "Client quantity", "Measured", "Variance", "%"],
        [
            [
                row.client_ref or "(not in the bill)",
                row.client_description or row.line_description or "",
                f"{_quantity(row.client_quantity)} {row.client_unit or ''}".strip(),
                f"{_quantity(row.measured_quantity)} {row.unit or ''}".strip(),
                _quantity(row.variance),
                f"{row.variance_percent}%" if row.variance_percent is not None else NONE,
            ]
            for row in flagged[:TOP]
        ],
        note=f"{len(flagged)} flagged in all." if flagged else "None flagged.",
    )

    asked = [
        row for row in clarifications.register(session, bid.id) if row.state in OPEN_CLARIFICATIONS
    ]
    issues = spec_analysis.list_issues(session, bid.id, "open")
    open_items = Section(
        "open_items",
        "Open clarifications and issues",
        f"{base}/clarifications",
        ["Kind", "Reference", "Subject", "Status", "Due"],
        [
            [
                "clarification",
                clarifications.label(row),
                row.subject,
                clarifications.STATE_WORDS[row.state],
                row.due_at.date().isoformat() if row.due_at else "",
            ]
            for row in asked
        ]
        + [["specification issue", issue.severity, issue.title, "open", ""] for issue in issues],
        note=f"{len(asked)} clarification(s) unresolved; "
        f"{len(issues)} specification issue(s) open.",
    )

    prices = pricing.line_prices(session, bid, current) if current is not None else {}
    waiting = [
        line
        for line in lines
        if prices.get(line.id) and prices[line.id].status in ("unpriced", "proposed")
    ]
    unpriced = Section(
        "unpriced_lines",
        "Unpriced lines",
        f"{base}/boq",
        ["Item", "Description", "Quantity", "Status"],
        [
            [
                line.item_no or "",
                line.description,
                f"{_quantity(line.quantity)} {line.unit}",
                prices[line.id].status,
            ]
            for line in waiting
        ],
        note=f"{len(waiting)} line(s) carry no price and are not in the total."
        if waiting
        else "Every line is priced.",
    )

    covered = review.coverage(session, bid.id)
    blockers = qto.g1_blockers(session, bid.id)
    coverage = Section(
        "g1_coverage",
        "G1 coverage",
        f"{base}/workbench",
        ["Measure", "Value"],
        [
            ["Items verified", f"{covered['items_verified']} of {covered['items_total']}"],
            ["Items verified (%)", f"{covered['items_percent']}"],
            ["Value verified (%)", f"{covered['value_percent']}"],
            ["Value basis", str(covered["value_basis"])],
            ["Coverage policy (%)", f"{covered['policy_percent']}"],
            ["Unresolved duplicate groups", str(len(blockers.unresolved_groups))],
            ["Items with incomplete evidence", str(len(blockers.incomplete_items))],
        ],
        note="Coverage meets the policy." if covered["met"] else "Coverage is below the policy.",
    )

    return Pack(
        bid_id=str(bid.id),
        human_id=bid.human_id,
        client=bid.client_name,
        reference=bid.tender_reference,
        generated_at=datetime.now(UTC),
        priced_on=built.priced_on,
        sections=[
            by_component,
            by_system,
            drivers,
            risks,
            variances,
            open_items,
            unpriced,
            coverage,
        ],
        figures={
            "direct": str(built.direct.amount),
            "total_excluding_gst": str(built.total.amount),
            "gst": str(built.gst.gst.amount),
            "total_including_gst": str(built.gst.inclusive.amount),
            "margin": str(margin_amount) if margin_amount is not None else "",
            "priced_bill": str(totals.grand.amount) if totals else "0.00",
            "labour_hours": str(estimate.hours),
            "labour_cost": str(estimate.cost.amount),
            "risk_allowances": str(allowance.quantize(Decimal("0.01"))),
            "unpriced_lines": str(len(waiting)),
            "open_clarifications": str(len(asked)),
            "open_issues": str(len(issues)),
            "items_verified_percent": str(covered["items_percent"]),
        },
    )


FIGURE_LABELS = {
    "direct": "Direct cost (SGD)",
    "total_excluding_gst": "Total excluding GST (SGD)",
    "gst": "GST (SGD)",
    "total_including_gst": "Total including GST (SGD)",
    "margin": "Margin (SGD)",
    "priced_bill": "Priced bill (SGD)",
    "labour_hours": "Labour hours",
    "labour_cost": "Labour cost (SGD)",
    "risk_allowances": "Risk allowances (SGD)",
    "unpriced_lines": "Unpriced lines",
    "open_clarifications": "Open clarifications",
    "open_issues": "Open specification issues",
    "items_verified_percent": "Takeoff items verified (%)",
}


def as_workbook(pack: Pack) -> bytes:
    """The pack as a workbook: a summary sheet, then a sheet a section."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    book = Workbook()
    summary = book.active
    assert summary is not None  # noqa: S101 - a new workbook has one
    summary.title = "Summary"
    summary.append([f"Review pack: {pack.human_id}"])
    summary.append([f"{pack.client} · {pack.reference}"])
    summary.append(
        [f"Generated {pack.generated_at:%Y-%m-%d %H:%M} UTC · priced on {pack.priced_on}"]
    )
    summary.append([])
    for key, label in FIGURE_LABELS.items():
        summary.append([label, pack.figures.get(key, "")])
    summary.append([])
    summary.append(["Section", "In the platform"])
    for section in pack.sections:
        summary.append([section.title, section.link])
    summary["A1"].font = Font(bold=True, size=14)
    summary.column_dimensions["A"].width = 46
    summary.column_dimensions["B"].width = 46
    for section in pack.sections:
        sheet = book.create_sheet(section.title[:31])
        sheet.append([section.title])
        sheet.append([f"In the platform: {section.link}"])
        sheet.append([section.note])
        sheet.append(section.columns)
        for row in section.rows:
            sheet.append(row)
        sheet["A1"].font = Font(bold=True, size=12)
        for cell in sheet[4]:
            cell.font = Font(bold=True)
        for index in range(1, len(section.columns) + 1):
            sheet.column_dimensions[get_column_letter(index)].width = 28
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def _fit(text: str, width: int) -> str:
    return text if len(text) <= width else text[: width - 1] + "…"


def as_pdf(pack: Pack) -> bytes:
    """The pack as a PDF: the summary, then each section as a table, a page or more each."""
    import matplotlib

    matplotlib.use("Agg")
    from matplotlib import pyplot as plt
    from matplotlib.backends.backend_pdf import PdfPages

    buffer = io.BytesIO()
    per_page = 38

    def page(pages: PdfPages, title: str, link: str, lines: list[str], note: str) -> None:
        for start in range(0, max(len(lines), 1), per_page):
            figure = plt.figure(figsize=(11.69, 8.27))
            figure.text(0.05, 0.94, title, fontsize=14, weight="bold")
            figure.text(0.05, 0.91, f"In the platform: {link}", fontsize=8, color="#444444")
            if note:
                figure.text(0.05, 0.885, _fit(note, 150), fontsize=8)
            y = 0.85
            for line in lines[start : start + per_page]:
                figure.text(0.05, y, line, fontsize=7.5, family="monospace")
                y -= 0.021
            figure.text(
                0.05,
                0.03,
                f"{pack.human_id} · {pack.client} · "
                f"generated {pack.generated_at:%Y-%m-%d %H:%M} UTC",
                fontsize=7,
                color="#666666",
            )
            pages.savefig(figure)  # type: ignore[no-untyped-call]
            plt.close(figure)

    with PdfPages(buffer) as pages:
        page(
            pages,
            f"Review pack: {pack.human_id}",
            f"/bids/{pack.bid_id}",
            [f"{label:<34} {pack.figures.get(key, '')}" for key, label in FIGURE_LABELS.items()],
            f"{pack.client} · {pack.reference} · priced on {pack.priced_on}",
        )
        for section in pack.sections:
            width = max(10, 150 // max(len(section.columns), 1))
            lines = [
                " ".join(_fit(str(cell), width - 1).ljust(width) for cell in row)
                for row in [section.columns, *section.rows]
            ]
            page(pages, section.title, section.link, lines, section.note)
    return buffer.getvalue()
