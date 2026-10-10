"""
Generate the FireBid SG user manual as a Word document from docs/user-manual/*.md.
Executed with: uv run --project backend python scripts/generate_user_manual_docx.py

The Markdown files stay the source: edit them, then run this again. The table of contents
is a Word field: open the document in Word and press F9 (or answer "Yes" when asked to
update fields) to fill it in.
"""

from __future__ import annotations

import re
from collections.abc import Iterator
from datetime import date
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
MANUAL = ROOT / "docs" / "user-manual"
OUT = MANUAL / "FireBid_SG_User_Manual.docx"

# The parts, in the order a reader needs them. The reference's own contents list and
# opening are replaced by this document's.
PARTS = (
    ("Part 1 · A tender from upload to award", "happy-path.md", None),
    ("Part 2 · When something is not right", "exceptions.md", None),
    ("Part 3 · Screen reference", "README.md", "## 1. Signing in"),
)
PART_NAMES = {
    "happy-path.md": "Part 1",
    "exceptions.md": "Part 2",
    "README.md": "Part 3",
}

NAVY = RGBColor(0x0F, 0x1E, 0x3D)
RED = RGBColor(0xB9, 0x1C, 0x1C)
GREY = RGBColor(0x55, 0x60, 0x70)
HEADER_FILL = "E8EDF5"
TEXT_WIDTH = Cm(17.0)
MAX_IMAGE_HEIGHT = Cm(19.0)

INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|!?\[[^\]]*\]\([^)]*\))")
LINK = re.compile(r"\[([^\]]*)\]\(([^)]*)\)")
IMAGE = re.compile(r"^!\[([^\]]*)\]\(([^)]+)\)$")
ORDERED = re.compile(r"^(\d+)\.\s+(.*)$")


def shade(cell: Any, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def field(paragraph: Any, instruction: str, placeholder: str = "") -> None:
    """A Word field (a page number, a table of contents) that Word fills in."""
    run = paragraph.add_run()
    for kind, text in (("begin", None), (None, instruction), ("separate", None)):
        if kind:
            mark = OxmlElement("w:fldChar")
            mark.set(qn("w:fldCharType"), kind)
            if kind == "begin":
                mark.set(qn("w:dirty"), "true")
            run._r.append(mark)
        else:
            code = OxmlElement("w:instrText")
            code.set(qn("xml:space"), "preserve")
            code.text = text
            run._r.append(code)
    if placeholder:
        paragraph.add_run(placeholder)
    end = paragraph.add_run()
    mark = OxmlElement("w:fldChar")
    mark.set(qn("w:fldCharType"), "end")
    end._r.append(mark)


def write_inline(paragraph: Any, text: str, *, size: Pt | None = None, bold: bool = False) -> None:
    """Text with **bold**, `code` and [links](...): a link keeps its words, and one to
    another part of the manual says which part."""
    for piece in INLINE.split(text):
        if not piece:
            continue
        run = None
        if piece.startswith("**") and piece.endswith("**"):
            run = paragraph.add_run(piece[2:-2])
            run.bold = True
        elif piece.startswith("`") and piece.endswith("`"):
            run = paragraph.add_run(piece[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
        elif (found := LINK.fullmatch(piece)) is not None:
            words, target = found.group(1), found.group(2)
            part = PART_NAMES.get(target.split("#")[0])
            run = paragraph.add_run(f"{words} ({part})" if part else words)
        else:
            run = paragraph.add_run(piece)
        if bold:
            run.bold = True
        if size is not None and run.font.size is None:
            run.font.size = size


def blocks(lines: list[str]) -> Iterator[tuple[str, Any]]:
    """The Markdown as blocks: heading, paragraph, image, table, ordered or bullet list."""
    index = 0
    while index < len(lines):
        line = lines[index]
        stripped = line.strip()
        if not stripped:
            index += 1
            continue
        if stripped.startswith("#"):
            level = len(stripped) - len(stripped.lstrip("#"))
            yield "heading", (level, stripped[level:].strip())
            index += 1
        elif IMAGE.match(stripped):
            alt, path = IMAGE.findall(stripped)[0]
            yield "image", (alt, path)
            index += 1
        elif stripped.startswith("|"):
            rows = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                cells = [c.strip() for c in lines[index].strip().strip("|").split("|")]
                if not all(set(c) <= set("-: ") for c in cells):
                    rows.append(cells)
                index += 1
            yield "table", rows
        elif ORDERED.match(stripped) or stripped.startswith("- "):
            ordered = bool(ORDERED.match(stripped))
            items: list[tuple[str, str]] = []
            while index < len(lines):
                current = lines[index]
                text = current.strip()
                numbered = ORDERED.match(text)
                if ordered and numbered and not current.startswith(" "):
                    items.append((numbered.group(1), numbered.group(2)))
                elif not ordered and text.startswith("- ") and not current.startswith(" "):
                    items.append(("", text[2:]))
                elif text and current.startswith(" ") and items:
                    items[-1] = (items[-1][0], items[-1][1] + " " + text)
                else:
                    break
                index += 1
            yield ("ordered" if ordered else "bullets"), items
        else:
            words = []
            while index < len(lines):
                text = lines[index].strip()
                if (
                    not text
                    or text.startswith(("#", "|", "- "))
                    or IMAGE.match(text)
                    or ORDERED.match(text)
                ):
                    break
                words.append(text)
                index += 1
            yield "paragraph", " ".join(words)


def add_image(document: Any, alt: str, relative: str) -> None:
    from PIL import Image

    path = MANUAL / relative
    with Image.open(path) as image:
        width, height = image.size
    shown: int = TEXT_WIDTH
    if height / width * shown > MAX_IMAGE_HEIGHT:
        shown = int(MAX_IMAGE_HEIGHT * width / height)
    picture = document.add_paragraph()
    picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture.paragraph_format.space_before = Pt(6)
    picture.paragraph_format.space_after = Pt(2)
    picture.paragraph_format.keep_with_next = True
    inline = picture.add_run().add_picture(str(path), width=shown)
    # A thin border, so a white screen does not run into the page.
    properties = inline._inline.graphic.graphicData.pic.spPr
    outline = OxmlElement("a:ln")
    outline.set("w", "6350")
    fill = OxmlElement("a:solidFill")
    colour = OxmlElement("a:srgbClr")
    colour.set("val", "C5CCD6")
    fill.append(colour)
    outline.append(fill)
    properties.append(outline)
    caption = document.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_after = Pt(10)
    run = caption.add_run(alt)
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = GREY


def add_table(document: Any, rows: list[list[str]]) -> None:
    columns = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=columns)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r, row in enumerate(rows):
        for c in range(columns):
            cell = table.cell(r, c)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(2)
            paragraph.paragraph_format.space_before = Pt(2)
            write_inline(paragraph, row[c] if c < len(row) else "", size=Pt(9.5), bold=r == 0)
            if r == 0:
                shade(cell, HEADER_FILL)
    # The header row repeats where a table runs over a page.
    header = table.rows[0]._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    header.append(repeat)
    document.add_paragraph().paragraph_format.space_after = Pt(4)


def add_list(document: Any, items: list[tuple[str, str]], ordered: bool) -> None:
    for number, text in items:
        if ordered:
            # Numbered by hand: Word's own numbering would run on from one list to the next.
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Cm(0.9)
            paragraph.paragraph_format.first_line_indent = Cm(-0.7)
            paragraph.paragraph_format.tab_stops.add_tab_stop(Cm(0.9))
            paragraph.add_run(f"{number}.\t")
        else:
            paragraph = document.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(3)
        write_inline(paragraph, text)


def add_part(document: Any, title: str, name: str, start: str | None) -> None:
    lines = (MANUAL / name).read_text(encoding="utf-8").splitlines()
    if start is not None:
        lines = lines[next(i for i, line in enumerate(lines) if line.strip() == start) :]
    document.add_page_break()
    document.add_heading(title, level=1)
    for kind, value in blocks(lines):
        if kind == "heading":
            level, text = value
            if level == 1:
                continue  # the part's title stands for it
            document.add_heading(text, level=min(level, 4))
        elif kind == "paragraph":
            write_inline(document.add_paragraph(), value)
        elif kind == "image":
            add_image(document, *value)
        elif kind == "table":
            add_table(document, value)
        else:
            add_list(document, value, kind == "ordered")


def style(document: Any) -> None:
    section = document.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width, section.page_height = Cm(21.0), Cm(29.7)
    section.left_margin = section.right_margin = Cm(2.0)
    section.top_margin, section.bottom_margin = Cm(2.2), Cm(2.0)
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12
    for name, size, colour, before in (
        ("Heading 1", 20, NAVY, 0),
        ("Heading 2", 14, NAVY, 14),
        ("Heading 3", 12, NAVY, 10),
        ("Heading 4", 11, GREY, 8),
    ):
        heading = document.styles[name]
        heading.font.name = "Calibri"
        heading.font.size = Pt(size)
        heading.font.bold = True
        heading.font.color.rgb = colour
        heading.paragraph_format.space_before = Pt(before)
        heading.paragraph_format.space_after = Pt(6)
        heading.paragraph_format.keep_with_next = True
        # The theme's heading font would override the name set above.
        fonts = heading.element.get_or_add_rPr().get_or_add_rFonts()
        for attribute in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            if fonts.get(qn(attribute)) is not None:
                del fonts.attrib[qn(attribute)]
        fonts.set(qn("w:ascii"), "Calibri")
        fonts.set(qn("w:hAnsi"), "Calibri")
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run("FireBid SG · User Manual")
    run.font.size = Pt(8.5)
    run.font.color.rgb = GREY
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    lead = footer.add_run("Page ")
    lead.font.size = Pt(8.5)
    lead.font.color.rgb = GREY
    field(footer, "PAGE", "1")
    section.different_first_page_header_footer = True


def cover(document: Any) -> None:
    for _ in range(7):
        document.add_paragraph()
    brand = document.add_paragraph()
    run = brand.add_run("FireBid SG")
    run.font.size = Pt(40)
    run.bold = True
    run.font.color.rgb = RED
    title = document.add_paragraph()
    run = title.add_run("User Manual")
    run.font.size = Pt(28)
    run.font.color.rgb = NAVY
    rule = document.add_paragraph()
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for key, value in (("val", "single"), ("sz", "12"), ("space", "1"), ("color", "B91C1C")):
        bottom.set(qn(f"w:{key}"), value)
    border.append(bottom)
    rule._p.get_or_add_pPr().append(border)
    about = document.add_paragraph()
    about.paragraph_format.space_before = Pt(14)
    run = about.add_run(
        "Tendering for fire protection contractors: from uploaded tender documents to a "
        "priced, reviewed bid."
    )
    run.font.size = Pt(13)
    run.font.color.rgb = GREY
    for _ in range(12):
        document.add_paragraph()
    stamp = document.add_paragraph()
    run = stamp.add_run(
        f"Issued {date.today().strftime('%d %B %Y')} · screenshots from synthetic test bids"
    )
    run.font.size = Pt(10)
    run.font.color.rgb = GREY


def front(document: Any) -> None:
    document.add_page_break()
    document.add_heading("About this manual", level=1)
    document.add_paragraph(
        "FireBid SG takes a fire-protection tender from uploaded documents to a priced, "
        "reviewed bid. Nothing the platform proposes counts until a person accepts it, and "
        "every quantity and price traces back to where it came from."
    )
    document.add_paragraph("The manual has three parts:")
    add_table(
        document,
        [
            ["Part", "Read it to"],
            [
                "**Part 1 · A tender from upload to award**",
                "Follow one bid through every step, with who does each and what they press",
            ],
            [
                "**Part 2 · When something is not right**",
                "Find out what a warning or a blocked step means, and what to do",
            ],
            ["**Part 3 · Screen reference**", "Look up what a screen offers"],
        ],
    )
    document.add_paragraph(
        "The screenshots were taken on 8 and 10 October 2026 from synthetic test bids on a "
        "local installation. No client's tender appears in them."
    )
    heading = document.add_paragraph()
    heading.paragraph_format.space_before = Pt(14)
    run = heading.add_run("Contents")
    run.bold = True
    run.font.size = Pt(14)
    run.font.color.rgb = NAVY
    contents = document.add_paragraph()
    field(
        contents,
        r'TOC \o "1-2" \h \z \u',
        "Open this document in Word and press F9 to fill in the contents.",
    )


def main() -> None:
    document = Document()
    style(document)
    document.core_properties.title = "FireBid SG User Manual"
    document.core_properties.subject = "User manual"
    document.core_properties.author = "FireBid SG"
    cover(document)
    front(document)
    for title, name, start in PARTS:
        add_part(document, title, name, start)
    document.save(str(OUT))
    print(f"written {OUT.relative_to(ROOT)} ({OUT.stat().st_size // 1024} KB)")


if __name__ == "__main__":
    main()
